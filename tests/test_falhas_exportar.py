# -*- coding: utf-8 -*-
"""Exportar Excel da aba Strings e trackers (Levi, 20/10/2026: "1.4 botão de download"; o "Exportar Excel" que ele já
usa nas outras telas). A TELA manda as linhas que mostra e o servidor só formata (POST /api/painel/falhas/exportar):
número e data viram célula de número e de data de verdade, no padrão das exportações da casa (cabeçalho navy, painel
congelado, autofiltro, largura de coluna). Pedido fora do formato ou dos limites: 400 com o motivo. O gestor também
baixa (consulta, não gravação)."""
from datetime import datetime
from io import BytesIO

import openpyxl
import pytest

import app
import leitura_nexus as ln
from test_porta_nexus import _sem_rede, cli, entrar  # noqa: F401  (fixtures e passos da porta, para o gestor)

COLS = [{"titulo": "Cliente", "tipo": "texto"}, {"titulo": "Usina", "tipo": "texto"},
        {"titulo": "Strings", "tipo": "inteiro"}, {"titulo": "Parou", "tipo": "datahora"},
        {"titulo": "Voltou", "tipo": "datahora"}, {"titulo": "Dias", "tipo": "inteiro"},
        {"titulo": "Perda (kWh)", "tipo": "numero"}, {"titulo": "Dia", "tipo": "data"}]
# Inhapi 3.3 e Guaratinguetá V 1.3 como no pacote de outubro (20/10 13:45); a 3ª linha é o tracker que voltou em hora
# que o registro não guarda (só o dia) e um texto que o Excel leria como fórmula
LINHAS = [["Thopen", "Inhapi", 4, "2026-10-01 06:30", None, 20, 8024.8, "2026-10-01"],
          ["Thopen", "Guaratinguetá V", 10, "2026-10-07 08:50", "2026-10-14 11:00", 8, 1273.8, "2026-10-07"],
          ["Athon", "=1+1", 1, "2026-10-05", "2026-10-06", 2, 12, None]]
CORPO = {"nome": "strings_2026-10_Guaratinguetá V",
         "abas": [{"nome": "Ocorrências", "colunas": COLS, "linhas": LINHAS},
                  {"nome": "Dia a dia", "colunas": [{"titulo": "Dia", "tipo": "data"}, {"titulo": "Perda", "tipo": "numero"}],
                   "linhas": [["2026-10-07", 382.4], ["2026-10-08", 329.7]]}]}


@pytest.fixture
def c(monkeypatch):
    monkeypatch.setattr(app, "DASH_PASSWORD", "")
    return app.app.test_client()


def _abre(r):
    return openpyxl.load_workbook(BytesIO(r.data))


def test_o_excel_tem_numero_e_data_de_verdade_no_padrao_da_casa(c):
    r = c.post("/api/painel/falhas/exportar", json=CORPO)
    assert r.status_code == 200 and r.mimetype.endswith("spreadsheetml.sheet")
    assert r.headers["Content-Disposition"] == 'attachment; filename="strings_2026-10_Guaratingueta_V.xlsx"'
    wb = _abre(r)
    assert wb.sheetnames == ["Ocorrências", "Dia a dia"]
    ws = wb["Ocorrências"]
    assert [x.value for x in ws[1]] == [x["titulo"] for x in COLS]
    assert ws["A1"].fill.fgColor.rgb == "001A1B2E" and ws["A1"].font.bold and ws["A1"].font.color.rgb == "00FFFFFF"
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:H4"
    assert ws.column_dimensions["B"].width >= len("Guaratinguetá V")
    assert (ws["C2"].value, ws["C2"].number_format) == (4, "0")
    assert (ws["D2"].value, ws["D2"].number_format) == (datetime(2026, 10, 1, 6, 30), "dd/mm/yyyy hh:mm")
    assert ws["E2"].value is None
    assert (ws["G2"].value, ws["G2"].number_format) == (8024.8, "#,##0.0")
    assert ws["G4"].value == 12 and ws["G4"].data_type == "n" and ws["G4"].number_format == "#,##0.0"   # o 12 inteiro
    assert (ws["H2"].value, ws["H2"].number_format) == (datetime(2026, 10, 1), "dd/mm/yyyy")
    # a datahora que veio só com o dia vira data, sem 00:00 inventado; o texto com "=" fica texto
    assert (ws["D4"].value, ws["D4"].number_format) == (datetime(2026, 10, 5), "dd/mm/yyyy")
    assert ws["B4"].value == "=1+1" and ws["B4"].data_type == "s"
    assert ws["H4"].value is None
    assert [[x.value for x in lin] for lin in wb["Dia a dia"]["A2:B3"]] == [[datetime(2026, 10, 7), 382.4],
                                                                          [datetime(2026, 10, 8), 329.7]]


def test_nome_do_arquivo_saneado():
    assert app._falhas_export_nome_arquivo("trackers_2026-10_São Bento do Una") == "trackers_2026-10_Sao_Bento_do_Una.xlsx"
    assert app._falhas_export_nome_arquivo('a"b/../c.xlsx') == "a_b_.._c.xlsx"
    assert app._falhas_export_nome_arquivo("") == "falhas.xlsx"


def _com(**troca):
    aba = dict(CORPO["abas"][0], **troca)
    return {"nome": "x", "abas": [aba]}


@pytest.mark.parametrize("corpo, trecho", [
    ({"nome": "x", "abas": [dict(CORPO["abas"][1], nome=f"A{i}") for i in range(5)]}, "de 1 a 4 abas"),
    ({"nome": "x", "abas": []}, "de 1 a 4 abas"),
    (_com(nome="Ocorrências/outubro"), "regra do Excel"),
    (_com(nome="x" * 32), "regra do Excel"),
    ({"nome": "x", "abas": [CORPO["abas"][1], dict(CORPO["abas"][1], nome="dia A DIA")]}, "nome repetido"),
    (_com(colunas=[{"titulo": f"c{i}", "tipo": "texto"} for i in range(41)], linhas=[]), "de 1 a 40 colunas"),
    (_com(colunas=[{"titulo": "c", "tipo": "moeda"}], linhas=[]), "coluna 1"),
    (_com(linhas=[LINHAS[0][:-1]]), "8 valores"),
    (_com(linhas=[["Thopen", "x" * 501] + LINHAS[0][2:]]), "mais de 500 caracteres"),
    (_com(linhas=[LINHAS[0][:6] + ["12,5", "2026-10-01"]]), "numero"),
    (_com(linhas=[LINHAS[0][:2] + [1.5] + LINHAS[0][3:]]), "inteiro, veio 1.5"),
    (_com(linhas=[LINHAS[0][:7] + ["01/10/2026"]]), "AAAA-MM-DD"),
    (_com(linhas=[LINHAS[0][:3] + ["2026-10-01 6h30"] + LINHAS[0][4:]]), "AAAA-MM-DD HH:MM"),
    (_com(linhas=[LINHAS[0][:2] + [True] + LINHAS[0][3:]]), "verdadeiro/falso"),
    ({"nome": "x"}, "mande {nome, abas"),
])
def test_fora_do_formato_ou_dos_limites_e_400_com_o_motivo(c, corpo, trecho):
    r = c.post("/api/painel/falhas/exportar", json=corpo)
    assert r.status_code == 400 and r.get_json()["ok"] is False and trecho in r.get_json()["motivo"], r.get_json()


def test_mais_de_30_mil_linhas_somando_as_abas_e_400(c):
    col = [{"titulo": "n", "tipo": "inteiro"}]
    corpo = {"nome": "x", "abas": [{"nome": "a", "colunas": col, "linhas": [[1]] * 20000},
                                   {"nome": "b", "colunas": col, "linhas": [[1]] * 10001}]}
    r = c.post("/api/painel/falhas/exportar", json=corpo)
    assert r.status_code == 400 and "30000 linhas" in r.get_json()["motivo"]
    corpo["abas"][1]["linhas"] = [[1]] * 10000
    assert c.post("/api/painel/falhas/exportar", json=corpo).status_code == 200


def test_mais_de_600_mil_celulas_e_400(c):
    # o teto de células (revisão de 20/10/2026): 30 mil linhas × 40 colunas levavam ~62 s com a thread do waitress presa
    cols = [{"titulo": f"c{j}", "tipo": "numero"} for j in range(40)]
    corpo = {"nome": "x", "abas": [{"nome": "a", "colunas": cols, "linhas": [[1.0] * 40] * 15001}]}
    r = c.post("/api/painel/falhas/exportar", json=corpo)
    assert r.status_code == 400 and "600.000 células" in r.get_json()["motivo"]


def test_titulo_com_igual_fica_texto_e_caractere_de_controle_e_400(c):
    # revisão de 20/10/2026: o título "=HYPERLINK(...)" saía como fórmula em A1; o \x01 dava 500 em HTML, sem motivo
    corpo = {"nome": "x", "abas": [{"nome": "a", "colunas": [{"titulo": '=HYPERLINK("http://x","clique")', "tipo": "texto"}],
                                    "linhas": [["ok"]]}]}
    r = c.post("/api/painel/falhas/exportar", json=corpo)
    assert r.status_code == 200
    a1 = _abre(r)["a"]["A1"]
    assert a1.data_type == "s" and a1.value == '=HYPERLINK("http://x","clique")'
    for troca in ({"linhas": [["abc\x01def"]]}, {"linhas": [["linha\x0bvertical"]]},
                  {"colunas": [{"titulo": "t\x01x", "tipo": "texto"}]}, {"nome": "a\x01"}):
        aba = dict(corpo["abas"][0], **troca)
        r = c.post("/api/painel/falhas/exportar", json={"nome": "x", "abas": [aba]})
        assert r.status_code == 400 and r.get_json()["ok"] is False, troca


def test_corpo_grande_demais_ou_que_nao_e_json(c, monkeypatch):
    monkeypatch.setattr(app, "FALHAS_EXPORT_CORPO_MAX", 200)
    r = c.post("/api/painel/falhas/exportar", json=CORPO)
    assert r.status_code == 413 and "corpo de até" in r.get_json()["motivo"]
    monkeypatch.setattr(app, "FALHAS_EXPORT_CORPO_MAX", 16 * 1024 * 1024)
    r = c.post("/api/painel/falhas/exportar", data="isto não é json", content_type="application/json")
    assert r.status_code == 400


def test_o_gestor_tambem_baixa_e_a_rota_e_consulta_nao_gravacao(cli):
    assert ln.permitido_gestor("POST", "/api/painel/falhas/exportar")
    assert not any(g.rstrip("/") == "/api/painel/falhas/exportar" for g in ln.GRAVACOES_GESTOR)
    entrar(cli, "gestor@exemplo.com")
    r = cli.post("/api/painel/falhas/exportar", json=CORPO)
    assert r.status_code == 200 and _abre(r).sheetnames == ["Ocorrências", "Dia a dia"]
