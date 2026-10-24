# -*- coding: utf-8 -*-
"""A aba Registros (Levi, 23/10/2026: "Quero uma visão de drill down / Carteira, usinas, ao expandir usinas aparece as
colunas que tem no anexo / a pessoa vai poder exportar, ao exportar vai ter só usina/data/geração/irradiação/
disponibilidade / Geração e irradiação continua sendo da BD_Thopen, disponibilidade vai ser da forma nova!"). Dados,
meta e disponibilidade falsos, nada de rede."""
import datetime as dt
import io
import os
import sys

import pytest
from openpyxl import load_workbook

AQUI = os.path.dirname(__file__)
sys.path.insert(0, os.path.join(AQUI, "..", "thopen"))
import acessos  # noqa: E402
import dashboard_thopen as dth  # noqa: E402
import fonte_operacao  # noqa: E402
import operacao_mes  # noqa: E402

OUT = [{"data": dt.date(2026, 10, 1), "ger": 900.0, "ipoa": 6.0, "disp": 0.0, "com": "x"},
       {"data": dt.date(2026, 10, 2), "ger": 800.0, "ipoa": 5.5, "disp": 0.0, "com": None},
       {"data": dt.date(2026, 10, 3), "ger": None, "ipoa": None, "disp": None, "com": None},     # dia sem nada: fora
       {"data": dt.date(2026, 10, 4), "ger": 0.0, "ipoa": 4.0, "disp": None, "com": None},
       {"data": dt.date(2026, 10, 5), "ger": 0.0, "ipoa": None, "disp": None, "com": None}]     # ainda não coletado
SET = [{"data": dt.date(2026, 9, d), "ger": 1000.0, "ipoa": 5.0, "disp": 0.99, "com": None} for d in range(1, 4)]
DISP = {dt.date(2026, 10, 1): 100.0, dt.date(2026, 10, 2): 80.0, dt.date(2026, 10, 4): 90.0}
# o 0 do BD depois do último dia coletado é "sem dado"; no meio do mês, geração zero (o 02/10 da Piancó abaixo)
PIANCO = [dict(OUT[0]), dict(OUT[1], ger=0.0, ipoa=None), dict(OUT[0], data=dt.date(2026, 10, 3)),
          dict(OUT[1], data=dt.date(2026, 10, 4), ger=0.0, ipoa=None)]


@pytest.fixture
def cli(monkeypatch):
    monkeypatch.delenv("THOPEN_LOGIN", raising=False)
    monkeypatch.delenv("THOPEN_OPERACAO", raising=False)
    regs = {"Altair": OUT + SET, "Assis": SET, "Piancó 1": PIANCO}
    monkeypatch.setattr(dth, "_daily_records", lambda u: regs.get(u, []))
    monkeypatch.setattr(fonte_operacao, "dados", lambda *a, **k: {"disponibilidade_dia": [1]})
    monkeypatch.setattr(operacao_mes, "disp_por_dia",
                        lambda dados, u, a, m: dict(DISP) if (u, a, m) == ("Altair", 2026, 10) else {})
    monkeypatch.setattr(dth, "_wb", lambda: None)
    monkeypatch.setattr(dth, "_state", dict(dth._state, daily={u: ("x", "A1") for u in regs}))
    monkeypatch.setattr(dth, "_polaris_records", lambda: {})
    monkeypatch.setattr(dth, "_sheet_records", lambda: {})
    monkeypatch.setattr(dth, "CARTEIRAS", dict(dth.CARTEIRAS, Thopen=["Altair", "Assis", "Piancó 1"]))
    return dth.app.test_client()


def _usina(r, nome):
    return next(u for u in r["usinas"] if u["usina"] == nome)


def test_outubro_geracao_e_irradiacao_do_bd_e_disponibilidade_das_os(cli):
    r = cli.get("/api/t/registros?carteira=Thopen&ano=2026&mes=10").get_json()
    assert r["carteira"] == "Thopen" and r["fonte_disp"] == "os"
    alt = _usina(r, "Altair")
    assert [d["data"] for d in alt["dias"]] == ["2026-10-01", "2026-10-02", "2026-10-04"]
    assert [d["ger"] for d in alt["dias"]] == [900.0, 800.0, 0.0]
    assert [d["ipoa"] for d in alt["dias"]] == [6.0, 5.5, 4.0]
    # a disponibilidade das OS (0–100 na plataforma) vira fração, como a do BD: o 0,0 "não preenchido" do BD não entra
    assert [d["disp"] for d in alt["dias"]] == pytest.approx([1.0, 0.8, 0.9])
    assert alt["ger"] == 1700.0 and alt["ipoa"] == pytest.approx(15.5) and alt["disp"] == pytest.approx(0.9)
    assert [u["usina"] for u in r["usinas"]] == ["Altair", "Piancó 1"]            # a Assis não tem outubro
    pia = _usina(r, "Piancó 1")
    assert all(d["disp"] is None for d in pia["dias"])                            # sem OS publicada: vazio, não zero
    assert [(d["data"][-2:], d["ger"]) for d in pia["dias"]] == [("01", 900.0), ("02", 0.0), ("03", 900.0)]


def test_dia_ainda_nao_coletado_sai_sem_geracao_e_fica_pela_disponibilidade(cli, monkeypatch):
    disp = {**DISP, dt.date(2026, 10, 5): 100.0}
    monkeypatch.setattr(operacao_mes, "disp_por_dia",
                        lambda dados, u, a, m: dict(disp) if (u, a, m) == ("Altair", 2026, 10) else {})
    r = cli.get("/api/t/registros?carteira=Thopen&ano=2026&mes=10").get_json()
    assert _usina(r, "Altair")["dias"][-1] == {"data": "2026-10-05", "ger": None, "ipoa": None, "disp": 1.0}


def test_setembro_com_a_disponibilidade_do_bd(cli):
    r = cli.get("/api/t/registros?carteira=Thopen&ano=2026&mes=9").get_json()
    assert r["fonte_disp"] == "bd"
    assert [d["disp"] for d in _usina(r, "Assis")["dias"]] == [0.99, 0.99, 0.99]


def test_operacao_carregando_diz_que_a_disponibilidade_ainda_nao_chegou(cli, monkeypatch):
    monkeypatch.setattr(fonte_operacao, "dados", lambda *a, **k: None)
    r = cli.get("/api/t/registros?carteira=Thopen&ano=2026&mes=10").get_json()
    assert r["disp_carregando"] is True and all(d["disp"] is None for d in _usina(r, "Altair")["dias"])


@pytest.mark.parametrize("q", ["carteira=Thopen&ano=2026&mes=13", "carteira=Thopen&ano=x&mes=10",
                               "carteira=Nada&ano=2026&mes=10"])
def test_parametro_invalido_e_400(cli, q):
    assert cli.get("/api/t/registros?" + q).status_code == 400
    assert cli.get("/api/t/registros.xlsx?" + q).status_code == 400


def test_exportar_so_as_cinco_colunas(cli):
    r = cli.get("/api/t/registros.xlsx?carteira=Thopen&ano=2026&mes=10")
    assert r.status_code == 200
    assert r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "Registros_Thopen_2026-10.xlsx" in r.headers["Content-Disposition"]
    ws = load_workbook(io.BytesIO(r.data)).active
    linhas = list(ws.iter_rows(values_only=True))
    assert linhas[0] == ("Usina", "Data", "Geração (kWh)", "Irradiação (kWh/m²)", "Disponibilidade (%)")
    assert linhas[1][0] == "Altair" and linhas[1][1].date() == dt.date(2026, 10, 1)
    assert linhas[1][2:] == (900.0, 6.0, 1.0)
    assert len(linhas) == 1 + 3 + 3                                              # Altair e Piancó, 3 dias cada
    assert ws["B2"].number_format == "dd/mm/yyyy" and ws["E2"].number_format == "0.00%"


def test_cliente_so_ve_e_exporta_as_carteiras_dele(cli, monkeypatch):
    monkeypatch.setenv("THOPEN_LOGIN", "1")
    monkeypatch.setitem(dth.app.config, "SESSION_COOKIE_SECURE", False)
    conta = {"conta_id": "c1", "nome": "Cliente Teste", "email": "cliente@exemplo.com", "carteiras": ["Thopen"],
             "papel": "cliente"}
    monkeypatch.setattr(acessos, "conta", lambda cid: conta if cid == "c1" else None)
    with cli.session_transaction() as s:
        s["conta_id"] = "c1"
    assert cli.get("/api/t/registros?carteira=Thopen&ano=2026&mes=10").status_code == 200
    assert cli.get("/api/t/registros?carteira=Polaris&ano=2026&mes=10").status_code == 403
    assert cli.get("/api/t/registros.xlsx?carteira=Polaris&ano=2026&mes=10").status_code == 403


# ── a tela ──────────────────────────────────────────────────────────────────────────────────────────────────────────
def _html():
    return open(os.path.join(AQUI, "..", "thopen", "templates", "dashboard_thopen.html"), encoding="utf-8").read()


def test_aba_registros_ao_lado_da_metodologia_e_fora_do_pdf():
    h = _html()
    assert h.index('data-page="metodologia"') < h.index('data-page="registros"') < h.index('id="fldCarteira"')
    assert 'id="pageRegistros"' in h and "#pageRegistros{ display:none !important; }" in h
    assert "/api/t/registros?" in h and "/api/t/registros.xlsx?" in h


def test_cantos_em_bevel_com_arredondado_de_reserva():
    h = _html()
    assert "@supports (corner-shape: bevel)" in h and "corner-shape:bevel" in h
