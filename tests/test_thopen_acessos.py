# -*- coding: utf-8 -*-
"""Login por pessoa no Dashboard Thopen (Levi, 14/10/2026; spec docs/superpowers/specs/2026-10-14-thopen-login-por-pessoa-
design.md). As contas vêm do Nexus (livro nexus_acessos_clientes, nome e e-mail selados para a chave PÚBLICA do 5080); a
senha, o 5080 grava (livro thopen_acessos, o hash cifrado numa chave derivada da PRIVADA dele). Nada em claro no banco
(leitura aberta), link de uso único, conta desativada não entra, limite de tentativas. API falsa, nada de rede."""
import hashlib
import io
import json
import os
import sys
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "thopen"))
import acessos as A  # noqa: E402

PRIV, PUB = A.gerar_par()
AGORA = datetime(2026, 10, 14, 12, 0, tzinfo=timezone.utc)


def conta(cid, email, carteiras="Thopen", papel="cliente", status="ativa", token=None, expira_h=72, nome="Ana Souza",
          usuario=None):
    """Uma linha do livro de contas como o Nexus grava (com `usuario`, como o Nexus grava desde 15/10/2026; sem ele, as
    contas de antes, só com o e-mail)."""
    d = {"nome": nome, "email": email}
    if usuario is not None:
        d["usuario"] = usuario
    dados = json.dumps(d, ensure_ascii=False)
    return {"conta_id": cid, "carteiras": carteiras, "papel": papel, "status": status,
            "convite_hash": hashlib.sha256(token.encode()).hexdigest() if token else "",
            "convite_expira": (AGORA + timedelta(hours=expira_h)).strftime("%Y-%m-%dT%H:%M:%SZ") if token else "",
            "criada_em": "2026-10-14T11:00:00Z", "alterada_em": "2026-10-14T11:00:00Z", "alterada_por_hmac": "x",
            "dados_5080": A.selar(dados, A.ctx_conta(cid), PUB), "dados_nexus": "cf1:qualquer"}


class ApiFalsa:
    """Os dois livros: o de contas (só leitura para o 5080) e o de senhas (o 5080 troca inteiro)."""

    def __init__(self, contas, senhas=()):
        self.livros = {A.LIVRO_CONTAS: {A.ABA_CONTAS: list(contas)}, A.LIVRO_SENHAS: {A.ABA_SENHAS: list(senhas)}}
        self.gravacoes = 0

    def _abas(self):
        out, i = [], 1
        for lv, abas in self.livros.items():
            for aba, linhas in abas.items():
                if lv == A.LIVRO_SENHAS and not linhas and not self.gravacoes:
                    continue                                    # o livro de senhas nasce na 1ª senha
                out.append({"id": i, "workbook_key": lv, "sheet_name": aba, "row_count": len(linhas)})
                i += 1
        return out

    def get(self, caminho, tok):
        if caminho == "/api/sheets":
            return self._abas()
        if caminho.startswith("/api/workbooks"):
            return [{"key": lv} for lv in self.livros if lv != A.LIVRO_SENHAS or self.gravacoes]
        sid = int(caminho.split("/")[3])
        s = next(x for x in self._abas() if x["id"] == sid)
        linhas = self.livros[s["workbook_key"]][s["sheet_name"]]
        off = int(caminho.split("offset=")[1].split("&")[0])
        cab = list(linhas[0]) if linhas else []
        return {"rows": [{"headers": cab, "values": [l.get(c) for c in cab]} for l in linhas[off:off + 1000]]}

    def criar(self, livro, nome, tok):
        pass

    def trocar(self, livro, xlsx, tok):
        assert livro == A.LIVRO_SENHAS                          # o 5080 nunca grava o livro de contas do Nexus
        ws = load_workbook(io.BytesIO(xlsx)).active
        cab = [c.value for c in ws[1]]
        self.livros[livro][ws.title] = [dict(zip(cab, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
        self.gravacoes += 1


@pytest.fixture
def api(monkeypatch):
    a = ApiFalsa([conta("c1", "ana@thopen.com.br", token="TOKEN-ANA"),
                  conta("c2", "beto@polaris.com", carteiras="Polaris;Copel", status="desativada"),
                  conta("c3", "levi@gridco.com.br", carteiras="", papel="grid", token="TOKEN-LEVI", nome="Levi")])
    monkeypatch.setattr(A.livro_banco, "_get", a.get)
    monkeypatch.setattr(A.livro_banco, "_criar_livro", a.criar)
    monkeypatch.setattr(A.livro_banco, "_trocar_livro", a.trocar)
    monkeypatch.setattr(A, "_privada", lambda: PRIV)
    monkeypatch.setattr(A, "_token", lambda: "tok")
    A.limpar()
    return a


def test_selo_so_abre_com_a_privada_e_preso_a_conta():
    s = A.selar("ana@thopen.com.br", A.ctx_conta("c1"), PUB)
    assert s.startswith("cx1:") and "ana@thopen" not in s      # o e-mail inteiro: "ana" sai por acaso no base64
    assert A.abrir(s, A.ctx_conta("c1"), PRIV) == "ana@thopen.com.br"
    with pytest.raises(A.CifraErro):
        A.abrir(s, A.ctx_conta("c9"), PRIV)                     # copiada para outra conta: não abre
    with pytest.raises(A.CifraErro):
        A.abrir(s, A.ctx_conta("c1"), A.gerar_par()[0])         # outra chave: não abre


def test_convite_define_a_senha_uma_vez_e_entra(api):
    c, motivo = A.convite("TOKEN-ANA", agora=AGORA)
    assert motivo == "ok" and c["email"] == "ana@thopen.com.br" and c["carteiras"] == ["Thopen"]
    assert A.verificar("ana@thopen.com.br", "senha-forte-123", agora=AGORA) is None          # sem senha ainda
    A.definir_senha("TOKEN-ANA", "senha-forte-123", agora=AGORA)
    entrou = A.verificar("ANA@thopen.com.br ", "senha-forte-123", agora=AGORA)              # e-mail sem caixa/espaço
    assert entrou and entrou["conta_id"] == "c1" and entrou["nome"] == "Ana Souza"
    assert A.verificar("ana@thopen.com.br", "errada-errada", agora=AGORA) is None
    assert A.convite("TOKEN-ANA", agora=AGORA)[1] == "usado"                                  # uso único
    # nada da senha nem do e-mail em claro no livro do 5080
    claro = json.dumps(api.livros[A.LIVRO_SENHAS], ensure_ascii=False, default=str)
    assert "senha-forte" not in claro and "ana@" not in claro and "scrypt" not in claro


def test_link_vencido_de_conta_desativada_ou_falso_nao_vale(api):
    assert A.convite("TOKEN-ANA", agora=AGORA + timedelta(hours=73))[1] == "vencido"
    assert A.convite("NAO-EXISTE", agora=AGORA)[1] == "invalido"
    api.livros[A.LIVRO_CONTAS][A.ABA_CONTAS][0]["status"] = "desativada"
    A.limpar()
    assert A.convite("TOKEN-ANA", agora=AGORA)[1] == "desativada"
    with pytest.raises(A.Recusada):
        A.definir_senha("TOKEN-ANA", "senha-forte-123", agora=AGORA)


def test_senha_curta_ou_diferente_e_recusada(api):
    with pytest.raises(A.Recusada):
        A.definir_senha("TOKEN-ANA", "curta", agora=AGORA)
    assert api.gravacoes == 0


def test_conta_desativada_depois_da_senha_nao_entra(api):
    A.definir_senha("TOKEN-ANA", "senha-forte-123", agora=AGORA)
    assert A.verificar("ana@thopen.com.br", "senha-forte-123", agora=AGORA)
    api.livros[A.LIVRO_CONTAS][A.ABA_CONTAS][0]["status"] = "desativada"
    A.limpar()
    assert A.verificar("ana@thopen.com.br", "senha-forte-123", agora=AGORA) is None


def test_link_novo_troca_a_senha_e_a_mais_nova_vale(api):
    A.definir_senha("TOKEN-ANA", "senha-antiga-123", agora=AGORA)
    novo = conta("c1", "ana@thopen.com.br", token="TOKEN-NOVO")
    api.livros[A.LIVRO_CONTAS][A.ABA_CONTAS][0] = novo
    A.limpar()
    A.definir_senha("TOKEN-NOVO", "senha-nova-4567", agora=AGORA + timedelta(minutes=5))
    assert A.verificar("ana@thopen.com.br", "senha-nova-4567", agora=AGORA + timedelta(minutes=6))
    assert A.verificar("ana@thopen.com.br", "senha-antiga-123", agora=AGORA + timedelta(minutes=6)) is None


def test_papel_grid_ve_todas_as_carteiras(api):
    c, _ = A.convite("TOKEN-LEVI", agora=AGORA)
    assert c["papel"] == "grid" and set(c["carteiras"]) == set(A.CARTEIRAS)


def test_limite_de_tentativas():
    lim = A.Limite(tentativas=5, janela=900)
    for _ in range(5):
        lim.erro("ip:1.2.3.4", agora=100)
    assert lim.bloqueado("ip:1.2.3.4", agora=101) and not lim.bloqueado("ip:5.6.7.8", agora=101)
    assert not lim.bloqueado("ip:1.2.3.4", agora=100 + 901)


def test_sem_a_chave_privada_ninguem_entra(api, monkeypatch):
    monkeypatch.setattr(A, "_privada", lambda: "")
    A.limpar()
    assert A.verificar("ana@thopen.com.br", "qualquer-coisa", agora=AGORA) is None
    assert A.estado()["erro"]


def test_conta_pelo_usuario_de_hoje_do_cliente(api):
    # Levi, 15/10/2026: "Mesmo usuário e senha de hoje" — a conta entra pelo usuário (sem e-mail), sem caixa nem espaço
    api.livros[A.LIVRO_CONTAS][A.ABA_CONTAS].append(
        conta("c4", "", nome="Thopen — acesso de hoje", usuario="thopen.operacao", token="TOKEN-THOPEN"))
    A.limpar()
    c, motivo = A.convite("TOKEN-THOPEN", agora=AGORA)
    assert motivo == "ok" and (c["usuario"], c["email"]) == ("thopen.operacao", "")
    A.definir_senha("TOKEN-THOPEN", "senha1", agora=AGORA)              # 6 caracteres: a senha de hoje serve
    entrou = A.verificar(" Thopen.Operacao ", "senha1", agora=AGORA)
    assert entrou and entrou["conta_id"] == "c4" and entrou["carteiras"] == ["Thopen"]
    with pytest.raises(A.Recusada, match="6 caracteres"):
        A.definir_senha("TOKEN-ANA", "12345", agora=AGORA)
    # a conta de antes (só e-mail) segue entrando pelo e-mail
    assert A.convite("TOKEN-ANA", agora=AGORA)[0]["usuario"] == "ana@thopen.com.br"
