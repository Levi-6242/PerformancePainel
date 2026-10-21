# -*- coding: utf-8 -*-
"""A trava contra apagar o histórico dos livros de decisão (cobertura.Decisoes.acrescentar, 20/10/2026, governança do
Nexus, ao abrir o livro nexus_falhas). Vale para os três livros que a plataforma grava: nexus_tempo_real, nexus_depara e
nexus_falhas.

O defeito: o `_ler_banco` devolve [] quando a listagem /api/sheets omite a aba por um instante; o `acrescentar` gravava
então só as linhas novas com replace=true — o livro inteiro trocado pela decisão do momento — e a conferência de depois
passava, porque compara com o que acabou de enviar. Agora: livro que já existe (está no /api/workbooks) e voltou vazio,
ou com menos linhas do que a última leitura boa (memória ou disco), não é gravado. Livro novo de verdade grava.
"""
import json
from io import BytesIO

import pytest

import cobertura as cb

Q1, Q2, Q3 = "2026-10-20T15:00:00Z", "2026-10-20T15:01:00Z", "2026-10-20T15:02:00Z"


class _R:
    def __init__(self, d, code=200):
        self._d, self.status_code = d, code

    def json(self):
        return self._d

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(self.status_code)


class _API:
    """O banco de mentira; `omite` = a listagem /api/sheets sem a aba (o soluço); `corta` = a leitura das linhas volta
    com menos do que o banco tem."""

    def __init__(self):
        self.livros, self.linhas, self.omite, self.corta, self.syncs = set(), None, False, 0, 0
        self.sem_contagem, self.omite_livros = False, False

    def get(self, url, headers=None, params=None, timeout=None):
        if url.endswith("/api/workbooks"):
            if self.omite_livros:                                  # o soluço nas duas listagens ao mesmo tempo
                return _R([])
            # o banco diz quantas abas o livro tem (sheet_count): 0 = criado e nunca gravado
            return _R([{"key": k} if self.sem_contagem else {"key": k, "sheet_count": 0 if self.linhas is None else 1}
                       for k in self.livros])
        if url.endswith("/api/sheets"):
            ok = self.linhas is not None and not self.omite
            return _R([{"id": 5, "workbook_key": cb.LIVRO, "sheet_name": cb.ABA}] if ok else [])
        if url.endswith("/api/sheets/5/rows"):
            off, lim = params["offset"], params["limit"]
            vis = (self.linhas or [])[:len(self.linhas or []) - self.corta]
            return _R({"rows": [{"headers": cb.CAB, "values": lin} for lin in vis[off:off + lim]]})
        raise AssertionError(url)

    def post(self, url, headers=None, json=None, params=None, files=None, timeout=None):
        if url.endswith("/api/workbooks"):
            self.livros.add(json["key"])
            return _R({})
        if url.endswith(f"/api/workbooks/{cb.LIVRO}/sync-xlsx"):
            from openpyxl import load_workbook
            self.syncs += 1
            ws = load_workbook(BytesIO(files["file"][1]))[cb.ABA]
            self.linhas = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
            return _R({})
        raise AssertionError(url)


def _livro_com_duas(tmp_path, api, nome="dec.json"):
    d = cb.Decisoes("https://x/db", "tok", str(tmp_path / nome), ttl=0, sessao=api)
    assert d.acrescentar(cb.semente("pv", [1, 2], [1], Q1)) == 2
    return d


def test_livro_existente_com_a_aba_sumida_da_listagem_nao_e_regravado(tmp_path):
    api = _API()
    d = _livro_com_duas(tmp_path, api)
    api.omite = True                                              # o soluço: o livro está lá, a aba não veio na lista
    with pytest.raises(cb.GravacaoErro, match=r"veio com 0 linhas e eu conhecia 2: não gravei para não apagar"):
        d.acrescentar([cb.linha("pv", 1, "excluir", "tela", Q2)])
    assert api.syncs == 1 and len(api.linhas) == 2                 # o histórico ficou
    api.omite = False
    assert d.acrescentar([cb.linha("pv", 1, "excluir", "tela", Q2)]) == 3       # passado o soluço, grava


def test_maquina_que_nunca_leu_o_livro_tambem_nao_grava_por_cima(tmp_path):
    # o servidor recém-ligado: nem memória nem cópia em disco — mas o livro existe e voltou vazio
    api = _API()
    _livro_com_duas(tmp_path, api)
    api.omite = True
    novo = cb.Decisoes("https://x/db", "tok", str(tmp_path / "outra_maquina.json"), ttl=0, sessao=api)
    with pytest.raises(cb.GravacaoErro, match="veio com 0 linhas e o banco diz que ele tem abas"):
        novo.acrescentar([cb.linha("pv", 2, "excluir", "tela", Q2)])
    assert api.syncs == 1 and len(api.linhas) == 2


def test_releitura_com_menos_linhas_do_que_a_ultima_leitura_boa_nao_grava(tmp_path):
    api = _API()
    d = _livro_com_duas(tmp_path, api)
    assert d.acrescentar([cb.linha("pv", 2, "excluir", "tela", Q2)]) == 3
    api.corta = 1                                                 # a leitura volta com 2 das 3
    with pytest.raises(cb.GravacaoErro, match="veio com 2 linhas e eu conhecia 3"):
        d.acrescentar([cb.linha("pv", 1, "excluir", "tela", Q3)])
    # o processo novo (memória vazia) conhece as 3 pela cópia em disco que a última leitura boa deixou
    (tmp_path / "copia.json").write_text(json.dumps({"lido_em": Q2, "linhas": api.linhas}), encoding="utf-8")
    disco = cb.Decisoes("https://x/db", "tok", str(tmp_path / "copia.json"), ttl=0, sessao=api)
    with pytest.raises(cb.GravacaoErro, match="veio com 2 linhas e eu conhecia 3"):
        disco.acrescentar([cb.linha("pv", 1, "excluir", "tela", Q3)])
    assert api.syncs == 2 and len(api.linhas) == 3


def test_livro_novo_de_verdade_grava(tmp_path):
    api = _API()                                                   # fora do /api/workbooks e sem aba: nasce agora
    d = cb.Decisoes("https://x/db", "tok", str(tmp_path / "dec.json"), ttl=0, sessao=api)
    assert d.acrescentar([cb.linha("pv", 1, "incluir", "tela", Q1)]) == 1
    assert cb.LIVRO in api.livros and api.syncs == 1


def test_livro_criado_cuja_1a_gravacao_caiu_no_sync_volta_a_gravar(tmp_path):
    # o De-para em 11/10/2026: o POST /api/workbooks criou o livro e o sync caiu; sem aba nenhuma (sheet_count 0) não
    # há histórico a perder, e a trava não pode prender o livro para sempre
    api = _API()
    api.livros.add(cb.LIVRO)                                      # criado, nunca gravado
    d = cb.Decisoes("https://x/db", "tok", str(tmp_path / "dec.json"), ttl=0, sessao=api)
    assert d.acrescentar([cb.linha("pv", 1, "incluir", "tela", Q1)]) == 1
    assert api.syncs == 1


def test_livro_existente_sem_a_contagem_de_abas_e_vazio_nao_grava(tmp_path):
    # sem o sheet_count o banco não prova que o livro está vazio: soluço, não grava
    api = _API()
    _livro_com_duas(tmp_path, api)
    api.omite, api.sem_contagem = True, True
    novo = cb.Decisoes("https://x/db", "tok", str(tmp_path / "outra.json"), ttl=0, sessao=api)
    with pytest.raises(cb.GravacaoErro, match="não diz quantas abas"):
        novo.acrescentar([cb.linha("pv", 2, "excluir", "tela", Q2)])
    assert api.syncs == 1 and len(api.linhas) == 2


def test_as_duas_listagens_somem_juntas_e_a_maquina_que_conhece_o_livro_nao_grava(tmp_path):
    # /api/workbooks e /api/sheets vazios no mesmo instante: o livro pareceria novo e o replace apagaria as 2 linhas
    api = _API()
    d = _livro_com_duas(tmp_path, api)
    api.omite, api.omite_livros = True, True
    with pytest.raises(cb.GravacaoErro, match="veio com 0 linhas e eu conhecia 2"):
        d.acrescentar([cb.linha("pv", 1, "excluir", "tela", Q2)])
    assert api.syncs == 1 and len(api.linhas) == 2
