# -*- coding: utf-8 -*-
"""A meta do mês por usina no banco (12/10/2026): o Gerencial do worker publicado no livro `nexus_performance ·
fato_meta_usina_mes`, para a Visão do responsável do Nexus ler os MESMOS números da Visão gerencial."""
from datetime import datetime

import gerencial_publicar as gp

U1 = {"usina": "Altair 1", "cliente": "Thopen", "carteira": "Thopen", "prod": 812.4, "p50": 900.0, "atingimento": 90.3,
      "pr": 78.1, "pr_meta": 80.0, "ipoa": 150.2, "ipoa_prev": 160.0, "dias_cob": 11, "dias_mes": 31, "meta_dias": 11,
      "p50_src": "BD_Thopen", "src": "pg", "full_oem": True, "pr_motivo": "texto livre que não vai ao banco"}
U2 = {"usina": "Barretos 1", "cliente": "—", "carteira": "—", "prod": 300.0, "p50": None, "atingimento": None,
      "pr": None, "pr_meta": None, "ipoa": 0, "ipoa_prev": None, "dias_cob": None, "dias_mes": 31, "meta_dias": 0,
      "p50_src": None, "src": "bd", "full_oem": False}


def test_as_linhas_do_mes_sem_texto_livre_e_sem_o_traco():
    out = gp.linhas({"usinas": [U2, U1, {"usina": " "}]}, 202610)
    assert [r[1] for r in out] == ["Altair 1", "Barretos 1"]
    assert dict(zip(gp.CAB, out[0])) == {
        "mes_id": 202610, "usina": "Altair 1", "cliente": "Thopen", "carteira": "Thopen", "prod_kwh": 812400,
        "meta_kwh": 900000, "atingimento_pct": 90.3, "pr_pct": 78.1, "pr_meta_pct": 80.0, "ipoa_kwh_m2": 150.2,
        "ipoa_meta_kwh_m2": 160.0, "dias_cobertos_qtd": 11, "dias_mes_qtd": 31, "meta_dias_qtd": 11, "full_om": 1,
        "fonte": "pg", "meta_fonte": "BD_Thopen", "anomalia": None}
    assert out[1][2] is None and out[1][14] == 0                       # "—" vira vazio; sem Full O&M = 0
    assert "texto livre" not in str(out)                               # o pr_motivo não vai (leitura aberta)


def test_a_anomalia_vai_no_fim_da_linha():
    # 18/10/2026 (Levi: "apenas sinalize que o número está anômalo"): a frase da régua do app, na ÚLTIMA coluna
    assert gp.CAB[-1] == "anomalia" and gp.CAB[:-1][-1] == "meta_fonte"
    (r,) = gp.linhas({"usinas": [{**U1, "anomalia": "meta_baixa,geracao_impossivel", "anomalia_txt": "frase"}]}, 202610)
    assert r[-1] == "meta_baixa,geracao_impossivel" and "frase" not in r    # código, não frase (regra 8 do Nexus)


def test_mesclar_troca_so_os_meses_refeitos():
    antigas = [[202609, "Altair 1"] + [1] * 15, [202610, "Altair 1"] + [2] * 15, [202608, "Zeta"] + [3] * 15]
    novas = {202610: [[202610, "Altair 1"] + [9] * 15, [202610, "Barretos 1"] + [9] * 15]}
    out = gp.mesclar(antigas, novas)
    assert [(r[0], r[1], r[2]) for r in out] == [(202608, "Zeta", 3), (202609, "Altair 1", 1), (202610, "Altair 1", 9),
                                                  (202610, "Barretos 1", 9)]


SERIE = {"Altair 1": {"meta": 0.827, "pot": 3.5, "meta_ger": 18.57, "meta_poa": 7.4,
                      "pontos": [{"dia": 2, "realizado": 0.8491965, "geracao": 6.27, "poa": 2.09},
                                 {"dia": 1, "realizado": None, "geracao": 0.0, "poa": 0.1},
                                 {"dia": "x", "realizado": 1, "geracao": 1, "poa": 1}]}}


def test_a_serie_diaria_do_painel_noc_em_linhas():
    """A série do PR diário do Painel NOC (`/api/g/diario`): a meta do dia é a mesma todo dia; o PR em %; o dia vazio de
    PR fica vazio; ponto sem dia válido sai."""
    out = gp.linhas_dia(SERIE, 202610)
    assert [dict(zip(gp.CAB_DIA, r)) for r in out] == [
        {"mes_id": 202610, "usina": "Altair 1", "data_id": 20261001, "geracao_kwh": 0, "meta_kwh": 18570,
         "ipoa_kwh_m2": 0.1, "ipoa_meta_kwh_m2": 7.4, "pr_pct": None, "pr_meta_pct": 82.7},
        {"mes_id": 202610, "usina": "Altair 1", "data_id": 20261002, "geracao_kwh": 6270, "meta_kwh": 18570,
         "ipoa_kwh_m2": 2.09, "ipoa_meta_kwh_m2": 7.4, "pr_pct": 84.9, "pr_meta_pct": 82.7}]
    assert gp.linhas_dia({"X": None, " ": SERIE["Altair 1"]}, 202610) == []


def test_a_geracao_do_dia_e_a_do_gerencial_quando_vem():
    """13/10/2026: o ponto traz a geração do Gerencial em kWh (`geracao_kwh`); None = dia fora da conta dele, mesmo com
    o `geracao` do PR diário."""
    serie = {"Altair 1": {**SERIE["Altair 1"], "pontos": [
        {"dia": 1, "realizado": None, "poa": None, "geracao": 9.99, "geracao_kwh": 4321.4},
        {"dia": 2, "realizado": 0.8491965, "geracao": 6.27, "poa": 2.09, "geracao_kwh": None}]}}
    assert [(r[2], r[3], r[7]) for r in gp.linhas_dia(serie, 202610)] == [(20261001, 4321, None), (20261002, None, 84.9)]


def test_o_xlsx_leva_as_duas_abas_e_vazio_de_verdade():
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(gp.xlsx_bytes({gp.ABA: (gp.CAB, [[202610, "A"] + [""] * 15]),
                                              gp.ABA_DIA: (gp.CAB_DIA, gp.linhas_dia(SERIE, 202610))})))
    assert wb.sheetnames == [gp.ABA, gp.ABA_DIA]
    assert [c.value for c in wb[gp.ABA][1]] == gp.CAB and wb[gp.ABA]["C2"].value is None
    assert wb[gp.ABA_DIA].max_row == 3


def test_o_mes_anterior_se_refaz_ate_o_dia_5():
    assert gp.meses_a_refazer(2026, 10, 3) == [(2026, 10), (2026, 9)]
    assert gp.meses_a_refazer(2026, 10, 6) == [(2026, 10)]
    assert gp.meses_a_refazer(2027, 1, 2) == [(2027, 1), (2026, 12)]


# ── o laço no app ─────────────────────────────────────────────────────────────────────────────────────────────────────
import app  # noqa: E402


class _R:
    def __init__(self, d):
        self._d, self.status_code = d, 200

    def json(self):
        return self._d

    def raise_for_status(self):
        pass


class _Banco:
    """O livro `nexus_performance` com as duas abas (a do mês e a diária), lidas pelo `cobertura.Decisoes`."""

    def __init__(self):
        self.linhas, self.dia, self.fora, self.envios = None, None, False, 0

    def get(self, url, headers=None, params=None, timeout=None):
        if self.fora:
            raise ConnectionError("banco fora")
        if url.endswith("/api/sheets"):
            abas = [(7, gp.ABA, self.linhas), (8, gp.ABA_DIA, self.dia)]
            return _R([{"id": i, "workbook_key": gp.LIVRO, "sheet_name": a} for i, a, lin in abas if lin is not None])
        off, lim = params["offset"], params["limit"]
        cab, lin = (gp.CAB, self.linhas) if "/7/" in url else (gp.CAB_DIA, self.dia)
        return _R({"rows": [{"headers": cab, "values": x} for x in (lin or [])[off:off + lim]]})

    def sincronizar(self, conteudo, **k):
        from io import BytesIO
        from openpyxl import load_workbook
        assert k["chave"] == gp.LIVRO
        wb = load_workbook(BytesIO(conteudo))
        self.linhas = [list(r) for r in wb[gp.ABA].iter_rows(min_row=2, values_only=True)]
        self.dia = [list(r) for r in wb[gp.ABA_DIA].iter_rows(min_row=2, values_only=True)]
        self.envios += 1
        return {"inserted": len(self.linhas) + len(self.dia)}


def _ligar(monkeypatch, banco, gerencial, series=None):
    import falhas_publicar
    monkeypatch.setattr(app, "_ger_publicar_ligado", lambda: True)
    monkeypatch.setattr(app, "_ger_series_diarias", lambda g, a, m: series if series is not None else SERIE)
    monkeypatch.setenv("GRIDCO_SQL_TOKEN", "tok")
    monkeypatch.setattr(app, "_ger_pub", {"ts": 0.0, "marca": None})
    monkeypatch.setattr(app, "_ger_cache", {"ts": 1.0, "data": gerencial})
    original = app._cob.Decisoes
    monkeypatch.setattr(app._cob, "Decisoes", lambda *a, **k: original(*a, **dict(k, sessao=banco)))
    monkeypatch.setattr(falhas_publicar, "sincronizar", banco.sincronizar)


def test_publica_o_mes_do_cache_e_so_quando_muda(monkeypatch):
    banco = _Banco()
    banco.linhas = [[202609, "Altair 1"] + [None] * 15]                 # setembro já publicado fica
    _ligar(monkeypatch, banco, {"usinas": [U1, U2]})
    assert app._ger_publicar(agora=datetime(2026, 10, 12, 10)) == 3
    assert sorted((r[0], r[1]) for r in banco.linhas) == [(202609, "Altair 1"), (202610, "Altair 1"), (202610, "Barretos 1")]
    # a série diária do Painel NOC vai junto, na aba própria
    assert [(r[1], r[2], r[3]) for r in banco.dia] == [("Altair 1", 20261001, 0), ("Altair 1", 20261002, 6270)]
    assert app._ger_publicar(agora=datetime(2026, 10, 12, 10)) is None  # a hora não passou
    app._ger_pub["ts"] = 0.0
    assert app._ger_publicar(agora=datetime(2026, 10, 12, 11)) is None and banco.envios == 1   # nada mudou: não regrava


def test_gerencial_que_falhou_nao_apaga_o_mes_e_banco_fora_nao_grava(monkeypatch):
    banco = _Banco()
    banco.linhas = [[202610, "Altair 1"] + [5] * 15]
    _ligar(monkeypatch, banco, {"erro": "PG fora", "usinas": []})
    assert app._ger_publicar(agora=datetime(2026, 10, 12, 10)) is None and banco.envios == 0
    _ligar(monkeypatch, banco, {"usinas": [U1]})
    banco.fora = True
    assert app._ger_publicar(agora=datetime(2026, 10, 12, 10)) is None and banco.envios == 0


def test_mes_sem_serie_diaria_nao_apaga_a_do_banco(monkeypatch):
    banco = _Banco()
    banco.linhas = []
    banco.dia = [[202610, "Altair 1", 20261001] + [1] * 6]
    _ligar(monkeypatch, banco, {"usinas": [U1]}, series={})
    assert app._ger_publicar(agora=datetime(2026, 10, 12, 10)) == 1 and banco.dia == [[202610, "Altair 1", 20261001] + [1] * 6]


def test_desligado_nao_publica(monkeypatch):
    banco = _Banco()
    _ligar(monkeypatch, banco, {"usinas": [U1]})
    monkeypatch.setattr(app, "_ger_publicar_ligado", lambda: False)
    assert app._ger_publicar(agora=datetime(2026, 10, 12, 10)) is None and banco.linhas is None


# ── a série diária pela geração do Gerencial (13/10/2026) ────────────────────────────────────────────────────────────
# A tabela da janela da Meta do Nexus não fechava no % do Gerencial em 78 das 131 usinas: na Thopen o PR diário pula o
# dia com POA até 0,3 (Indaiatuba: 5 dias e 87% contra 12 dias e 66%), e no BD_Performance escreve 0 no dia em branco
# (Castelo do Piauí: 31 dias e 54,9% contra 9 dias e 189,1%).
def _diario(pontos, meta_ger=10.0):
    return {"usina": "X", "meta": 0.8, "pot": 1.0, "meta_ger": meta_ger, "meta_poa": 5.0, "pontos": pontos}


def test_a_serie_diaria_tem_os_dias_e_a_geracao_do_gerencial(monkeypatch):
    pr_th = [{"dia": 2, "realizado": 0.7, "geracao": 9.0, "poa": 5.0}, {"dia": 3, "realizado": 0.8, "geracao": 11.0, "poa": 6.0},
             {"dia": 4, "realizado": 0.9, "geracao": 3.0, "poa": 1.0}]
    pr_bp = [{"dia": d, "realizado": 0.8 if d <= 2 else None, "geracao": 12.0 if d <= 2 else 0.0,
              "poa": 6.0 if d <= 2 else 0.0} for d in range(1, 32)]
    chamadas = []

    def falso(cli, usina, ano, mes):
        chamadas.append(usina)
        return _diario(pr_th) if usina == "Indaiatuba" else (None if usina == "Só PG" else _diario(pr_bp))
    monkeypatch.setattr(app, "_ger_diario", falso)
    g = {"usinas": [
        # Thopen: o dia 1 (POA até 0,3) não está no PR diário e entra com a geração e sem PR; o dia 4 só do PR diário
        # (com PR) entra sem geração
        {"usina": "Indaiatuba", "cliente": "Thopen", "src": "bdthopen", "ger_dia": {1: 7000.0, 2: 9000.0, 3: 11000.0},
         "p50": 30.0, "meta_dias": 3},
        # BD_Performance: o dia em branco (0 no PR diário, sem PR) não vira linha; a chave em texto (JSON do snapshot)
        {"usina": "Castelo", "cliente": "Greenyellow", "src": "bdperf", "ger_dia": {"1": 12000.0, "2": 12000.0},
         "p50": 20.0, "meta_dias": 2},
        # sem PR diário: a meta do dia vem do Gerencial (p50 ÷ dias da meta)
        {"usina": "Só PG", "cliente": "Thopen", "src": "pg", "ger_dia": {5: 1500.0}, "p50": 9.0, "meta_dias": 3},
        # a meta do dia exata do Gerencial vence a do PR diário (arredondada em 0,01 MWh)
        {"usina": "Exata", "cliente": "Thopen", "src": "bdthopen", "ger_dia": {2: 9000.0}, "meta_dia_kwh": 10004.6},
        {"usina": "Sem geração", "cliente": "Thopen", "src": "cadastro", "ger_dia": None},
    ]}
    s = app._ger_series_diarias(g, 2026, 10)
    assert set(s) == {"Indaiatuba", "Castelo", "Só PG", "Exata"} and "Sem geração" not in chamadas
    assert [r[4] for r in gp.linhas_dia({"Exata": s["Exata"]}, 202610) if r[3] is not None] == [10005]
    assert [(p["dia"], p["geracao_kwh"], p["realizado"]) for p in s["Indaiatuba"]["pontos"]] == [
        (1, 7000.0, None), (2, 9000.0, 0.7), (3, 11000.0, 0.8), (4, None, 0.9)]
    assert [(p["dia"], p["geracao_kwh"]) for p in s["Castelo"]["pontos"]] == [(1, 12000.0), (2, 12000.0)]
    assert s["Só PG"]["meta_ger"] == 3.0 and [(p["dia"], p["geracao_kwh"]) for p in s["Só PG"]["pontos"]] == [(5, 1500.0)]
    # a soma dos dias com geração e a meta dos mesmos dias = o Gerencial (27 MWh de 30 nos 3 dias da Indaiatuba)
    lin = [r for r in gp.linhas_dia(s, 202610) if r[1] == "Indaiatuba" and r[3] is not None]
    assert sum(r[3] for r in lin) == 27000 and sum(r[4] for r in lin) == 30000


def test_gerencial_de_antes_do_ger_dia_nao_publica_serie(monkeypatch):
    """O snapshot do disco ao subir (de antes do `ger_dia`): nenhuma série, e o mês do banco fica como está."""
    monkeypatch.setattr(app, "_ger_diario", lambda *a: _diario([{"dia": 1, "realizado": 0.8, "geracao": 1.0, "poa": 5.0}]))
    assert app._ger_series_diarias({"usinas": [{"usina": "A", "cliente": "Thopen", "src": "bdthopen"}]}, 2026, 10) == {}


def test_o_thopen_guarda_o_dia_a_dia_que_conta(monkeypatch):
    """`ger_dia` = os dias do `dias_cob` (geração > 0, ou zero confirmado pelo sol medido), com a geração de cada um: a
    soma é o `prod_mwh`."""
    import dashboard_thopen as dth
    from datetime import datetime as _dt
    hoje = _dt.now()
    d1 = hoje.replace(day=1)
    monkeypatch.setattr(dth, "_registro", lambda: {"Usina W": {"cliente": "Thopen", "pot_mwp": 1.0}})
    monkeypatch.setattr(dth, "_CARTEIRA_DE", {"Usina W": "Thopen"})
    monkeypatch.setattr(dth, "_daily_records", lambda u: [
        {"data": d1, "ger": 1000.0, "ipoa": 5.0},       # sol medido: conta
        {"data": hoje, "ger": 900.0, "ipoa": None},     # sensor sem leitura, mas gerou: conta
        {"data": hoje, "ger": 0.0, "ipoa": None},       # zero sem sol medido: lacuna, não conta
    ])
    monkeypatch.setattr(app, "_thopen_prod_cache", {"ts": 0.0, "ym": None, "data": {}, "warming": True})
    r = app._thopen_prod_build()[app._nrm("Usina W")]
    assert len(r["ger_dia"]) == r["dias_cob"] and round(sum(r["ger_dia"].values()) / 1000, 3) == round(r["prod_mwh"], 3)


def test_o_bd_performance_guarda_o_dia_a_dia_que_conta(monkeypatch, tmp_path):
    """Dia com número de inversor conta (zero = usina parada); dia em branco e linha lixo não contam."""
    from openpyxl import Workbook
    from datetime import datetime as _dt
    wb = Workbook()
    ws = wb.active
    ws.title = "Usina BP"
    ws.append(["Usina", "Data", "Inversor 1", "Inversor 2", "Validação", "IPOA DEF"])
    ws.append(["Usina BP", _dt(2026, 9, 1), 100.0, 200.0, 1, 5.0])          # conta: 300
    ws.append(["Usina BP", _dt(2026, 9, 2), None, None, None, None])        # em branco: não conta
    ws.append(["Usina BP", _dt(2026, 9, 3), 0.0, 0.0, 1, 5.0])              # usina parada: conta com 0
    ws.append(["Usina BP", _dt(2026, 9, 4), 1.3e6, 10.0, 0, 5.0])           # lixo (não validado, impossível): não conta
    arq = tmp_path / "bd.xlsx"
    wb.save(arq)
    monkeypatch.setattr(app, "_bd_readable", lambda: str(arq))
    monkeypatch.setattr(app, "INFO_GERAL", {app._nrm("Usina BP"): {"usina": "Usina BP", "cliente": "Greenyellow",
                                                                   "potencia_mwp": 1.0}})
    r = app._bdperf_prod_build(2026, 9)[app._nrm("Usina BP")]
    assert r["ger_dia"] == {1: 300.0, 3: 0.0} and r["dias_cob"] == 2 and round(r["prod_mwh"], 3) == 0.3


def test_a_resposta_do_gerencial_nao_leva_o_dia_a_dia(monkeypatch):
    monkeypatch.setattr(app, "_gerencial_payload", lambda **k: {
        "usinas": [{"usina": "A", "prod": 1.0, "ger_dia": {1: 1000.0}, "meta_dia_kwh": 9.5}], "portfolio": {}})
    with app.app.test_request_context("/api/gerencial"):
        d = app.api_gerencial().get_json()
    assert d["usinas"] == [{"usina": "A", "prod": 1.0}]
