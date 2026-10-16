# -*- coding: utf-8 -*-
"""Contestações do cliente no banco (rodada 5, Levi 13/10/2026: "contestação ligada ao Nexus"; "ficaria na API do PG,
pode aplicar e construir!"). O 5080 grava o livro thopen_contestacoes; o Nexus lê na torre Contratos e clientes.

Trava o que não pode escapar: nada pessoal em claro no banco (leitura aberta), o 5080 só tem a chave PÚBLICA (não abre
nada), a cifra presa à linha, uma linha por setor, e a gravação que nunca regrava o livro a partir de uma leitura
parcial (o livro é trocado inteiro). API falsa, nada de rede."""
import io
import json
import os
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "thopen"))
import contestacao_banco as cb  # noqa: E402

PRIV, PUB = cb.gerar_par()
D = {"itens": ["geracao", "disponibilidade", "ronda"], "nome": "Ana Souza", "contato": "ana@cliente.com.br",
     "mensagem": "A geração do dia 06 parece baixa.", "tela": {"geracao": "4.429 kWh", "disp": "100,0%"}}


def test_cifra_so_abre_com_a_privada_e_presa_a_linha():
    c = cb.cifrar("Ana Souza", "thopen_contestacoes/abc/dados", PUB)
    assert c.startswith("cx1:") and "Ana" not in c
    assert cb.decifrar(c, "thopen_contestacoes/abc/dados", PRIV) == "Ana Souza"
    with pytest.raises(cb.CifraErro):
        cb.decifrar(c, "thopen_contestacoes/OUTRA/dados", PRIV)       # copiada para outra linha: não abre
    with pytest.raises(cb.CifraErro):
        cb.decifrar(c, "thopen_contestacoes/abc/dados", cb.gerar_par()[0])   # outra chave: não abre
    assert cb.cifrar("x", "c", PUB) != cb.cifrar("x", "c", PUB)       # cada cifra tem a sua chave efêmera


def test_uma_linha_por_setor_sem_nada_pessoal_em_claro():
    ls = cb.linhas(D, "Guaratinguetá V", 86, 2026, 10, 6, 6, {"performance": True, "cos": True, "zeladoria": False},
                   PUB)
    perf, cos, zel = ls
    assert (perf["setor"], perf["itens"], perf["avisado_teams"]) == ("performance", "geracao", "sim")
    # 15/10/2026: a disponibilidade é do COS
    assert (cos["setor"], cos["itens"], cos["avisado_teams"]) == ("cos", "disponibilidade", "sim")
    assert (zel["setor"], zel["itens"], zel["avisado_teams"]) == ("zeladoria", "ronda", "não")
    assert perf["contestacao_id"] == zel["contestacao_id"] and perf["id"] != zel["id"]
    assert (perf["usina"], perf["usina_id"], perf["data_id_contestada_ini"], perf["data_id_contestada_fim"]) == (
        "Guaratinguetá V", 86, 20261006, 20261006)
    assert set(perf) == set(cb.CAB)
    claro = json.dumps(ls, ensure_ascii=False)
    for s in ("Ana", "cliente.com.br", "baixa", "4.429"):
        assert s not in claro, s                                      # a API do banco tem leitura aberta
    dados = json.loads(cb.decifrar(perf["dados_cifrados"], cb.contexto(perf["id"]), PRIV))
    assert dados == {"nome": "Ana Souza", "contato": "ana@cliente.com.br", "mensagem": D["mensagem"],
                     "tela": D["tela"], "itens": ["geracao"]}


class ApiFalsa:
    """O pedaço da API do banco que o 5080 usa: listar abas, ler linhas e trocar o livro inteiro (sync-xlsx)."""

    def __init__(self, linhas=None, mente_contagem=False):
        self.linhas = list(linhas or [])
        self.mente = mente_contagem
        self.livro_existe = bool(linhas)
        self.gravacoes = []

    def get(self, caminho, tok):
        if caminho == "/api/sheets":
            if not self.livro_existe:
                return []
            n = len(self.linhas) + (1 if self.mente else 0)
            return [{"id": 7, "workbook_key": cb.LIVRO, "sheet_name": cb.ABA, "row_count": n}]
        if caminho.startswith("/api/workbooks"):
            return [{"key": cb.LIVRO}] if self.livro_existe else []
        off = int(caminho.split("offset=")[1].split("&")[0])
        return {"rows": [{"headers": list(cb.CAB), "values": [l.get(c) for c in cb.CAB]}
                         for l in self.linhas[off:off + 1000]]}

    def criar(self, chave, nome, tok):
        self.livro_existe = True

    def trocar(self, xlsx, tok):
        ws = load_workbook(io.BytesIO(xlsx)).active
        assert ws.title == cb.ABA
        cab = [c.value for c in ws[1]]
        self.linhas = [dict(zip(cab, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
        self.gravacoes.append(len(self.linhas))


def _usa(monkeypatch, api):
    monkeypatch.setattr(cb, "_get", api.get)
    monkeypatch.setattr(cb, "_criar_livro", api.criar)
    monkeypatch.setattr(cb, "_trocar_livro", api.trocar)


def test_grava_somando_ao_que_ja_esta_no_banco_e_confere(monkeypatch):
    antigas = cb.linhas(D, "Altair", 1, 2026, 10, 2, 2, {"performance": True}, PUB)
    api = ApiFalsa(antigas)
    _usa(monkeypatch, api)
    novas = cb.linhas(D, "Guaratinguetá V", 86, 2026, 10, 6, 6, {"performance": True, "zeladoria": True}, PUB)
    assert cb.registrar(novas, token="t") == 3
    assert api.gravacoes == [3] and [l["id"] for l in api.linhas] == [l["id"] for l in antigas + novas]


def test_livro_novo_nasce_na_primeira_contestacao(monkeypatch):
    api = ApiFalsa()
    _usa(monkeypatch, api)
    assert cb.registrar(cb.linhas(D, "Altair", 1, 2026, 10, 2, 2, {"performance": True}, PUB), token="t") == 1
    assert api.livro_existe and api.gravacoes == [1]


def test_leitura_parcial_nunca_regrava_o_livro(monkeypatch):
    # o livro é trocado inteiro: regravar a partir de uma leitura que perdeu linhas apagaria contestações
    api = ApiFalsa(cb.linhas(D, "Altair", 1, 2026, 10, 2, 2, {"performance": True}, PUB), mente_contagem=True)
    _usa(monkeypatch, api)
    with pytest.raises(cb.GravacaoErro):
        cb.registrar(cb.linhas(D, "Altair", 1, 2026, 10, 3, 3, {"performance": True}, PUB), token="t")
    assert api.gravacoes == []


def test_sem_token_ou_sem_chave_nao_grava(monkeypatch):
    api = ApiFalsa()
    _usa(monkeypatch, api)
    with pytest.raises(cb.GravacaoErro):
        cb.registrar(cb.linhas(D, "Altair", 1, 2026, 10, 2, 2, {"performance": True}, PUB), token="")
    with pytest.raises(cb.CifraErro):
        cb.linhas(D, "Altair", 1, 2026, 10, 2, 2, {"performance": True}, "")
    assert api.gravacoes == []
