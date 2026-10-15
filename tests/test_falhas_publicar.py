# -*- coding: utf-8 -*-
"""Falhas de strings e trackers no workbook `falhas_performance` da Gridco API (Levi, 25/09/2026: "pode criar o
workbook também").

São as tabelas que ele pediu no começo — STRINGS: cliente, usina, inversor, quantidade, strings em texto e perda;
TRACKERS: cliente, usina, tracker, inversor, parou, voltou e perda — mais o saiu/voltou por string. O caminho é o que
o gêmeo já usa em produção: xlsx com o cabeçalho na linha 1, vazio de verdade (o sync rejeita texto vazio) e
sync-xlsx?replace=true, criando o workbook só se faltar (a API não apaga workbook).

Duas regras que só existem por causa do sync de hora em hora:
- as linhas vão pela data de início: o que é novo entra no FIM e o replace não reescreve o mês inteiro;
- nenhuma coluna muda sozinha a cada volta (o "atualizado em" mora na aba `atualizacao`, uma linha por mês) —
  senão cada sync regravaria todas as linhas e encheria o histórico por linha da API.
"""
import io

import pytest
from openpyxl import load_workbook

import falhas_publicar as pub


def pacote(mes="2026-09", gerado="2026-09-24 22:48"):
    return {"mes": mes, "gerado_em": gerado, "periodo": [f"{mes}-01", f"{mes}-24"],
            "regua": {"janela_strings": "06–18h, 2 h seguidas", "fator_tracker": "1 − cos(desvio)"},
            "strings": {
                "rows": [
                    {"dia": f"{mes}-05", "fonte": "pv", "cliente": "Thopen", "usina": "Inhapi", "inversor": "Inversor 3.3",
                     "qtd": 4, "strings": "Ipv1, Ipv2, Ipv3, Ipv4", "h_sol": 42.5, "perda_kwh": 410.2,
                     "perda_nominal_kwh": 800.0, "kwp_string": 11.2, "yield": 5.1, "flags": []},
                    {"dia": f"{mes}-02", "fonte": "sunop", "cliente": "Athon", "usina": "SMP100", "inversor": "Inversor 5.2",
                     "qtd": 1, "strings": "ST 19", "h_sol": 9.0, "perda_kwh": 80.0, "perda_nominal_kwh": 100.0,
                     "kwp_string": 10.0, "yield": 4.5, "flags": ["sem geração no dia (4,5 kWh/kWp)", "captura parcial"]},
                ],
                "episodios": [
                    {"fonte": "pv", "cliente": "Thopen", "usina": "Inhapi", "inversor": "Inversor 3.3", "string": "Ipv1",
                     "inicio": f"{mes}-05 06:30", "fim": None, "fim_motivo": "em aberto", "dias": 20, "h_sol": 211.0,
                     "perda_kwh": 2165.0, "perda_inferida_kwh": 300.0, "kwp_string": 11.2, "metodo": "quedas gravadas",
                     "flags": ["sem registro desde 23/09"]},
                    {"fonte": "sunop", "cliente": "Athon", "usina": "SMP100", "inversor": "Inversor 5.2", "string": "ST 19",
                     "inicio": f"{mes}-02 09:10", "fim": f"{mes}-02 15:00", "fim_motivo": "voltou", "dias": 1,
                     "h_sol": 5.8, "perda_kwh": 80.0, "perda_inferida_kwh": 0, "kwp_string": 10.0, "metodo": "curva",
                     "flags": []},
                ]},
            "trackers": {"rows": [
                {"fonte": "pv", "cliente": "Thopen", "usina": "Guatambu", "tracker": "TRK1", "inversor": "",
                 "inicio": f"{mes}-11 17:00", "fim": None, "h_sol": 127.6, "desvio_pico": 78.6, "fator": 0.8,
                 "kwp_tracker": 53.0, "perda_kwh": 2100.0, "cronico": False, "flags": ["em aberto", "sem inversor no BD_Trackers"]},
                {"fonte": "sunop", "cliente": "Athon", "usina": "CPP100", "tracker": "Tracker 41", "inversor": "Inversor 2.1",
                 "inicio": f"{mes}-03 08:00", "fim": f"{mes}-09 17:10", "h_sol": 60.0, "desvio_pico": 40.0, "fator": 0.234,
                 "kwp_tracker": 89.1, "perda_kwh": 700.0, "cronico": True, "flags": ["saiu: severo"]},
                {"fonte": "pv", "cliente": "Thopen", "usina": "Barretos", "tracker": "TRK9", "inversor": "Inversor 1.1",
                 "inicio": f"{mes}-03 08:00", "fim": f"{mes}-03 12:00", "h_sol": 4.0, "desvio_pico": 30.0, "fator": 0.134,
                 "kwp_tracker": 50.0, "perda_kwh": 20.0, "cronico": False, "flags": []},
            ]}}


def linhas(tabs, aba):
    cab = pub.CABECALHOS[aba]
    return [dict(zip(cab, lin)) for lin in tabs[aba]]


def test_cada_aba_com_o_cabecalho_na_linha_1():
    wb = load_workbook(io.BytesIO(pub.xlsx_bytes(pub.tabelas([pacote()]))))
    assert wb.sheetnames == list(pub.CABECALHOS)
    for aba, cab in pub.CABECALHOS.items():
        assert [c.value for c in wb[aba][1]] == cab
        assert wb[aba].freeze_panes == "A2"


def test_strings_por_inversor_e_dia_tem_as_colunas_pedidas_na_ordem_do_dia():
    l = linhas(pub.tabelas([pacote()]), "strings_inversor_dia")
    assert [x["dia"] for x in l] == ["2026-09-02", "2026-09-05"]            # pela data: o novo entra no fim
    assert (l[1]["cliente"], l[1]["usina"], l[1]["inversor"], l[1]["qtd_strings"], l[1]["strings"], l[1]["perda_kwh"]) == \
        ("Thopen", "Inhapi", "Inversor 3.3", 4, "Ipv1, Ipv2, Ipv3, Ipv4", 410.2)
    assert l[1]["fonte"] == "Thopen · API PV" and l[0]["fonte"] == "Athon"
    assert l[0]["avisos"] == "sem geração no dia (4,5 kWh/kWp); captura parcial" and l[1]["avisos"] is None


def test_string_em_aberto_fica_com_voltou_vazio():
    l = linhas(pub.tabelas([pacote()]), "strings_episodios")
    ab = next(x for x in l if x["string"] == "Ipv1")
    assert ab["saiu"] == "2026-09-05 06:30" and ab["voltou"] is None and ab["situacao"] == "em aberto"
    assert ab["avisos"] == "sem registro desde 23/09"


def test_situacao_do_tracker_e_ordem_pelo_inicio():
    l = linhas(pub.tabelas([pacote()]), "trackers_episodios")
    assert [(x["usina"], x["tracker"]) for x in l] == [("Barretos", "TRK9"), ("CPP100", "Tracker 41"), ("Guatambu", "TRK1")]
    sit = {x["tracker"]: x["situacao"] for x in l}
    assert sit == {"TRK9": "voltou a girar", "Tracker 41": "saiu: severo", "TRK1": "em aberto"}
    trk1 = next(x for x in l if x["tracker"] == "TRK1")
    assert trk1["parou"] == "2026-09-11 17:00" and trk1["voltou"] is None and trk1["inversor"] is None
    assert next(x for x in l if x["tracker"] == "Tracker 41")["cronico"] == "sim"


def test_varios_meses_viram_as_mesmas_abas_e_o_atualizado_em_fica_a_parte():
    tabs = pub.tabelas([pacote("2026-09"), pacote("2026-10", "2026-10-03 10:00")])
    assert [x["mes"] for x in linhas(tabs, "strings_episodios")] == ["2026-09", "2026-09", "2026-10", "2026-10"]
    assert all("atualizado" not in c for aba in ("strings_inversor_dia", "strings_episodios", "trackers_episodios")
               for c in pub.CABECALHOS[aba])
    at = linhas(tabs, "atualizacao")
    assert [(x["mes"], x["atualizado_em"], x["trackers_episodios"]) for x in at] == \
        [("2026-09", "2026-09-24 22:48", 3), ("2026-10", "2026-10-03 10:00", 3)]


def test_celula_vazia_e_vazia_de_verdade():
    wb = load_workbook(io.BytesIO(pub.xlsx_bytes(pub.tabelas([pacote()]))))
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            assert all(c.value != "" for c in row)          # o sync rejeita texto vazio


class _Resp:
    def __init__(self, status=200, js=None):
        self.status_code, self._js = status, js

    def json(self):
        return self._js

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(f"HTTP {self.status_code}")


class _Sessao:
    def __init__(self, existe=True, falha_sync=False):
        self.chamadas, self.existe, self.falha_sync = [], existe, falha_sync

    def get(self, url, **kw):
        self.chamadas.append(("GET", url, kw))
        return _Resp(200, [{"key": pub.WORKBOOK}] if self.existe else [{"key": "bd_performance"}])

    def post(self, url, **kw):
        self.chamadas.append(("POST", url, kw))
        if url.endswith("/sync-xlsx"):
            return _Resp(500 if self.falha_sync else 200, {"sheets": 4, "inserted": 9, "updated": 0, "deleted": 0})
        return _Resp(201, {"id": 50, "key": pub.WORKBOOK})


def test_sincronizar_cria_o_workbook_so_se_faltar_e_manda_replace_true():
    s = _Sessao(existe=False)
    assert pub.sincronizar(b"xlsx", base="https://api.teste/db/", token="tok", sessao=s)["sheets"] == 4
    assert [(m, u.split("/db")[1]) for m, u, _ in s.chamadas] == [
        ("GET", "/api/workbooks"), ("POST", "/api/workbooks"), ("POST", f"/api/workbooks/{pub.WORKBOOK}/sync-xlsx")]
    cria, sync = s.chamadas[1][2], s.chamadas[2][2]
    assert cria["json"]["key"] == pub.WORKBOOK and cria["json"]["display_name"]
    assert sync["params"] == {"replace": "true"} and sync["files"]["file"][1] == b"xlsx"
    assert sync["headers"]["Authorization"] == "Bearer tok" and sync["files"]["file"][2] == pub.MIME_XLSX
    s2 = _Sessao(existe=True)
    pub.sincronizar(b"x", base="https://api.teste/db", token="tok", sessao=s2)
    assert [m for m, _, _ in s2.chamadas] == ["GET", "POST"]            # já existe: não cria de novo
    with pytest.raises(RuntimeError):
        pub.sincronizar(b"x", base="https://api.teste/db", token="tok", sessao=_Sessao(falha_sync=True))


# ── Rodada 4 do Dashboard Thopen (Levi, 09/10/2026: "coluna para disponibilidade, calculada pelo Fracttal, da mesma forma
# que calculamos pelo gerencial da plataforma"): a disponibilidade de cada usina por dia fechado, pela régua da aba
# Disponibilidade do Gerencial. O 100 do dia sem OS vai ESCRITO — o 5080 nunca deduz um 100 pela falta da linha.
DISP = {"periodo": ["2026-10-01", "2026-10-04"],          # fim exclusivo: até D-1
        "usinas": [{"usina": "Guaratinguetá V", "fracttal": "Thopen - Guaratingueta 5 - SP", "cliente": "Thopen",
                    "pot_kwp": 2700.0},
                   {"usina": "Altair", "fracttal": "Thopen - Altair 1 - SP", "cliente": "Thopen", "pot_kwp": 6898.0}],
        "diario": {"Guaratinguetá V": {"2026-10-02": {"disp": 95.8, "h_eq": 0.509}}}}


def test_disponibilidade_por_dia_com_o_100_escrito_e_a_contagem_na_atualizacao():
    p = pacote("2026-10")
    p["disponibilidade"] = DISP
    tabs = pub.tabelas([p])
    d = linhas(tabs, "disponibilidade_dia")
    assert len(d) == 6                                     # 2 usinas × 3 dias fechados (01 a 03/10)
    gv = [(x["dia"], x["disp"], x["h_perdidas"]) for x in d if x["usina"] == "Guaratinguetá V"]
    assert gv == [("2026-10-01", 100.0, 0), ("2026-10-02", 95.8, 0.509), ("2026-10-03", 100.0, 0)]
    assert d[0]["usina_fracttal"] in ("Thopen - Altair 1 - SP", "Thopen - Guaratingueta 5 - SP")
    assert [x["dia"] for x in d] == sorted(x["dia"] for x in d)          # dia novo entra no fim
    (a,) = linhas(tabs, "atualizacao")
    assert a["disponibilidade_dia"] == 6


def test_mes_sem_disponibilidade_conta_vazio_e_nao_zero():
    (a,) = linhas(pub.tabelas([pacote("2026-09")]), "atualizacao")
    assert a["disponibilidade_dia"] is None and a["disponibilidade_os"] is None
    # pacote de antes de 12/10/2026 (com o diário e sem as OS): as OS contam vazio, não 0
    p = pacote("2026-10")
    p["disponibilidade"] = DISP
    (a,) = linhas(pub.tabelas([p]), "atualizacao")
    assert a["disponibilidade_dia"] == 6 and a["disponibilidade_os"] is None


def test_as_os_que_formaram_a_disponibilidade_uma_linha_por_os_e_usina():
    """Levi, 12/10/2026: "Quero também que apareça todas as OSs listadas que formaram aquele valor (...) a OS, que tipo
    de religamento foi, quanto tempo demorou, qual foi o total de perda de energia". Sem descrição nem quem criou.
    As OS daqui são do índice de antes do `kwh_usina` (13/10/2026): a de duas usinas reparte a perda pela potência
    afetada, que é o que dá para fazer sem a perda de cada usina."""
    p = pacote("2026-10")
    p["disponibilidade"] = dict(DISP, oss=[
        {"folio": 13001, "tipos": ["Religamento Remoto"], "origem": "queda", "code": "THPN-GTA500",
         "ini": "2026-10-02T09:00:00", "fim": "2026-10-02T15:06:00", "aberta": False, "usinas": ["Guaratinguetá V"],
         "escopo": [{"usina": "Guaratinguetá V", "rotulo": "Usina inteira", "kwp": 2700.0}], "h_solar": 6.1,
         "kwh": 16470},
        {"folio": 12990, "tipos": ["Corretiva Emergencial"], "origem": "equipamento", "code": "X-CABN1",
         "ini": "2026-10-01T08:00:00", "fim": None, "aberta": True, "usinas": ["Guaratinguetá V", "Altair"],
         "escopo": [{"usina": "Guaratinguetá V", "rotulo": "Cabine 1", "kwp": 300.0},
                    {"usina": "Altair", "rotulo": "Cabine 2", "kwp": 100.0}], "h_solar": 2.0, "kwh": 800}])
    tabs = pub.tabelas([p])
    d = linhas(tabs, "disponibilidade_os")
    assert [(x["os"], x["usina"], x["perda_kwh"]) for x in d] == [("12990", "Altair", 200), ("12990", "Guaratinguetá V", 600),
                                                                ("13001", "Guaratinguetá V", 16470)]
    g = d[2]
    assert (g["tipo"], g["origem"], g["escopo"], g["kwp"], g["duracao_sol_min"], g["aberta"]) == (
        "Religamento Remoto", "queda", "Usina inteira", 2700.0, 366, "não")
    assert g["usina_fracttal"] == "Thopen - Guaratingueta 5 - SP" and d[0]["aberta"] == "sim" and d[0]["fim"] is None
    (a,) = linhas(tabs, "atualizacao")
    assert a["disponibilidade_os"] == 3


def test_a_perda_de_cada_usina_e_a_do_indice_e_nao_a_reparticao_pela_potencia():
    """A OS de duas usinas leva a perda que o índice mediu em cada uma (`kwh_usina`, 13/10/2026), não a repartição pela
    potência: o teto de 100% pega uma usina e não a outra (e, até 15/10/2026, a geração desmentia o dia numa só) — em
    setembro a Rodrigues 1 e 2 trocavam ~47,9 MWh por essa conta. A usina do escopo em que a OS não perdeu nada (nenhuma
    hora de sol nela) segue com a linha, em 0: a linha não some (a API guarda histórico por linha)."""
    p = pacote("2026-10")
    beta = [{"usina": "Beta 1", "rotulo": "Usina inteira", "kwp": 500.0},
            {"usina": "Beta 2", "rotulo": "Usina inteira", "kwp": 500.0}]
    p["disponibilidade"] = dict(DISP, oss=[
        {"folio": 20, "tipos": ["Religamento"], "origem": "queda", "code": "TC-BET100", "ini": "2026-10-01T06:00:00",
         "fim": "2026-10-02T18:00:00", "aberta": False, "usinas": ["Beta 1", "Beta 2"], "escopo": beta,
         "h_solar": 24.0, "kwh": 18000, "kwh_usina": {"Beta 1": 6000, "Beta 2": 12000}},
        {"folio": 21, "tipos": ["Religamento"], "origem": "queda", "code": "TC-BET100", "ini": "2026-10-03T06:00:00",
         "fim": "2026-10-03T12:00:00", "aberta": False, "usinas": ["Beta 1", "Beta 2"], "escopo": beta,
         "h_solar": 6.0, "kwh": 3000, "kwh_usina": {"Beta 2": 3000}}])
    d = linhas(pub.tabelas([p]), "disponibilidade_os")
    assert [(x["os"], x["usina"], x["kwp"], x["perda_kwh"]) for x in d] == [
        ("20", "Beta 1", 500.0, 6000), ("20", "Beta 2", 500.0, 12000),
        ("21", "Beta 1", 500.0, 0), ("21", "Beta 2", 500.0, 3000)]


def test_a_parcela_de_cada_os_em_cada_dia_soma_a_perda_do_dia():
    """Dashboard Thopen, 14/10/2026 (Levi: "A partir de outubro quero os comentários 100% OS Fracttal. Comentários e
    disponibilidade!"): a aba `disponibilidade_os_dia` leva a parcela de cada OS em cada dia, da MESMA conta da
    `disponibilidade_dia` — o dia que a geração desmentiu não tem linha, e o escopo e o kWp são os da OS naquela usina."""
    p = pacote("2026-10")
    oss = [{"folio": 13001, "tipos": ["Religamento Remoto"], "origem": "queda", "code": "THPN-GTA500",
            "ini": "2026-10-02T09:00:00", "fim": "2026-10-02T15:06:00", "aberta": False, "usinas": ["Guaratinguetá V"],
            "escopo": [{"usina": "Guaratinguetá V", "rotulo": "Usina inteira", "kwp": 2700.0}], "h_solar": 6.1,
            "kwh": 16470},
           {"folio": 12990, "tipos": ["Corretiva Emergencial"], "origem": "equipamento", "code": "X-CABN1",
            "ini": "2026-10-01T08:00:00", "fim": None, "aberta": True, "usinas": ["Guaratinguetá V", "Altair"],
            "escopo": [{"usina": "Guaratinguetá V", "rotulo": "Cabine 1", "kwp": 300.0},
                       {"usina": "Altair", "rotulo": "Cabine 2", "kwp": 100.0}], "h_solar": 2.0, "kwh": 800}]
    diario = {"Guaratinguetá V": {"2026-10-02": {"disp": 49.2, "h_eq": 6.1, "os": [
        {"folio": 12990, "h": 0.0, "trechos": "06:00-09:00"},
        {"folio": 13001, "h": 6.1, "trechos": "09:00-15:06"}]}},
        "Altair": {"2026-10-01": {"disp": 99.5, "h_eq": 0.058, "os": [
            {"folio": 12990, "h": 0.058, "trechos": "08:00-12:00"}]}}}
    p["disponibilidade"] = dict(DISP, oss=oss, diario=diario)
    tabs = pub.tabelas([p])
    d = linhas(tabs, "disponibilidade_os_dia")
    assert [(x["dia"], x["usina"], x["os"], x["escopo"], x["kwp"], x["trechos"], x["h_perdidas"]) for x in d] == [
        ("2026-10-01", "Altair", "12990", "Cabine 2", 100.0, "08:00-12:00", 0.058),
        ("2026-10-02", "Guaratinguetá V", "12990", "Cabine 1", 300.0, "06:00-09:00", 0.0),
        ("2026-10-02", "Guaratinguetá V", "13001", "Usina inteira", 2700.0, "09:00-15:06", 6.1)]
    g = d[2]
    assert (g["tipo"], g["origem"], g["inicio"], g["fim"], g["aberta"], g["usina_fracttal"]) == (
        "Religamento Remoto", "queda", "2026-10-02T09:00:00", "2026-10-02T15:06:00", "não", "Thopen - Guaratingueta 5 - SP")
    (a,) = linhas(tabs, "atualizacao")
    assert a["disponibilidade_os_dia"] == 3


def test_indice_sem_a_parcela_do_dia_conta_vazio_e_nao_zero():
    # o índice publicado antes do código de 14/10/2026 não tem a chave "os" nos dias: a aba sai vazia e a atualizacao
    # diz VAZIO (não publicada) — o 5080 não pode ler isso como "nenhuma OS"
    p = pacote("2026-10")
    p["disponibilidade"] = DISP
    tabs = pub.tabelas([p])
    assert linhas(tabs, "disponibilidade_os_dia") == []
    (a,) = linhas(tabs, "atualizacao")
    assert a["disponibilidade_os_dia"] is None
