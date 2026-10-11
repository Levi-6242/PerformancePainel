# -*- coding: utf-8 -*-
"""Quem entra no tempo real, fonte por fonte (Levi, 11/10/2026): a regra pura de `cobertura.py`, sem rede.

Regras do Levi: usina nova espera ele liberar (e o sino avisa); a escolha vale no Tempo real, Monitoramento, Painel NOC e
ronda; usina tirada para de ser coletada. A regra de segurança: fonte sem as linhas de migração segue como antes, e
decisões que não puderam ser lidas não filtram nada (sumiço em massa é pior que mostrar a mais)."""
import json

import pytest

import cobertura as cb

Q1, Q2, Q3 = "2026-10-11T12:00:00Z", "2026-10-11T13:00:00Z", "2026-10-11T14:00:00Z"


def test_a_linha_tem_o_dia_de_brasilia_e_recusa_o_que_nao_e_do_dominio():
    assert cb.linha("pv", 123, "incluir", "tela", "2026-10-12T02:30:00Z") == \
        ["2026-10-12T02:30:00Z", 20261011, "pv", "123", "incluir", "tela"]      # 02:30Z é o dia anterior em Brasília
    for ruim in (("xx", 1, "incluir", "tela"), ("pv", 1, "talvez", "tela"), ("pv", 1, "incluir", "eu"),
                 ("pv", "  ", "incluir", "tela")):
        with pytest.raises(ValueError):
            cb.linha(*ruim, Q1)


def test_vale_a_ultima_decisao_de_cada_fonte_e_chave():
    linhas = [cb.linha("pv", 1, "incluir", "migracao", Q1), cb.linha("pv", 1, "excluir", "tela", Q2),
              cb.linha("pv", 2, "excluir", "migracao", Q1), cb.linha("pv", 2, "incluir", "tela", Q3),
              cb.linha("sunop", 1, "incluir", "migracao", Q1)]
    assert cb.vigente(linhas) == {("pv", "1"): "excluir", ("pv", "2"): "incluir", ("sunop", "1"): "incluir"}
    # no mesmo instante, vale a que veio depois
    assert cb.vigente([cb.linha("pg", 9, "incluir", "tela", Q1), cb.linha("pg", 9, "excluir", "tela", Q1)]) == \
        {("pg", "9"): "excluir"}
    # linha estranha (fonte fora do domínio, decisão vazia) não conta
    assert cb.vigente([["x", 1, "zz", "1", "incluir", "tela"], [Q1, 1, "pv", "3", None, "tela"]]) == {}


def test_fonte_sem_migracao_nao_e_filtrada():
    usinas = [{"id": 1}, {"id": 2}]
    linhas = [cb.linha("pv", 1, "excluir", "tela", Q1)]                         # decisão solta, sem migração
    assert cb.filtrar("pv", usinas, lambda u: u["id"], linhas) == (usinas, [])
    assert cb.governadas(linhas) == set()


def test_decisoes_indisponiveis_nao_filtram_nada():
    usinas = [{"id": 1}, {"id": 2}]
    assert cb.filtrar("pv", usinas, lambda u: u["id"], None) == (usinas, [])


def test_fonte_governada_mostra_so_as_incluidas_e_separa_as_novas():
    linhas = cb.semente("pv", inventario=[1, 2, 3], visiveis=[1, 2], quando=Q1) + [cb.linha("pv", 2, "excluir", "tela", Q2)]
    usinas = [{"id": i} for i in (1, 2, 3, 4)]
    entram, novas = cb.filtrar("pv", usinas, lambda u: u["id"], linhas)
    assert entram == [{"id": 1}]                     # 2 saiu pela tela, 3 já estava fora no dia da troca
    assert novas == [{"id": 4}]                      # apareceu depois: espera a liberação
    assert cb.governadas(linhas) == {"pv"}


def test_a_semente_reproduz_o_que_aparecia_no_dia_da_troca():
    linhas = cb.semente("solaredge", inventario=["a", "b", "c"], visiveis=["a", "c"], quando=Q1)
    assert [(lin[3], lin[4], lin[5]) for lin in linhas] == \
        [("a", "incluir", "migracao"), ("b", "excluir", "migracao"), ("c", "incluir", "migracao")]
    itens = ["a", "b", "c"]
    assert cb.filtrar("solaredge", itens, str, linhas) == (["a", "c"], [])          # nada muda na tela
    # visível que não estava no inventário (lista fixa do código) também entra
    assert [(lin[3], lin[4]) for lin in cb.semente("semp", inventario=[], visiveis=[7], quando=Q1)] == [("7", "incluir")]


def test_o_inventario_troca_a_lista_da_fonte_e_avisa_cada_nova_uma_vez(tmp_path):
    inv = cb.Inventario(str(tmp_path / "inv.json"))
    inv.registrar("pv", [(1, "Altair 1"), (2, "Altair 2")])
    inv.registrar("pv", [(1, "Altair 1"), (3, "Nova Usina")])                    # a 2 saiu da conta
    d = inv.ler()
    assert sorted(d["fontes"]["pv"]) == ["1", "3"] and d["fontes"]["pv"]["3"]["nome"] == "Nova Usina"
    assert inv.a_avisar("pv", ["3"]) == ["3"]
    assert inv.a_avisar("pv", ["3"]) == []                                       # avisada uma vez só
    assert inv.a_avisar("sunop", ["3"]) == ["3"]                                 # outra fonte, outro aviso
    inv.registrar("xx", [(9, "x")])                                              # fonte fora do domínio: ignorada
    assert "xx" not in inv.ler()["fontes"]


# ── as decisões no banco (a API de mentira) ──────────────────────────────────────────────────────────────────────────
class _R:
    def __init__(self, dados, status=200):
        self._d, self.status_code = dados, status

    def json(self):
        return self._d

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(f"HTTP {self.status_code}")


class _API:
    """A db_performace de mentira: workbooks, sheets, rows e o sync-xlsx (lê o xlsx enviado de verdade)."""

    def __init__(self):
        self.livros, self.linhas, self.fora, self.engole = set(), None, False, 0

    def get(self, url, headers=None, params=None, timeout=None):
        if self.fora:
            raise ConnectionError("banco fora")
        if url.endswith("/api/workbooks"):
            return _R([{"key": k} for k in self.livros])
        if url.endswith("/api/sheets"):
            return _R([{"id": 7, "workbook_key": cb.LIVRO, "sheet_name": cb.ABA}] if self.linhas is not None else [])
        if url.endswith("/api/sheets/7/rows"):
            off, lim = params["offset"], params["limit"]
            pedaco = (self.linhas or [])[off:off + lim]
            return _R({"rows": [{"headers": cb.CAB, "values": lin} for lin in pedaco]})
        raise AssertionError(url)

    def post(self, url, headers=None, json=None, params=None, files=None, timeout=None):
        assert headers["Authorization"] == "Bearer tok"
        if url.endswith("/api/workbooks"):
            self.livros.add(json["key"])
            return _R({})
        if url.endswith(f"/api/workbooks/{cb.LIVRO}/sync-xlsx"):
            assert params == {"replace": "true"}
            from io import BytesIO
            from openpyxl import load_workbook
            ws = load_workbook(BytesIO(files["file"][1]))[cb.ABA]
            linhas = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
            self.linhas = linhas[:len(linhas) - self.engole]
            return _R({"inserted": len(linhas)})
        raise AssertionError(url)


def test_acrescentar_rele_grava_inteiro_e_confere(tmp_path):
    api = _API()
    d = cb.Decisoes("https://x/db", "tok", str(tmp_path / "dec.json"), sessao=api)
    assert d.linhas() == [] and cb.governadas(d.linhas()) == set()               # livro ainda não existe
    assert d.acrescentar(cb.semente("pv", [1, 2], [1], Q1)) == 2
    assert cb.LIVRO in api.livros
    assert d.acrescentar([cb.linha("pv", 2, "incluir", "tela", Q2)]) == 3        # só acrescenta: as 2 de antes ficam
    assert cb.vigente(d.linhas()) == {("pv", "1"): "incluir", ("pv", "2"): "incluir"}
    api.engole = 1                                                               # o banco perdeu uma linha na gravação
    with pytest.raises(cb.GravacaoErro):
        d.acrescentar([cb.linha("pv", 1, "excluir", "tela", Q3)])


def test_sem_token_nao_grava(tmp_path):
    d = cb.Decisoes("https://x/db", "", str(tmp_path / "dec.json"), sessao=_API())
    with pytest.raises(cb.GravacaoErro, match="GRIDCO_SQL_TOKEN"):
        d.acrescentar([cb.linha("pv", 1, "incluir", "tela", Q1)])


def test_banco_fora_vale_a_ultima_leitura_boa_da_memoria_e_do_disco(tmp_path):
    api = _API()
    api.linhas = cb.semente("pv", [1, 2], [1], Q1)
    caminho = str(tmp_path / "dec.json")
    d = cb.Decisoes("https://x/db", "tok", caminho, ttl=0, sessao=api)
    assert len(d.linhas()) == 2 and json.loads(open(caminho, encoding="utf-8").read())["linhas"]
    api.fora = True
    assert len(d.linhas()) == 2 and d.erro == "ConnectionError"                  # memória
    outro = cb.Decisoes("https://x/db", "tok", caminho, ttl=0, sessao=api)       # processo novo, banco fora: o disco
    assert len(outro.linhas()) == 2
    nunca = cb.Decisoes("https://x/db", "tok", str(tmp_path / "nada.json"), sessao=api)
    assert nunca.linhas() is None                                                # nada lido: quem chama não filtra


# ── a ligação com o app ──────────────────────────────────────────────────────────────────────────────────────────────
import app  # noqa: E402
from test_porta_nexus import _sem_rede, cli, entrar, entrar_com_senha  # noqa: E402,F401  (fixtures e passos da porta)


@pytest.fixture
def ligado(monkeypatch, tmp_path):
    """O filtro ligado, com o banco de mentira e a chave de gravação."""
    api = _API()
    monkeypatch.setattr(app, "COBERTURA_LIGADA", True)
    monkeypatch.setattr(app, "_COB_DECISOES", cb.Decisoes("https://x/db", "tok", str(tmp_path / "dec.json"), ttl=0, sessao=api))
    monkeypatch.setenv("GRIDCO_SQL_TOKEN", "tok")
    monkeypatch.setattr(app, "FULL_OM", {"Altair 1", "Altair 2"})
    return api


PLANTAS = [{"id": 11, "nome": "Altair 1"}, {"id": 12, "nome": "Altair 2"}, {"id": 13, "nome": "Sem Full OM"},
           {"id": 18758732, "nome": "Tucano 1"}]                               # a última é da SEMP (PV_FONTES)


def test_a_api_pv_segue_o_full_om_ate_a_foto_de_hoje_e_depois_obedece_a_tela(ligado):
    assert [p["id"] for p in app._pv_plantas_da_fonte(PLANTAS, None)] == [11, 12]          # regra de antes
    inv = app._COB_INV.ler()["fontes"]["pv"]
    assert sorted(inv) == ["11", "12", "13"] and inv["13"]["nome"] == "Sem Full OM"        # o inventário vê tudo
    ligado.linhas = cb.semente("pv", ["11", "12", "13"], ["11", "12"], Q1) + [cb.linha("pv", 12, "excluir", "tela", Q2)]
    assert [p["id"] for p in app._pv_plantas_da_fonte(PLANTAS, None)] == [11]              # a 12 saiu pela tela
    # os outros montadores da API PV perguntam a mesma coisa
    # a da SEMP (outra conta) segue a regra de antes desses montadores: o Full O&M
    assert [app._pv_monitorada(p) for p in PLANTAS] == [True, False, False, False]
    # a da SEMP (lista fixa) não muda enquanto a fonte semp não passa pela foto
    assert [p["id"] for p in app._pv_plantas_da_fonte(PLANTAS, "semp")] == [18758732]


def test_usina_nova_numa_fonte_controlada_espera_e_o_sino_avisa_uma_vez(ligado):
    ligado.linhas = cb.semente("pv", ["11", "12", "13"], ["11", "12"], Q1)
    novas = PLANTAS[:3] + [{"id": 14, "nome": "Usina Nova"}]
    assert [p["id"] for p in app._pv_plantas_da_fonte(novas, None)] == [11, 12]            # a 14 espera a liberação
    app._pv_plantas_da_fonte(novas, None)
    evs = [e for e in (app._load_state().get("notif") or {}).get("eventos", []) if e.get("tipo") == "usina_nova"]
    assert [(e["fonte"], e["plant_id"], e["usina"]) for e in evs] == [("pv", "14", "Usina Nova")]   # uma vez só
    assert evs[0]["rotulo"] == "Thopen · API PV"


def test_desligado_nada_muda_nem_le_o_banco(monkeypatch):
    monkeypatch.setattr(app, "FULL_OM", {"Altair 1"})
    assert [p["id"] for p in app._pv_plantas_da_fonte(PLANTAS, None)] == [11]
    assert app._cob_linhas() is None and app._cob_entra("pv", 13) is True


def test_a_tela_lista_migra_e_decide(cli, ligado):
    entrar_com_senha(cli)
    app._pv_plantas_da_fonte(PLANTAS, None)                                               # a coleta leu a fonte
    d = cli.get("/api/admin/tempo-real/fontes").get_json()
    pv = next(f for f in d["fontes"] if f["id"] == "pv")
    assert d["decisoes"]["ok"] and d["decisoes"]["pode_gravar"] and pv["governada"] is False
    assert {u["chave"]: u["situacao"] for u in pv["usinas"]} == {"11": "hoje", "12": "hoje", "13": "fora_antes"}
    # decidir antes da foto não vale
    r = cli.post("/api/admin/tempo-real/decidir", json={"decisoes": [{"fonte": "pv", "chave": "13", "decisao": "incluir"}]})
    assert r.status_code == 409
    # a prévia não grava; a confirmação grava a foto
    p = cli.post("/api/admin/tempo-real/migrar", json={"fonte": "pv"}).get_json()
    assert p["previa"] and (p["incluir"], p["excluir"]) == (2, 1) and ligado.linhas is None
    assert cli.post("/api/admin/tempo-real/migrar", json={"fonte": "pv", "confirmar": True}).get_json()["total"] == 3
    assert cli.post("/api/admin/tempo-real/migrar", json={"fonte": "pv", "confirmar": True}).status_code == 409
    r = cli.post("/api/admin/tempo-real/decidir", json={"decisoes": [{"fonte": "pv", "chave": "13", "decisao": "incluir"}]})
    assert r.status_code == 200 and r.get_json()["total"] == 4
    pv = next(f for f in cli.get("/api/admin/tempo-real/fontes").get_json()["fontes"] if f["id"] == "pv")
    assert pv["governada"] and {u["chave"]: u["situacao"] for u in pv["usinas"]} == \
        {"11": "incluida", "12": "incluida", "13": "incluida"}
    assert next(u for u in pv["usinas"] if u["chave"] == "13")["origem"] == "tela"
    assert [p["id"] for p in app._pv_plantas_da_fonte(PLANTAS, None)] == [11, 12, 13]
    # decisão fora do domínio
    r = cli.post("/api/admin/tempo-real/decidir", json={"decisoes": [{"fonte": "pv", "chave": "13", "decisao": "talvez"}]})
    assert r.status_code == 400


def test_a_tela_e_so_de_administrador_do_nexus(cli, ligado):
    entrar(cli, "analista@exemplo.com", destino="/admin/tempo-real", admin=False)
    assert cli.get("/admin/tempo-real").status_code == 403
    assert cli.get("/api/admin/tempo-real/fontes").status_code == 403
    assert cli.post("/api/admin/tempo-real/decidir", json={"decisoes": []}).status_code == 403
    entrar(cli, "admin@exemplo.com", destino="/admin/tempo-real", admin=True)
    h = cli.get("/admin/tempo-real").get_data(as_text=True)
    assert "<h1>Fontes do tempo real</h1>" in h and 'href="/" data-casca' in h
    assert cli.get("/api/admin/tempo-real/fontes").status_code == 200


def test_a_2c_tira_o_codigo_so_quando_todas_as_plantas_dele_saem(ligado):
    plantas = [str(p) for ps in app._2C_STR_PLANTAS.values() for p in ps]
    ligado.linhas = cb.semente("2capi", plantas, plantas, Q1)
    assert app._2c_codigos() == list(app.OWEN_UFVS)
    ligado.linhas += [cb.linha("2capi", 18771898, "excluir", "tela", Q2),            # Araputanga (1 planta)
                      cb.linha("2capi", 18771915, "excluir", "tela", Q2)]            # 1 das 3 da Ipixuna
    assert "ARA" not in app._2c_codigos() and "IPX" in app._2c_codigos()


def test_sunop_os_lacos_perguntam_ao_meta_filtrado(ligado, monkeypatch):
    meta = {"MRO100": {"trackers": 1}, "MAB100": {"trackers": 1}}
    monkeypatch.setattr(app, "_si", lambda inst="gridco": {"meta": meta})
    assert app._sunop_meta_monitorada("gridco") is meta                             # não governada: o próprio meta
    ligado.linhas = cb.semente("sunop", list(meta), list(meta), Q1) + [cb.linha("sunop", "MAB100", "excluir", "tela", Q2)]
    assert list(app._sunop_meta_monitorada("gridco")) == ["MRO100"]
    assert app._sunop_meta_monitorada("axis") is meta                               # a Axis é outra fonte
    assert app._sunop_fonte("axis") == "axis" and app._sunop_fonte("gridco") == "sunop"


def test_trackers_da_api_pv_e_o_rollup_seguem_a_decisao(ligado, monkeypatch):
    monkeypatch.setattr(app, "_pv_trk_plant", {11: {"payload": {"total": 3}}, 12: {"payload": {"total": 2}}})
    ligado.linhas = cb.semente("pv", ["11", "12"], ["11", "12"], Q1) + [cb.linha("pv", 12, "excluir", "tela", Q2)]
    assert [pid for pid, _e in app._pv_trk_itens()] == [11]
    assert app._rollup_cob_fonte("API PV", 11) == "pv" and app._rollup_cob_fonte("API PV", 18758732) == "semp"
    assert app._rollup_cob_fonte("Thopen", 5) == "pg" and app._rollup_cob_fonte("RenoGrid", 9) == "solaredge"
    assert app._cob_entra("pv", 12) is False and app._cob_entra("pg", 99) is True        # pg não governada


def test_sunop_lista_em_cache_tambem_registra_o_inventario(ligado, monkeypatch):
    """Na pausa de domingo/teto (e com a lista de 10 min em cache) o ensure_sunop_meta volta antes da rede: o inventário
    da Athon tem de sair dali também, senão a foto de hoje da Athon nunca teria o que fotografar num domingo."""
    import time as _t
    S = {"meta": {"MRO100": {}, "MAB100": {}}}
    monkeypatch.setattr(app, "_si", lambda inst="gridco": S)
    monkeypatch.setitem(app._sunop_plants_lista, "gridco", {"nomes": ["MRO100", "MAB100"], "ts": _t.time()})
    app.ensure_sunop_meta("gridco")
    assert sorted(app._COB_INV.ler()["fontes"]["sunop"]) == ["MAB100", "MRO100"]
