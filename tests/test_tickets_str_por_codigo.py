# -*- coding: utf-8 -*-
"""Tickets de strings em TODAS as fontes: o ticket que o nome não acha vai pelo CÓDIGO da usina (Levi, 24/09/2026).

O pedido: *"Como estávamos usando só Athon como teste, implemente para os demais!"*. Medido em 24/09 nas 9 fontes da
tabela de strings: a coluna, o card no inversor e a OS de recomposição já funcionam em Thopen API PV, Thopen Banco,
Athon e RenoGrid (todo ticket de inversor casou com um inversor do drill). O que faltava era o CASAMENTO da usina:
  - CNN100 (2 tickets): a planilha escreve o código, e a aba "Base de dados - Usinas" tem a Canarana com código 0;
  - ADS100 (1 ticket): a aba de usinas chama de "Araçoiaba da Serra IA"; a plataforma, de "Araçoiaba da Serra 1".
O de-para mestre nome → código é a aba "Info Geral" do BD_Performance (USINA_COD), e ele estava VAZIO desde 30/07:
alguém inseriu uma linha acima do cabeçalho, o leitor de metas passou a procurar o cabeçalho, e o load_usina_codigos
não — lia a linha 1, não achava "Código Fractal" e saía calado. Com o cabeçalho certo: 142 códigos.

A reserva pelo código só vale para a linha cujo código é SÓ DELA na tabela: Altair 1 a 5 dividem ALT100, e cada parte
tem os tickets dela pelo nome + o 1º número do inversor. Pelo código, o ticket da Altair 1 cairia nas cinco.
"""
import io

import openpyxl
import pytest

import app

COLS = ["", "Usina", "Inversor", "Quantidade de strings no afetadas", "Causa raiz", "Início da ocorrência",
        "Fim da ocorrência", "Comentários gerais"]
LINHAS = [   # (usina como a planilha escreve, inversor, afetadas, comentário)
    ("CNN100", "Inversor 1.2", 2, "Strings 3 e 4 com corrente nula"),       # linha 5: código sem de-para na aba de usinas
    ("ADS100", "Inversor 1.1", 1, "String 7 com corrente nula"),            # linha 6: a aba chama de "Araçoiaba da Serra IA"
    ("ALT100", "Inversor 1.3", 1, "String 2 com corrente nula"),            # linha 7: Altair 1 pelo nome + bloco
    ("ALT100", "Inversor 9.1", 1, "String 5 com corrente nula"),            # linha 8: bloco que nenhuma parte tem
    ("Canarana 2", "Inversor 2.1", 1, "a verificar"),                       # linha 9: casa pelo nome
    ("CNN200", "Inversor 2.4", 1, "String 1 com corrente nula"),            # linha 10: órfão da Canarana 2 (código)
    ("CGU100", "Inversor 1.1", 1, "String 12 com corrente nula em 09/09/2026 08:00"),   # linha 11: Cipó Guaçu, em 3 partes
]
USINAS = [("Araçoiaba da Serra IA", "THPN-ADS100"), ("Altair", "THPN-ALT100"), ("Canarana 1", "0"), ("Canarana 2", "0")]
INFO_GERAL = [("Canarana 1", "Thopen", "CNN100", "UFV Canarana 1"), ("Canarana 2", "Thopen", "CNN200", "UFV Canarana 2"),
              ("Araçoiaba da Serra 1", "Thopen", "ADS100", "UFV Araçoiaba IA"), ("Altair 1", "Thopen", "ALT100", "UFV Altair"),
              ("Altair 2", "Thopen", "ALT100", "UFV Altair"), ("Cipó Guaçu", "Thopen", "CGU100", "UFV Cipó Guaçu")]


def _tickets_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Strings indisp"
    for _ in range(3):
        ws.append(["rascunho"])
    ws.append(COLS)
    for u, inv, q, com in LINHAS:
        ws.append(["", u, inv, q, None, "2026-09-20 06:00", None, com])
    wu = wb.create_sheet("Base de dados - Usinas")
    for _ in range(3):
        wu.append(["rascunho"])
    wu.append(["", "Usina", "Código da usina"])
    for n, c in USINAS:
        wu.append(["", n, c])
    b = io.BytesIO()
    wb.save(b)
    b.seek(0)
    return b


def _bd_xlsx(linhas_acima=1):
    """Info Geral como está no BD_Performance desde 30/07: uma linha a mais ACIMA do cabeçalho."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Info Geral"
    for _ in range(linhas_acima):
        ws.append(["Relação de usinas — atualizar a cada onboarding"])
    ws.append(["Usina", "Cliente", "Código Fractal", "Descrição Fractal"])
    for r in INFO_GERAL:
        ws.append(list(r))
    b = io.BytesIO()
    wb.save(b)
    b.seek(0)
    return b


@pytest.fixture
def base(monkeypatch):
    monkeypatch.setattr(app, "_bd_readable",
                        lambda chave=None, fallback=None: _tickets_xlsx() if chave == "tickets_performance" else _bd_xlsx())
    monkeypatch.setattr(app, "_tickets_marca", lambda: 1.0)
    monkeypatch.setattr(app, "TICKETS_STR", {})
    monkeypatch.setattr(app, "TICKETS_STR_COD", {}, raising=False)
    monkeypatch.setattr(app, "USINA_COD", {})
    app.load_usina_codigos()
    app.load_tickets_strings()


def _linhas(*nomes):
    rows = [{"usina": n, "plant_id": i} for i, n in enumerate(nomes)]
    return {r["usina"]: r.get("tickets_str") for r in app._com_tickets_str({"rows": rows, "summary": {}})["rows"]}


@pytest.mark.parametrize("acima", [0, 1, 3])
def test_info_geral_acha_o_cabecalho_mesmo_com_linha_acima(monkeypatch, acima):
    """30/07: uma linha inserida acima do cabeçalho zerou o de-para em silêncio (o log de metas dizia "ajustado")."""
    monkeypatch.setattr(app, "_bd_readable", lambda chave=None, fallback=None: _bd_xlsx(acima))
    monkeypatch.setattr(app, "USINA_COD", {})
    app.load_usina_codigos()
    assert app.USINA_COD[app._usina_key("Canarana 1")] == "CNN100"
    assert app.USINA_COD[app._usina_key("UFV Araçoiaba IA")] == "ADS100", "a descrição Fracttal também é chave"
    assert app.USINA_COD[app._usina_key("CNN200")] == "CNN200", "o próprio código também"


def test_codigo_sem_de_para_na_aba_de_usinas_casa_pela_info_geral(base):
    """CNN100 em 24/09: 2 tickets que não apareciam em fonte nenhuma, com a Canarana 1 na tabela da API PV."""
    ts = _linhas("Canarana 1", "Altair 1 (73)")["Canarana 1"]
    assert ts and [t["linha"] for t in ts["lista"]] == [5] and ts["strings"] == 2


def test_nome_diferente_na_aba_de_usinas_casa_pela_info_geral(base):
    """"Araçoiaba da Serra IA" na aba de usinas, "Araçoiaba da Serra 1" na plataforma: o código ADS100 é dos dois."""
    ts = _linhas("Araçoiaba da Serra 1", "Araçoiaba da Serra 2")
    assert [t["linha"] for t in ts["Araçoiaba da Serra 1"]["lista"]] == [6] and ts["Araçoiaba da Serra 2"] is None


def test_codigo_dividido_entre_partes_nao_espalha_o_ticket(base):
    """Altair 1 e 2 dividem ALT100. O ticket do Inversor 1.3 é da Altair 1 pelo nome + bloco; o do Inversor 9.1 não é de
    nenhuma parte que a tabela mostra — pelo código ele cairia nas duas, e quem somasse contaria duas vezes."""
    ts = _linhas("Altair 1 (73)", "Altair 2 (74)")
    assert [t["linha"] for t in ts["Altair 1 (73)"]["lista"]] == [7]
    assert ts["Altair 2 (74)"] is None


def test_orfao_pelo_codigo_entra_junto_com_os_do_nome(base):
    """Canarana 2 tem um ticket pelo nome e um órfão pelo código (CNN200): os dois são dela."""
    ts = _linhas("Canarana 2")["Canarana 2"]
    assert sorted(t["linha"] for t in ts["lista"]) == [9, 10] and ts["tickets"] == 2


def test_ticket_que_o_nome_ja_achou_nao_volta_pelo_codigo(base):
    """Ticket casado pelo nome em uma linha não é órfão: não pode aparecer de novo em outra pelo código."""
    ts = _linhas("CNN100", "Canarana 1")
    assert [t["linha"] for t in ts["CNN100"]["lista"]] == [5] and ts["Canarana 1"] is None


def test_codigo_de_usina_em_partes_vai_para_a_parte_do_inversor(base):
    """CGU100 (ticket 315 em 24/09, "String 12 com corrente nula"): o Info Geral diz Cipó Guaçu, e a API PV mostra três
    partes, Cipó-Guaçu 1, 2 e 3. O código é das três, então a reserva pelo código único não servia; o 1º número do
    inversor diz a parte — o Inversor 1.1 é da Cipó-Guaçu 1. Confirmado pelo Levi: "É a Cipó Guaçu"."""
    ts = _linhas("Cipó-Guaçu 1 (106)", "Cipó-Guaçu 2 (107)", "Cipó-Guaçu 3 (108)")
    assert [t["linha"] for t in ts["Cipó-Guaçu 1 (106)"]["lista"]] == [11]
    assert ts["Cipó-Guaçu 2 (107)"] is None and ts["Cipó-Guaçu 3 (108)"] is None


def test_hifen_no_nome_da_linha_nao_esconde_o_codigo(base):
    """A API PV escreve "Cipó-Guaçu", o Info Geral "Cipó Guaçu"."""
    assert app._tk_str_cod_da_linha("Cipó-Guaçu 2 (107)") == "CGU100"


def test_parte_com_numero_composto_nao_adivinha(base):
    """"Ceilandia 1.2 (90)": o fim do nome não é o número de uma parte. Sem dizer qual parte, não cai em nenhuma."""
    assert app._tk_str_parte("Ceilandia 1.2 (90)") is None and app._tk_str_parte("Cipó-Guaçu 3 (108)") == 3
    assert app._tk_str_parte("Brodowski - Skid 2 (86)") == 2 and app._tk_str_parte("Canarana") is None
