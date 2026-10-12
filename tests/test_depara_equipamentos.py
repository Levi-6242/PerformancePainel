# -*- coding: utf-8 -*-
"""De-para de equipamentos (Levi, 11/10/2026: "quero que o de-para da equipamentos seja feito na plataforma de forma
intuitiva"): o diagnóstico (`depara_diag.py`), as decisões da tela (`depara.py`, livro nexus_depara) que vencem o de-para
do código, o editor, a sugestão pelo valor e a tela só de administrador.

Os casos são os do mesmo dia: a API renomeou a Ipixuna e a Morada Nova e as duas ficaram sem esperadas (string morta não
acusava); a União tinha 6 dos 12 inversores no cadastro da ponte; e a conta oem@ nomeia inversor pela ordem dos ids quando
não há de-para (na Morada Nova a ordem errava 6 de 8)."""
import pytest

import depara as dp
import depara_diag as dd

CTX = {"nome_api": {18771915: "Ipixuna do Pará 1"}, "ponte_ids": {18771915: "Santa Cecilia 1"},
       "apelidos": {"Sete Lagoa": "Sete Lagoas"}, "inv_nomes": {18771915: {400840: "INVERSOR0 1.1"},
                                                                 18772125: {401313: "UNI_Inv_2.1"}},
       "provisorios": {18772125: "União: o grupo 1 pela ordem"}, "oem": {18771915, 18772125, 99},
       "sg_por_ordem": {("salenergia", "boa"): True}}
Q1, Q2, Q3 = "2026-10-11T12:00:00Z", "2026-10-11T13:00:00Z", "2026-10-11T14:00:00Z"


def _r(pid, usina, str_esp=100, inv_esp=6, inv=6):
    return {"plant_id": pid, "usina": usina, "str_esp": str_esp, "inv_esp": inv_esp, "qtd_inversores": inv}


def test_a_planta_renomeada_liga_pelo_id_e_diz_o_nome_que_a_api_escreve():
    u = dd.linha("2capi", _r(18771915, "Santa Cecilia 1", 138, 8, 8), CTX)
    assert u["liga"] == "pelo id (ponte): a API escreve “Ipixuna do Pará 1”"
    assert u["nomes_inversor"] == "de-para no código" and u["alertas"] == [] and u["pior"] == "ok"


def test_sem_esperadas_e_critico_porque_string_morta_nao_acusa():
    u = dd.linha("pv", _r(5, "Usina X", str_esp=None), CTX)
    assert u["pior"] == "critico" and "string morta não acusa" in u["alertas"][0]["texto"]
    assert dd.linha("pv", _r(5, "Usina X", str_esp=0), CTX)["pior"] == "critico"


def test_vale_o_cadastro_nao_o_desconto_da_leitura():
    """1ª prova no PC (11/10, 19h): 20 usinas da Thopen "sem esperadas" que tinham cadastro. Eram String Box à noite: o
    `str_esp` da linha vem com os inversores sem visão descontados e fica 0."""
    noite = {**_r(60061, "Altair 2 (74)", 0, 8, 8), "stringbox": True, "sem_visao": True, "str_esp_cadastro": 152}
    u = dd.linha("pv", noite, CTX)
    assert u["str_esp"] == 152 and u["alertas"] == []
    # cadastro vazio de verdade: crítico; com a régua por inversor vigiando, atenção
    assert dd.linha("pv", {**noite, "str_esp_cadastro": None}, CTX)["pior"] == "critico"
    barretos = {**noite, "str_esp_cadastro": None, "inv_padrao": {"status": "ok"}}
    assert dd.linha("pv", barretos, CTX)["alertas"][0]["texto"].endswith("régua por inversor (padrão 30 dias)")
    # fonte sem o total do cadastro na linha: zero de leitura sem visão ou sem dados não acusa nada
    sem = {"plant_id": "MRO100", "usina": "Morro", "sem_dados": True}
    assert dd.linha("sunop", sem, CTX)["pior"] == "info"
    # o nome que a API escreve vem da própria linha
    assert dd.linha("2capi", {**_r(18771915, "Santa Cecilia 1"), "nome_api": "Ipixuna do Pará 1"},
                    {**CTX, "nome_api": {}})["liga"] == "pelo id (ponte): a API escreve “Ipixuna do Pará 1”"


def test_inversores_na_fonte_diferente_do_cadastro_e_o_provisorio_vem_junto():
    u = dd.linha("2capi", _r(18772125, "União 1 e 2", 108, 6, 12), CTX)
    assert [(a["nivel"], a["texto"]) for a in u["alertas"]] == \
        [("atencao", "12 inversor(es) na fonte, 6 no cadastro"), ("info", "União: o grupo 1 pela ordem")]
    assert u["pior"] == "atencao"
    # só o provisório: aparece na fonte, mas não vai para "Precisa de você"
    assert dd.linha("2capi", _r(18772125, "União 1 e 2", 216, 12, 12), CTX)["pior"] == "info"


def test_nome_de_inversor_pela_ordem_e_aviso():
    u = dd.linha("2capi", _r(99, "Nova da oem"), CTX)
    assert u["nomes_inversor"] == "pela ordem do cadastro" and u["pior"] == "atencao"
    assert dd.linha("pv", _r(7, "Altair 1"), CTX)["nomes_inversor"] == "o nome que a fonte manda"
    s = dd.linha("salenergia", _r("boa", "Boa Esperança do Sul"), CTX)               # scadaGridco sem de-para por valor
    assert s["nomes_inversor"] == "pela ordem do cadastro" and s["pior"] == "atencao"
    assert dd.linha("greenyellow", _r("irece2", "Irecê 2"), CTX)["nomes_inversor"] == "o nome que a fonte manda"


def test_apelido_vale_so_nas_contas_da_api_pv():
    assert dd.linha("2capi", _r(18771901, "Sete Lagoas"), CTX)["liga"] == "por apelido no código"
    assert dd.linha("sunop", _r("STL", "Sete Lagoas"), CTX)["liga"] == "pelo nome"
    assert dd.linha("pv", _r(1, "Altair 1"), CTX)["liga"] == "pelo nome"


def test_o_diagnostico_junta_por_fonte_poe_o_pior_na_frente_e_diz_quem_nao_foi_lido():
    tabelas = {"2capi": [_r(18772125, "União 1 e 2", 108, 6, 12), _r(18771915, "Santa Cecilia 1", 138, 8, 8),
                         _r(18766373, "MORADA NOVA", None, None, 8), {"usina": "sem id"}],
               "pg": []}
    d = dd.diagnostico(tabelas, CTX)
    por = {f["id"]: f for f in d["fontes"]}
    assert list(por) == list(dd.ROTULOS)                         # toda fonte do tempo real aparece, na mesma ordem
    assert por["sunop"]["lida"] is False and por["pg"]["lida"] is True and por["pg"]["usinas"] == []
    assert [u["usina"] for u in por["2capi"]["usinas"]] == ["MORADA NOVA", "União 1 e 2", "Santa Cecilia 1"]
    assert [(u["usina"], u["pior"]) for u in d["precisa"]] == [("MORADA NOVA", "critico"), ("União 1 e 2", "atencao")]
    assert d["contagem"] == {"usinas": 3, "criticos": 1, "atencao": 1, "sem_depara_usina": 0, "sem_depara_inversor": 0,
                             "ok": 1}


def test_o_diagnostico_diz_o_que_a_tela_decidiu_e_o_que_ela_edita():
    ctx = {**CTX, "ponte_tela": {18771915}, "revistas": {"18772125"}, "edita_usina": {"pv", "semp", "alveslima", "2capi"},
           "edita_inversor": {"semp", "alveslima", "2capi", "greenyellow", "salenergia"}}
    u = dd.linha("2capi", _r(18771915, "Santa Cecilia 1", 138, 8, 8), ctx)
    assert u["liga"] == "pelo id (ponte da tela): a API escreve “Ipixuna do Pará 1”"
    assert u["edita_usina"] and u["edita_inversor"]
    assert dd.linha("2capi", _r(18772125, "União 1 e 2", 216, 12, 12), ctx)["nomes_inversor"] == "decidido na tela"
    t = dd.linha("pv", _r(7, "Altair 1"), ctx)                  # a conta principal: a ponte sim, o inversor não
    assert t["edita_usina"] and not t["edita_inversor"]
    assert not any(dd.linha("sunop", _r("MRO100", "Morro"), ctx)[k] for k in ("edita_usina", "edita_inversor"))


# ── as decisões da tela (depara.py) ───────────────────────────────────────────────────────────────────────────────────
def test_a_linha_tem_o_dia_de_brasilia_e_so_onde_o_codigo_liga_pelo_id():
    assert dp.linha_usina("2capi", 18771915, "Santa Cecilia 1", "2026-10-12T02:30:00Z") == \
        ["2026-10-12T02:30:00Z", 20261011, "2capi", "18771915", "", "Santa Cecilia 1", "", "tela"]
    assert dp.linha_inversor("salenergia", "salvales", 3, "Inversor 1.3", Q1)[2:7] == \
        ["salenergia", "salvales", "3", "", "Inversor 1.3"]
    for ruim in (lambda: dp.linha_usina("sunop", "MRO100", "x"), lambda: dp.linha_inversor("pg", 1, 2, "x"),
                 lambda: dp.linha_inversor("2capi", 1, "", "x"), lambda: dp.linha_usina("pv", " ", "x")):
        with pytest.raises(ValueError):
            ruim()


def test_vale_a_ultima_decisao_e_o_vazio_desfaz():
    linhas = [dp.linha_usina("2capi", 18771915, "Santa Cecilia 1", Q1), dp.linha_usina("2capi", 18771915, "", Q2),
              dp.linha_inversor("2capi", 18772125, 401322, "UNI_Inv_1.2", Q1),
              dp.linha_inversor("2capi", 18772125, 401323, "", Q3),
              ["x", 1, "2capi", "1", "", "Lixo", "", "outra_origem"]]                 # origem estranha não conta
    est = dp.vigente(linhas)
    assert dp.pontes(est) == {18771915: ""}
    assert dp.inversores(est, ("2capi",), 18772125) == {"401322": "UNI_Inv_1.2", "401323": ""}
    assert dp.aplica_inversores({401322: "UNI_Inv_1.1", 401323: "UNI_Inv_1.2", 401313: "UNI_Inv_2.1"},
                                dp.inversores(est, ("2capi",), 18772125)) == {401322: "UNI_Inv_1.2", 401313: "UNI_Inv_2.1"}


def test_a_sugestao_casa_pelo_valor_um_para_um_e_diz_a_confianca():
    d1, d2, d3 = "2026-10-08", "2026-10-09", "2026-10-10"
    fonte = {388162: {d1: 1000.0, d2: 900.0, d3: 950.0}, 388159: {d1: 600.0, d2: 610.0, d3: 590.0},
             388158: {d1: 300.0, d2: 310.0}, 402513: {d1: 10.0}}
    bd = {"INVERSOR01": {d1: 1000.4, d2: 899.8, d3: 950.2}, "INVERSOR02": {d1: 630.0, d2: 640.0, d3: 600.0},
          "INVERSOR03": {d1: 360.0, d2: 380.0}, "INVERSOR08": {d1: 0, d2: 0}}
    r = dp.sugerir_por_valor(fonte, bd)
    assert r["pares"]["388162"]["coluna"] == "INVERSOR01" and r["pares"]["388162"]["confianca"] == "certa"
    assert r["pares"]["388159"]["coluna"] == "INVERSOR02" and r["pares"]["388159"]["confianca"] == "provável"
    assert r["pares"]["388158"]["confianca"] == "fraca"                       # 21% de erro: conferir antes
    assert r["sem_par"] == ["402513"] and r["colunas_vazias"] == ["INVERSOR08"]   # 1 dia só; coluna sem valor no BD


# ── a tela e a ligação com o app ──────────────────────────────────────────────────────────────────────────────────────
import app  # noqa: E402
from test_porta_nexus import _sem_rede, cli, entrar, entrar_com_senha  # noqa: E402,F401  (fixtures e passos da porta)


def test_todo_provisorio_e_de_um_de_para_que_existe():
    assert set(app.PV_INV_PROVISORIOS) <= set(app.PV_INV_NOMES)


def test_a_linha_da_api_pv_leva_o_cadastro_antes_dos_descontos_e_o_nome_da_api(monkeypatch):
    monkeypatch.setattr(app, "ESPERADO", {"Santa Cecilia 1": {"inv_esp": 8, "str_esp": 138}})
    planta = {"id": 18771915, "nome": "Santa Cecilia 1", "nome_api": "Ipixuna do Pará 1"}
    assert app._pv_cadastro_da_linha(planta) == {"str_esp_cadastro": 138, "nome_api": "Ipixuna do Pará 1"}
    assert app._pv_cadastro_da_linha({"id": 1, "nome": "Sem Cadastro "}) == {"str_esp_cadastro": None}


def test_a_rota_le_so_o_que_a_plataforma_ja_tem_em_cache(cli, monkeypatch):
    entrar_com_senha(cli)
    monkeypatch.setattr(app, "_2capi_cache", {"payload": {"rows": [_r(18766373, "MORADA NOVA", 128, 8, 8)]}})
    monkeypatch.setattr(app, "_pg_cache", {"summary": [_r(21480, "Indaiatuba", 400, 20, 20)], "detail": {}, "ts": 0.0})
    monkeypatch.setattr(app, "_plants_cache", {"ts": 0.0, "data": [{"id": 18766373, "nome": "MORADA NOVA",
                                                                    "nome_api": "Morada Nova"}]})
    monkeypatch.setattr(app, "_se_cache", {})                                         # não montada: "não lida"
    d = cli.get("/api/admin/de-para").get_json()
    por = {f["id"]: f for f in d["fontes"]}
    mn = por["2capi"]["usinas"][0]
    assert mn["liga"] == "pelo id (ponte): a API escreve “Morada Nova”" and mn["nomes_inversor"] == "de-para no código"
    assert [a["nivel"] for a in mn["alertas"]] == ["info"]                             # o INVERSOR08 por eliminação
    assert mn["edita_usina"] and mn["edita_inversor"]
    assert por["pg"]["usinas"][0]["liga"] == "pelo nome" and por["solaredge"]["lida"] is False
    assert d["rotulos"]["2capi"] == "2C" and set(d["decisoes"]) >= {"ok", "pode_gravar"}


def test_a_tela_e_so_de_administrador_do_nexus(cli):
    entrar(cli, "analista@exemplo.com", destino="/admin/de-para", admin=False)
    assert cli.get("/admin/de-para").status_code == 403
    assert cli.get("/api/admin/de-para").status_code == 403
    for url in ("/api/admin/de-para/usina?fonte=2capi&id=1", "/api/admin/de-para/sugerir?fonte=2capi&id=1"):
        assert cli.get(url).status_code == 403
    assert cli.post("/api/admin/de-para/decidir", json={"fonte": "2capi", "id": "1", "usina_supervisorio": ""}).status_code == 403
    entrar(cli, "admin@exemplo.com", destino="/admin/de-para", admin=True)
    h = cli.get("/admin/de-para").get_data(as_text=True)
    assert "<h1>De-para de equipamentos</h1>" in h and 'href="/" data-casca' in h
    assert cli.get("/api/admin/de-para").status_code == 200


class _R:
    def __init__(self, d, code=200):
        self._d, self.status_code = d, code

    def json(self):
        return self._d

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(self.status_code)


class _Banco:
    """A db_performace de mentira, com o livro do de-para (lê o xlsx enviado de verdade)."""

    def __init__(self):
        self.livros, self.linhas, self.fora, self.recargas = set(), None, False, []

    def get(self, url, headers=None, params=None, timeout=None):
        if self.fora:
            raise ConnectionError("banco fora")
        if url.endswith("/api/workbooks"):
            return _R([{"key": k} for k in self.livros])
        if url.endswith("/api/sheets"):
            return _R([{"id": 9, "workbook_key": dp.LIVRO, "sheet_name": dp.ABA}] if self.linhas is not None else [])
        if url.endswith("/api/sheets/9/rows"):
            off, lim = params["offset"], params["limit"]
            return _R({"rows": [{"headers": dp.CAB, "values": lin} for lin in (self.linhas or [])[off:off + lim]]})
        raise AssertionError(url)

    def post(self, url, headers=None, json=None, params=None, files=None, timeout=None):
        assert headers["Authorization"] == "Bearer tok"
        if url.endswith("/api/workbooks"):
            self.livros.add(json["key"])
            return _R({})
        if url.endswith(f"/api/workbooks/{dp.LIVRO}/sync-xlsx"):
            from io import BytesIO
            from openpyxl import load_workbook
            ws = load_workbook(BytesIO(files["file"][1]))[dp.ABA]
            self.linhas = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
            return _R({"inserted": len(self.linhas)})
        raise AssertionError(url)


@pytest.fixture
def ligado(monkeypatch, tmp_path):
    banco = _Banco()
    monkeypatch.setattr(app, "DEPARA_LIGADO", True)
    monkeypatch.setattr(app, "_DP_DECISOES", app._cob.Decisoes("https://x/db", lambda: app.os.environ.get("GRIDCO_SQL_TOKEN", ""),
                                                               str(tmp_path / "dp.json"), ttl=0, sessao=banco,
                                                               livro=dp.LIVRO, aba=dp.ABA, cab=dp.CAB))
    monkeypatch.setattr(app, "_DP_MEMO", {"linhas": None, "estado": {"usina": {}, "inversor": {}, "apelido": {}}})
    monkeypatch.setattr(app, "_dp_recarregar_cadastro", lambda: banco.recargas.append(1))
    monkeypatch.setenv("GRIDCO_SQL_TOKEN", "tok")
    return banco


def test_a_tela_vence_o_codigo_na_ponte_e_nos_inversores(ligado):
    api = [{"id": 18771915, "nome": "Ipixuna do Pará 1"}, {"id": 999, "nome": "Morada Nova 2"}]
    assert [p["nome"] for p in app._pv_nome_do_cadastro(api)] == ["Santa Cecilia 1", "Morada Nova 2"]     # só o código
    ligado.linhas = [dp.linha_usina("2capi", 18771915, "", Q1), dp.linha_usina("alveslima", 999, "MORADA NOVA", Q1),
                     dp.linha_inversor("2capi", 18772125, 401322, "UNI_Inv_1.2", Q1),
                     dp.linha_inversor("2capi", 18772125, 401323, "", Q1)]
    out = app._pv_nome_do_cadastro(api)
    assert [p["nome"] for p in out] == ["Ipixuna do Pará 1", "MORADA NOVA"] and out[1]["nome_api"] == "Morada Nova 2"
    assert app._pv_nome_2c(18772125, 401322) == "UNI_Inv_1.2"
    assert app._pv_nome_2c(18772125, 401323) is None                                   # "volta ao código" = sem nome
    assert app._pv_nome_2c(18772125, 401313) == "UNI_Inv_2.1"                          # o resto segue o código


def test_livro_ilegivel_vale_o_codigo(ligado):
    ligado.fora = True
    assert app._pv_nome_2c(18772125, 401322) == "UNI_Inv_1.1"
    assert app._pv_nome_do_cadastro([{"id": 18771915, "nome": "Ipixuna do Pará 1"}])[0]["nome"] == "Santa Cecilia 1"


def test_o_scadagridco_pergunta_a_tela_antes_da_ordem(ligado):
    cfg = app._sg.usina_cfg("salenergia", "salvales")
    cad = ["Inversor 1.1", "Inversor 1.2", "Inversor 1.3"]
    assert app._sg.nome_inversor(cfg, "1", cad) == "Inversor 1.1"                     # pela ordem, como antes
    ligado.linhas = [dp.linha_inversor("salenergia", "salvales", "1", "Inversor 1.3", Q1)]
    assert app._sg.nome_inversor(cfg, "1", cad) == "Inversor 1.3"
    assert app._sg.nome_inversor(cfg, "2", cad) == "Inversor 1.2"


def test_decidir_confere_no_cadastro_e_grava(cli, ligado, monkeypatch):
    entrar_com_senha(cli)
    monkeypatch.setattr(app, "ESPERADO", {"Santa Cecilia 1": {"inv_esp": 8, "str_esp": 138}, "União": {"inv_esp": 12}})
    monkeypatch.setattr(app, "ESPERADO_INV", {"União": {"UNI_Inv_1.1": 18, "UNI_Inv_1.2": 18}})
    monkeypatch.setattr(app, "_dp_planta", lambda pid: {"id": pid, "nome": "União"})
    r = cli.post("/api/admin/de-para/decidir", json={"fonte": "2capi", "id": "18771915", "usina_supervisorio": "Não Existe"})
    assert r.status_code == 400 and "não está no cadastro" in r.get_json()["error"]
    r = cli.post("/api/admin/de-para/decidir", json={"fonte": "2capi", "id": "18771915", "usina_supervisorio": "Santa Cecilia 1"})
    assert r.status_code == 200 and r.get_json()["total"] == 1
    r = cli.post("/api/admin/de-para/decidir", json={"fonte": "2capi", "id": "18772125", "inversores": {"401322": "Inventado"}})
    assert r.status_code == 400
    r = cli.post("/api/admin/de-para/decidir", json={"fonte": "2capi", "id": "18772125",
                                                     "inversores": {"401322": "UNI_Inv_1.2", "401323": ""}})
    assert r.status_code == 200 and r.get_json()["total"] == 3
    assert app._pv_nome_2c(18772125, 401322) == "UNI_Inv_1.2"
    assert cli.post("/api/admin/de-para/decidir", json={"fonte": "solaredge", "id": "77",
                                                        "usina_supervisorio": "Santa Cecilia 1"}).status_code == 400


def test_sem_a_chave_do_banco_nao_grava(cli, ligado, monkeypatch):
    entrar_com_senha(cli)
    monkeypatch.setenv("GRIDCO_SQL_TOKEN", "")
    monkeypatch.setattr(app, "ESPERADO", {"Santa Cecilia 1": {"inv_esp": 8}})
    r = cli.post("/api/admin/de-para/decidir", json={"fonte": "2capi", "id": "18771915", "usina_supervisorio": "Santa Cecilia 1"})
    assert r.status_code == 503 and "GRIDCO_SQL_TOKEN" in r.get_json()["error"] and ligado.linhas is None


def test_o_editor_diz_de_onde_vem_cada_nome(cli, ligado, monkeypatch):
    entrar_com_senha(cli)
    monkeypatch.setattr(app, "_dp_planta", lambda pid: {"id": pid, "nome": "União", "nome_api": "União "})
    monkeypatch.setattr(app, "_dp_ids_hoje_api_pv", lambda pid: (list(range(401313, 401319)) + list(range(401322, 401328)), None))
    d = cli.get("/api/admin/de-para/usina?fonte=2capi&id=18772125").get_json()
    assert d["edita_usina"] and d["edita_inversor"] and d["sugerir"] and "pela ordem" in d["provisorio"]
    assert {x["id"]: x["origem"] for x in d["inversores"]}["401322"] == "codigo" and len(d["inversores"]) == 12
    ligado.linhas = [dp.linha_inversor("2capi", 18772125, 401322, "UNI_Inv_1.2", Q1)]
    d = cli.get("/api/admin/de-para/usina?fonte=2capi&id=18772125").get_json()
    x = next(x for x in d["inversores"] if x["id"] == "401322")
    assert (x["origem"], x["nome"], x["codigo"]) == ("tela", "UNI_Inv_1.2", "UNI_Inv_1.1") and d["provisorio"] is None
    assert cli.get("/api/admin/de-para/usina?fonte=solaredge&id=77").status_code == 400


def test_a_sugestao_pelo_valor_le_a_fonte_e_a_aba_da_usina(cli, monkeypatch):
    entrar_com_senha(cli)
    monkeypatch.setattr(app, "_dp_planta", lambda pid: {"id": pid, "nome": "Santa Cecilia 2"})
    monkeypatch.setattr(app, "_pv_token_for", lambda pid: "t")
    monkeypatch.setattr(app, "EQUIP_NAMES", {"Santa Cecilia 2": {"INVERSOR01": "Inversor 2.1", "INVERSOR02": "Inversor 2.2"}})
    monkeypatch.setattr(app, "USINA_GRUPO", {"Santa Cecilia 2": "Ipixuna do Pará"})
    hoje = app.datetime.now().date()
    dias = [(hoje - app.timedelta(days=k)) for k in (1, 2, 3)]
    eday = {400882: [700.0, 710.0, 690.0], 400883: [500.0, 505.0, 495.0]}

    def hist(pid, token, data):                       # o Eday sobe ao longo do dia: vale o maior de cada inversor
        k = [d.strftime("%d/%m/%Y") for d in dias].index(data)
        return [{"idefinversor": i, "conteudojson": {"Eday": v[k] * 0.5}} for i, v in eday.items()] + \
               [{"idefinversor": i, "conteudojson": {"Eday": v[k]}} for i, v in eday.items()], None
    monkeypatch.setattr(app, "_spv_day_records_hist", hist)
    pedidas = []

    def bd(usina, ano, mes):
        pedidas.append(usina)
        if usina != "Ipixuna do Pará":
            return {}
        return {d.isoformat(): {"inv": {"Inversor 2.1": eday[400883][k], "Inversor 2.2": eday[400882][k],
                                        "Inversor 1.1": 999.0}} for k, d in enumerate(dias)}
    monkeypatch.setattr(app, "_bdperf_inv_dias", bd)
    r = cli.get("/api/admin/de-para/sugerir?fonte=2capi&id=18771929").get_json()
    assert pedidas[0] == "Ipixuna do Pará" and r["aba"] == "Ipixuna do Pará"
    assert {i: (p["nome"], p["coluna"], p["confianca"]) for i, p in r["pares"].items()} == \
        {"400882": ("INVERSOR02", "Inversor 2.2", "certa"), "400883": ("INVERSOR01", "Inversor 2.1", "certa")}
    assert cli.get("/api/admin/de-para/sugerir?fonte=pv&id=1").status_code == 400


def test_a_sugestao_le_o_bd_primeiro_e_nao_gasta_cota_sem_dois_dias_no_bd(cli, monkeypatch):
    """No PC o espelho do BD estava parado em 08/10: com os 3 últimos dias fixos, só 1 tinha valor e nada casava."""
    entrar_com_senha(cli)
    monkeypatch.setattr(app, "_dp_planta", lambda pid: {"id": pid, "nome": "UFV Tucano 1"})
    monkeypatch.setattr(app, "_pv_token_for", lambda pid: "t")
    monkeypatch.setattr(app, "EQUIP_NAMES", {})
    um_dia = (app.datetime.now().date() - app.timedelta(days=6)).isoformat()
    monkeypatch.setattr(app, "_bdperf_inv_dias", lambda usina, a, m: {um_dia: {"inv": {"Inversor 1.1": 900.0}}}
                        if usina == "UFV Tucano 1" else {})
    pedidos = []
    monkeypatch.setattr(app, "_spv_day_records_hist", lambda *a: pedidos.append(a) or ([], None))
    r = cli.get("/api/admin/de-para/sugerir?fonte=semp&id=18758732").get_json()
    assert pedidos == [] and r["pares"] == {} and "1 dia(s) com valor" in r["motivos"][0]


def test_texto_vazio_vai_como_celula_vazia_no_xlsx():
    """1ª gravação no servidor (11/10/2026, 20:28): o openpyxl gravou "" como inlineStr sem <is> e a API recusou o arquivo
    inteiro ("Célula F2 é inlineStr sem elemento <is>")."""
    import zipfile
    from io import BytesIO
    b = app._cob.xlsx_bytes([dp.linha_inversor("semp", 18758732, 378310, "", Q1)], dp.ABA, dp.CAB)
    xml = zipfile.ZipFile(BytesIO(b)).read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert xml.count('t="inlineStr"') == xml.count("<is>")
    from openpyxl import load_workbook
    assert [c.value for c in load_workbook(BytesIO(b))[dp.ABA][2]][4:7] == ["378310", None, None]


def test_a_gaveta_do_editor_obedece_ao_hidden():
    """11/10/2026: o display:flex da .gaveta vencia o hidden do navegador — a gaveta aparecia vazia na carga da página e o
    "Fechar" não fechava. Todo elemento com hidden some, venha a regra de display de onde vier."""
    import pathlib
    import re
    html = (pathlib.Path(__file__).resolve().parents[1] / "plataforma" / "templates" / "depara.html").read_text(encoding="utf-8")
    assert re.search(r"\[hidden\]\s*\{\s*display\s*:\s*none\s*!important", html)
    assert 'id="ed" role="dialog" aria-modal="true" aria-labelledby="edTit" hidden' in html and 'id="veu" hidden' in html


# ── 12/10/2026: sem de-para, a sugestão pelo nome, as já ligadas e o apelido das fontes que ligam pelo nome ─────────────
# Levi: "não deveria aparecer usina que já estão ligadas, a não ser que eu tirasse um filtro"; "em NOVO ao invés de manter
# tem que alertar que está sem de-para"; "precisamos trazer para o Nexus e deixar mais dinâmico com a plataforma". No dia,
# 5 usinas da Thopen estavam "sem esperadas" e o problema era outro: a API escreve "Itaja 1 (228)" e "Pharma 2 (644)", e o
# cadastro tem "Itaja 1 (551)" e "Pharma 2".
CAD = [{"chave": "Itaja 1 (551)", "nome": "Itaja 1 (551)"}, {"chave": "Itaja 2 (552)", "nome": "Itaja 2 (552)"},
       {"chave": "Pharma 2", "nome": "Pharma II"}, {"chave": "Ceu Azul 2", "nome": "Céu Azul II"}]


def test_a_sugestao_pelo_nome_exige_o_mesmo_numero_e_avisa_o_parenteses():
    s = dp.parecidas("Itaja 1 (228)", CAD)
    assert [x["chave"] for x in s] == ["Itaja 1 (551)"] and "228 na fonte, 551 no cadastro" in s[0]["confere"]
    assert dp.parecidas("Pharma 2 (644)", CAD) == [{"chave": "Pharma 2", "nota": 1.0, "confere": None}]
    assert [x["chave"] for x in dp.parecidas("UFV Céu Azul II", CAD)] == ["Ceu Azul 2"]
    assert dp.parecidas("Itaja 3 (230)", CAD) == [] and dp.parecidas("", CAD) == []


def test_sem_de_para_e_o_nome_fora_do_cadastro_e_leva_a_sugestao():
    sug = [{"chave": "Itaja 1 (551)", "nota": 1.0, "confere": "x"}]
    ctx = {**CTX, "fora": {("pv", 18749574)}, "sugestoes": {("pv", 18749574): sug}, "edita_usina": {"pv"}}
    u = dd.linha("pv", {**_r(18749574, "Itaja 1 (228)", None, 4, 4), "str_esp_cadastro": None}, ctx)
    assert u["liga"] == "sem de-para" and u["sem_depara"] == "usina" and u["sugestao"] == sug and u["edita_usina"]
    assert u["pior"] == "critico" and u["alertas"][0]["texto"].startswith("sem de-para: o cadastro não tem “Itaja 1 (228)”")
    # no cadastro, mas sem esperadas: o alerta de antes, sem sugestão
    u = dd.linha("pv", {**_r(5, "Usina X", None), "str_esp_cadastro": None}, ctx)
    assert u["sem_depara"] is None and u["sugestao"] == [] and "string morta não acusa" in u["alertas"][0]["texto"]
    assert dd.linha("salenergia", _r("boa", "Boa"), CTX)["sem_depara"] == "inversor"
    # Athon/Axis: o "Editar" só onde há o que decidir (sem de-para, ou o apelido da tela para desfazer)
    c2 = {**CTX, "edita_usina": {"sunop"}, "so_sem_depara": {"sunop"}}
    assert not dd.linha("sunop", _r("CPP100", "CPP100"), c2)["edita_usina"]
    assert dd.linha("sunop", _r("CPP101", "CPP101", None), {**c2, "fora": {("sunop", "CPP101")}})["edita_usina"]
    assert dd.linha("sunop", _r("X", "CPP100"), {**c2, "apelidos_tela": {"X": "CPP100"}})["edita_usina"]


def test_o_apelido_vale_so_nas_fontes_que_ligam_pelo_nome_e_o_vazio_desfaz():
    assert dp.linha_apelido("sunop", " MRO 1 NOVO ", "MRO100", Q1)[2:8] == ["sunop", "MRO 1 NOVO", "", "MRO100", "", "tela"]
    for f in ("pv", "2capi", "solaredge", "pg"):                 # o banco liga pelo id da usina nele, não pelo nome
        with pytest.raises(ValueError):
            dp.linha_apelido(f, "X", "Y", Q1)
    est = dp.vigente([dp.linha_apelido("sunop", "A", "MRO100", Q1), dp.linha_apelido("axis", "B", "Indaiatuba", Q1),
                      dp.linha_apelido("sunop", "A", "", Q2), dp.linha_usina("pv", 7, "Itaja 1 (551)", Q1)])
    assert dp.apelidos(est) == {"B": "Indaiatuba"} and dp.pontes(est) == {7: "Itaja 1 (551)"}
    u = dd.linha("axis", _r("B", "Indaiatuba"), {**CTX, "apelidos_tela": dp.apelidos(est)})
    assert u["liga"] == "por apelido da tela: a fonte escreve “B”"


def test_o_apelido_da_tela_faz_o_cadastro_responder_pelo_nome_da_fonte_sem_sobrescrever(monkeypatch):
    monkeypatch.setattr(app, "ESPERADO", {"MRO100": {"inv_esp": 4, "str_esp": 60}, "Outra": {"inv_esp": 1, "str_esp": 9}})
    monkeypatch.setattr(app, "ESPERADO_INV", {"MRO100": {"INV1": 15}})
    monkeypatch.setattr(app, "EQUIP_NAMES", {})
    monkeypatch.setattr(app, "POWER_INV", {})
    monkeypatch.setattr(app, "USINA_DISPLAY", {"MRO100": "Morro 100"})
    monkeypatch.setattr(app, "FULL_OM", {"MRO100"})
    monkeypatch.setattr(app, "STRING_BOX", set())
    est = dp.vigente([dp.linha_apelido("sunop", "MRO 1 NOVO", "MRO100", Q1), dp.linha_apelido("sunop", "Outra", "MRO100", Q1)])
    monkeypatch.setattr(app, "_DP_MEMO", {"linhas": [], "estado": est})
    app._aplica_alias_api_cadastro()
    assert app.ESPERADO["MRO 1 NOVO"] == {"inv_esp": 4, "str_esp": 60} and app.ESPERADO_INV["MRO 1 NOVO"] == {"INV1": 15}
    assert app.USINA_DISPLAY["MRO 1 NOVO"] == "Morro 100" and "MRO 1 NOVO" in app.FULL_OM
    assert app.ESPERADO["Outra"] == {"inv_esp": 1, "str_esp": 9}                       # o cadastro nunca é sobrescrito


def _tabelas_do_dia(monkeypatch):
    monkeypatch.setattr(app, "ESPERADO", {"Itaja 1 (551)": {"inv_esp": 4, "str_esp": 76},
                                          "Indaiatuba": {"inv_esp": 20, "str_esp": 400},
                                          "Sete Lagoa": {"inv_esp": 6}, "Sete Lagoas": {"inv_esp": 6}})
    monkeypatch.setattr(app, "USINA_DISPLAY", {})
    monkeypatch.setattr(app, "_cache", {"payload": {"rows": [
        {**_r(18749574, "Itaja 1 (228)", None, 4, 4), "str_esp_cadastro": None},
        {**_r(31, "Sete Lagoas", 100), "str_esp_cadastro": 100}]}})
    monkeypatch.setattr(app, "_pg_cache", {"summary": [_r("Indaiatuba", "Indaiatuba", 400, 20, 20)], "detail": {}, "ts": 0.0})
    monkeypatch.setattr(app, "_plants_cache", {"ts": 0.0, "data": [{"id": 18749574, "nome": "Itaja 1 (228)"},
                                                                    {"id": 31, "nome": "Sete Lagoa"}]})
    for c in ("_2capi_cache", "_semp_cache", "_alveslima_cache", "_se_cache"):
        monkeypatch.setattr(app, c, {})
    monkeypatch.setattr(app, "_pv_plant_devices", lambda pid: [])                  # sem API PV no teste
    monkeypatch.setattr(app, "_dp_ids_hoje_api_pv", lambda pid: ([], None))
    # o servidor sem o token do scadaGridco: a linha vem sem leitura, e o nome não casa letra a letra com o cadastro
    monkeypatch.setattr(app, "_SG_CACHE", {"greenyellow": {"payload": {"rows": [
        {"plant_id": "dermeval_lobao", "usina": "Demerval Lobao", "str_esp": None, "sem_dados": True}]}}})


def test_a_rota_diz_sem_de_para_com_a_sugestao_e_o_editor_esconde_o_que_ja_liga(cli, ligado, monkeypatch):
    entrar_com_senha(cli)
    _tabelas_do_dia(monkeypatch)
    d = cli.get("/api/admin/de-para").get_json()
    it = next(u for u in d["precisa"] if u["plant_id"] == 18749574)
    assert it["sem_depara"] == "usina" and it["liga"] == "sem de-para" and it["edita_usina"]
    assert it["sugestao"][0]["chave"] == "Itaja 1 (551)" and it["sugestao"][0]["ligada_a"] == []
    assert d["contagem"]["sem_depara_usina"] == 1
    gy = next(f for f in d["fontes"] if f["id"] == "greenyellow")["usinas"][0]               # liga pelo _sg_cadastro
    assert gy["sem_depara"] is None and gy["pior"] == "info" and "sem leitura" in gy["alertas"][0]["texto"]
    ed = cli.get("/api/admin/de-para/usina?fonte=pv&id=18749574").get_json()
    assert ed["sem_depara"] and ed["sugestao"][0]["chave"] == "Itaja 1 (551)"
    cad = {u["chave"]: u["ligada_a"] for u in ed["cadastro_usinas"]}
    assert cad == {"Itaja 1 (551)": [], "Indaiatuba": ["Thopen · Banco de dados · Indaiatuba"],
                   "Sete Lagoas": ["Thopen · API PV · Sete Lagoas"]}                    # sem a cópia "Sete Lagoa"
    ed = cli.get("/api/admin/de-para/usina?fonte=pv&id=31").get_json()             # a própria usina não conta
    assert not ed["sem_depara"] and {u["chave"]: u["ligada_a"] for u in ed["cadastro_usinas"]}["Sete Lagoas"] == []


def test_o_apelido_das_fontes_que_ligam_pelo_nome_grava_e_rele_o_cadastro(cli, ligado, monkeypatch):
    entrar_com_senha(cli)
    _tabelas_do_dia(monkeypatch)
    ed = cli.get("/api/admin/de-para/usina?fonte=sunop&id=MRO%201%20NOVO").get_json()
    assert ed["por_nome"] and ed["sem_depara"] and ed["edita_usina"] and not ed["edita_inversor"]
    r = cli.post("/api/admin/de-para/decidir", json={"fonte": "sunop", "id": "MRO 1 NOVO", "usina_supervisorio": "Indaiatuba"})
    assert r.status_code == 200 and ligado.recargas == [1]                     # relê o cadastro já, sem esperar 60 s
    assert app._dp_apelidos_da_tela() == {"MRO 1 NOVO": "Indaiatuba"}
    ed = cli.get("/api/admin/de-para/usina?fonte=sunop&id=MRO%201%20NOVO").get_json()
    assert ed["ponte"]["tela"] == "Indaiatuba" and not ed["sem_depara"]
    r = cli.post("/api/admin/de-para/decidir", json={"fonte": "axis", "id": "indaiatuba ", "usina_supervisorio": "Itaja 1 (551)"})
    assert r.status_code == 400 and "liga pelo nome" in r.get_json()["error"]          # o nome já alcança o cadastro
    ed = cli.get("/api/admin/de-para/usina?fonte=sunop&id=Indaiatuba").get_json()      # ligada: o editor não acusa
    assert ed["no_cadastro"] and not ed["sem_depara"] and ed["cadastro_chave"] == "Indaiatuba" and not ed["edita_usina"]
    assert cli.get("/api/admin/de-para/usina?fonte=pg&id=8").status_code == 400       # o banco segue no cadastro
    r = cli.post("/api/admin/de-para/decidir", json={"fonte": "sunop", "id": "MRO 1 NOVO", "usina_supervisorio": ""})
    assert r.status_code == 200 and app._dp_apelidos_da_tela() == {} and ligado.recargas == [1, 1]   # desfazer também relê


# ── 12/10/2026: inversor a inversor na conta principal da API PV ─────────────────────────────────────────────────────
# Levi: "no de-para equipamentos eu consigo ligar usinas mas não consigo ligar inversor a inversor, precisamos ter essa
# visualização". O nome do inversor da conta principal vem do plant_devices e só casava com o cadastro pela grafia.
def test_o_editor_da_conta_principal_mostra_cada_inversor_e_se_ele_casa_com_o_cadastro(cli, ligado, monkeypatch):
    entrar_com_senha(cli)
    monkeypatch.setattr(app, "ESPERADO_INV", {"Itaja 1 (551)": {"Inversor 1.1": 19, "Inversor 1.2": 19}})
    monkeypatch.setattr(app, "EQUIP_NAMES", {})
    monkeypatch.setattr(app, "_dp_planta", lambda pid: {"id": pid, "nome": "Itaja 1 (551)", "nome_api": "Itaja 1 (228)"})
    monkeypatch.setattr(app, "_dp_ids_hoje_api_pv", lambda pid: ([501, 502, 503], None))
    devs = [{"device_id": 501, "device_name": "Inversor 1.1"}, {"device_id": 502, "device_name": "INV-02"},
            {"device_id": 503, "device_name": "INV-03"}, {"device_id": 9, "device_name": "Medidor"}]
    monkeypatch.setattr(app, "_pv_devs_cache", {18749574: {"ts": app.time.time(), "devs": devs}})
    monkeypatch.setattr(app, "_pv_dev_names_cache", {18749574: {"ts": app.time.time(),
                                                                "names": {d["device_id"]: d["device_name"] for d in devs}}})
    d = cli.get("/api/admin/de-para/usina?fonte=pv&id=18749574").get_json()
    por = {x["id"]: x for x in d["inversores"]}
    assert d["edita_inversor"] and set(por) == {"501", "502", "503"}                  # o medidor não é inversor
    assert por["501"]["origem"] == "nome" and por["502"]["origem"] == "sem" and por["502"]["nome_api"] == "INV-02"
    assert {a["nome"] for a in d["alvos"]} == {"Inversor 1.1", "Inversor 1.2"}
    r = cli.post("/api/admin/de-para/decidir", json={"fonte": "pv", "id": "18749574", "inversores": {"502": "Inventado"}})
    assert r.status_code == 400
    r = cli.post("/api/admin/de-para/decidir", json={"fonte": "pv", "id": "18749574", "inversores": {"502": "Inversor 1.2"}})
    assert r.status_code == 200
    # a decisão vale nos dois lugares que dão nome ao inversor da conta principal, e o da API fica guardado
    assert app._pv_dev_names(18749574, "tok")[502] == "Inversor 1.2"
    d2 = {x["device_id"]: x for x in app._pv_plant_devices(18749574)}
    assert d2[502]["device_name"] == "Inversor 1.2" and d2[502]["device_name_api"] == "INV-02"
    assert d2[503]["device_name"] == "INV-03" and devs[1]["device_name"] == "INV-02"   # o cache segue com o da API
    por = {x["id"]: x for x in cli.get("/api/admin/de-para/usina?fonte=pv&id=18749574").get_json()["inversores"]}
    assert por["502"]["origem"] == "tela" and por["502"]["nome"] == "Inversor 1.2" and por["503"]["origem"] == "sem"


def test_sem_a_usina_no_cadastro_o_editor_pede_para_ligar_a_usina_antes(cli, ligado, monkeypatch):
    entrar_com_senha(cli)
    monkeypatch.setattr(app, "ESPERADO_INV", {})
    monkeypatch.setattr(app, "_dp_planta", lambda pid: {"id": pid, "nome": "Itaja 1 (228)"})
    monkeypatch.setattr(app, "_dp_ids_hoje_api_pv", lambda pid: ([501], None))
    monkeypatch.setattr(app, "_pv_devs_cache", {18749574: {"ts": app.time.time(), "devs": [{"device_id": 501, "device_name": "INV-01"}]}})
    d = cli.get("/api/admin/de-para/usina?fonte=pv&id=18749574").get_json()
    assert d["alvos"] == [] and "ligue a usina primeiro" in d["nota_inversor"] and d["inversores"][0]["origem"] == "sem"
