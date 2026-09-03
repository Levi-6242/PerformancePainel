# gemeo/tests/test_modelar_publicar.py
"""Publicacao no workbook da API da Performance: xlsx com cabecalho na linha 1 e vazio de verdade (o sync rejeita
inlineStr vazio), sincronizacao com replace=true criando o workbook so se faltar, e falha que NUNCA derruba o
modelar. O teste de banco (pula sem DSN) confere as sete consultas sobre a MRO100 semeada."""
import datetime as dt
import io
import json
import sys
import types
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
from gemeo.modelar import publicar as pub  # noqa: E402

UTC = dt.timezone.utc
G = Path(__file__).parent / "fixtures" / "golden"
CFG = types.SimpleNamespace(bd_api_base="https://api.teste/db/", bd_api_token="tok", publicar_workbook="gemeo_digital",
                            publicar_dias=90, publicar_ativo=True)
TABS = {"usina": [["MRO100", "sunop", "America/Belem", 6942.0, 5000.0, 25, None, ""]], "equipamento": [], "alias": [],
        "modelo": [["MRO100", "placa", {"gamma": -0.0035}, 0.08, False, True, None, {}]],
        "cascata_dia": [["MRO100", dt.date(2026, 9, 1), "placa", 42606.7, 40417.0, 2189.7, 1697.9, 103.4, 0.2, 388.1, 1.0, 1]],
        "perda_dia": [],
        "evento": [["MRO100", "Inversor parado", "INV_22", dt.datetime(2026, 9, 1, 6, 45), "aberto", 1697.9, "media", {"dia": "2026-09-01"}]]}


def test_xlsx_tem_7_abas_cabecalho_na_linha_1_e_vazio_de_verdade():
    wb = load_workbook(io.BytesIO(pub.xlsx_bytes(TABS)))
    assert wb.sheetnames == list(pub.CABECALHOS)
    ws = wb["usina"]
    assert [c.value for c in ws[1]] == pub.CABECALHOS["usina"] and ws.freeze_panes == "A2"
    assert ws["G2"].value is None and ws["H2"].value is None            # None e "" viram celula vazia
    assert wb["cascata_dia"]["B2"].value == "2026-09-01" and wb["evento"]["D2"].value == "2026-09-01 06:45"
    assert wb["modelo"]["C2"].value == '{"gamma": -0.0035}' and wb["modelo"]["E2"].value == "nao" and wb["modelo"]["F2"].value == "sim"
    assert wb["perda_dia"].max_row == 1                                 # so o cabecalho quando nao ha linhas


class _Resp:
    def __init__(self, status=200, js=None):
        self.status_code, self._js = status, js

    def json(self):
        return self._js

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(f"HTTP {self.status_code}")


class _Sessao:
    def __init__(self, existe=True, falha_sync=False):
        self.chamadas, self.existe, self.falha_sync = [], existe, falha_sync

    def get(self, url, **kw):
        self.chamadas.append(("GET", url, kw))
        return _Resp(200, [{"key": "gemeo_digital"}] if self.existe else [{"key": "bd_performance"}])

    def post(self, url, **kw):
        self.chamadas.append(("POST", url, kw))
        if url.endswith("/sync-xlsx"):
            return _Resp(500 if self.falha_sync else 200, {"sheets": 7, "inserted": 0, "updated": 10, "deleted": 0})
        return _Resp(201, {"id": 36, "key": "gemeo_digital"})


def test_sincronizar_cria_o_workbook_so_se_faltar_e_manda_replace_true():
    s = _Sessao(existe=False)
    assert pub.sincronizar(CFG, b"xlsx", sessao=s)["sheets"] == 7
    assert [(m, u.split("/db")[1]) for m, u, _ in s.chamadas] == [("GET", "/api/workbooks"), ("POST", "/api/workbooks"),
                                                                  ("POST", "/api/workbooks/gemeo_digital/sync-xlsx")]
    sync = s.chamadas[-1][2]
    assert sync["params"] == {"replace": "true"} and sync["files"]["file"][1] == b"xlsx"
    assert sync["headers"]["Authorization"] == "Bearer tok" and sync["files"]["file"][2] == pub.MIME_XLSX
    s2 = _Sessao(existe=True)
    pub.sincronizar(CFG, b"x", sessao=s2)
    assert [m for m, _, _ in s2.chamadas] == ["GET", "POST"]            # ja existe: nao cria de novo
    with pytest.raises(RuntimeError):
        pub.sincronizar(CFG, b"x", sessao=_Sessao(falha_sync=True))


def test_publicar_registra_o_estado_e_nunca_derruba_o_modelar(monkeypatch):
    gravado = {}
    monkeypatch.setattr(pub, "tabelas", lambda conn, dias: TABS)
    monkeypatch.setattr(pub.db, "gravar_estado", lambda conn, chave, valor: gravado.update({chave: json.loads(valor)}))

    def quebra(cfg, conteudo, sessao=None):
        raise RuntimeError("HTTP 500")

    monkeypatch.setattr(pub, "sincronizar", quebra)
    conn = types.SimpleNamespace(rollback=lambda: None)
    res = pub.publicar(conn, CFG, agora=dt.datetime(2026, 9, 3, 15, 0, tzinfo=UTC))
    assert res["ok"] is False and "HTTP 500" in res["erro"] and gravado["publicar.ultimo"]["ok"] is False
    monkeypatch.setattr(pub, "sincronizar", lambda cfg, conteudo, sessao=None: {"sheets": 7, "updated": 3})
    res = pub.publicar(conn, CFG)
    assert res["ok"] is True and res["linhas"]["cascata_dia"] == 1 and gravado["publicar.ultimo"]["api"]["sheets"] == 7


def test_tabelas_do_banco_semeado(conn):
    from semear import limpar_tudo, semear_fixture
    from gemeo.modelar import job
    limpar_tudo(conn)
    trk_inv = json.load(open(G / "mro100_trk_inv.json", encoding="utf-8"))
    usina, ids = semear_fixture(conn, G / "mro100_2026-08-31.json", trk_inv)
    try:
        job.modelar(conn, usina, dt.datetime(2026, 8, 31, 3, tzinfo=UTC), dt.datetime(2026, 9, 1, 3, tzinfo=UTC))
        t = pub.tabelas(conn, dias=3650)
        assert set(t) == set(pub.CABECALHOS)
        assert len(t["usina"]) == 1 and len(t["equipamento"]) == 182 and len(t["modelo"]) == 1 and len(t["cascata_dia"]) == 1
        assert len(t["perda_dia"]) >= 20 and len(t["evento"]) >= 3
        assert t["cascata_dia"][0][0] == "MRO100" and t["cascata_dia"][0][6] > 1000                   # inv_parado do INV_22
        assert any(e[1] == "Inversor parado" and e[2] == "INV_22" for e in t["evento"])
        assert all(len(l) == len(pub.CABECALHOS[k]) for k, ls in t.items() for l in ls)
        wb = load_workbook(io.BytesIO(pub.xlsx_bytes(t)))
        assert wb["cascata_dia"].max_row == 2 and wb["equipamento"].max_row == 183
        assert pub.tabelas(conn, dias=0)["cascata_dia"] == []                                        # janela: so o que cabe nos dias
    finally:
        limpar_tudo(conn)
