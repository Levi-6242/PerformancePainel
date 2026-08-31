# -*- coding: utf-8 -*-
"""Cópia de segurança, no PostgreSQL, do estado que ninguém reconstrói.

POR QUE EXISTE
Medido em 30/08/2026: dos 52 MB de estado em JSON da plataforma, **92% não se reconstrói** —
não é cache. São eventos de tracker, ocorrências com histerese, perdas de string por dia, o
acervo de e-mails da 2C (que já foram consumidos), strings trancadas e comentários de analista.
Hoje isso vive só em disco, na máquina do analista. Num servidor com disco efêmero, um
re-deploy apagaria tudo — e nada traria de volta.

POR QUE CÓPIA, E NÃO ESCRITA DIRETA NO BANCO
O caminho óbvio seria gravar no banco a cada ação. Seria errado: trancar uma string é ação
FREQUENTE (são 8.590 hoje) e HOJE é instantânea, porque escreve num arquivo local. Trocar isso
por uma chamada de rede põe latência e uma chance de falha no meio de um clique do usuário —
para resolver um problema (re-deploy) que não acontece no clique.

Então: o arquivo local continua sendo onde se ESCREVE, e este módulo publica no banco de tempos
em tempos. No boot, se o arquivo não existir (servidor novo, disco efêmero), restaura do banco.
É o mesmo princípio do `cache_snapshot.json`, só que para o que não pode ser perdido.

COMO O DADO VAI
Um workbook próprio (`plataforma_estado`), uma aba por estrutura, LINHA POR REGISTRO — não um
JSON enfiado numa célula. `strings_trancadas` vira 8.590 linhas de uma coluna; `comments_thread`
vira uma linha por comentário com usina/autor/data/texto. Assim o dado fica legível e
consultável no banco, em vez de um blob opaco.

O envio usa `sync-xlsx`, o mesmo caminho já provado do `sync_gridco_api.py`: monta o .xlsx em
memória e sobe de uma vez, em vez de milhares de POSTs de linha.

O QUE NÃO ENTRA
`whats_ronda.json` — tem `token` e `service_url`, que são SEGREDO. Configuração com credencial
não vai para um banco de planilhas; ela pertence ao volume persistente, junto com o
`tokens_runtime.json`.
"""
import datetime as dt
import io
import json
import os
import threading
import time
import urllib.error
import urllib.request

_AQUI = os.path.dirname(os.path.abspath(__file__))
API = os.environ.get("GRIDCO_DB_API", "https://app.gridco.com.br/db_performace").rstrip("/")
WORKBOOK = "plataforma_estado"
_lock = threading.Lock()


def _token():
    v = os.environ.get("GRIDCO_SQL_TOKEN")
    if v:
        return v.strip()
    import re
    for cand in (os.path.join(_AQUI, "..", "tokens.txt"), os.path.join(_AQUI, "tokens.txt")):
        try:
            for l in io.open(cand, encoding="utf-8", errors="surrogateescape"):
                m = re.match(r"\s*GRIDCO_SQL_TOKEN\s*=\s*(.+)", l)
                if m:
                    return m.group(1).strip()
        except OSError:
            continue
    return ""


def _req(caminho, dados=None, metodo=None, tipo="application/json", timeout=180):
    """`timeout` generoso no ENVIO: 206 mil linhas do grupo `series` levam bem mais que os 180 s
    das leituras, e um timeout do CLIENTE no meio de um sync deixa o workbook vazio — foi o que
    aconteceu na 1ª tentativa. O sync das bases usa 1800 s pelo mesmo motivo."""
    h = {"Content-Type": tipo}
    tk = _token()
    if tk:
        h["Authorization"] = "Bearer " + tk
    r = urllib.request.Request(API + caminho, data=dados, headers=h, method=metodo)
    with urllib.request.urlopen(r, timeout=timeout) as resp:
        b = resp.read()
        try:
            return json.loads(b.decode("utf-8"))
        except ValueError:
            return {"_raw": b[:200]}


# ── O que é copiado, e como cada coisa vira linha ─────────────────────────────
# Cada entrada: (aba, cabeçalho, extrator). O extrator recebe o JSON do arquivo e devolve
# uma lista de linhas. Manter isto declarativo é o que permite acrescentar uma estrutura nova
# sem mexer na mecânica de envio.
def _linhas_lista(v):
    return [[x] for x in (v or [])]


def _linhas_dict(v):
    return [[k, json.dumps(val, ensure_ascii=False)] for k, val in (v or {}).items()]


def _linhas_threads(v):
    """comments_thread: {chave: [comentário, ...]} → uma linha POR COMENTÁRIO.

    Os campos são `nick`, `texto` e `ts` — conferidos nos 112 comentários reais, não supostos.
    Na primeira versão eu procurei "autor", que não existe, e a coluna saiu vazia: a cópia
    teria o texto de todo mundo sem saber de quem é cada um. Os `.get` alternativos ficam como
    rede se o formato mudar, mas o nome certo vem primeiro."""
    out = []
    for chave, lista in (v or {}).items():
        for c in (lista or []):
            if isinstance(c, dict):
                out.append([chave,
                            c.get("nick") or c.get("autor") or c.get("user") or "",
                            c.get("ts") or c.get("data") or "",
                            c.get("texto") or c.get("text") or ""])
            else:
                out.append([chave, "", "", str(c)])
    return out


def _linhas_trk_eventos(d):
    """{dia: {plant_id: {nome, ts, cobertura, eventos:[{tracker,parada,retorno,dur_min}]}}}
    → uma linha por EVENTO, com o dia e a usina desnormalizados na frente.

    A cobertura fica na aba irmã (`trk_cobertura`): ela é do usina-dia, não do evento, e
    repeti-la em 122 mil linhas só engordaria o arquivo."""
    out = []
    for dia, usinas in (d or {}).items():
        for pid, u in (usinas or {}).items():
            for e in (u.get("eventos") or []):
                out.append([dia, pid, u.get("nome") or "", e.get("tracker") or "",
                            e.get("parada") or "", e.get("retorno") or "", e.get("dur_min")])
    return out


def _linhas_trk_cobertura(d):
    return [[dia, pid, u.get("nome") or "", u.get("ts"), u.get("cobertura"),
             len(u.get("eventos") or [])]
            for dia, usinas in (d or {}).items() for pid, u in (usinas or {}).items()]


def _linhas_perdas(d):
    """{dia: {fonte: {plant_id: {usina, ts, eventos:[...]}}}} → uma linha por evento."""
    out = []
    for dia, fontes in (d or {}).items():
        for fonte, usinas in (fontes or {}).items():
            if not isinstance(usinas, dict):
                continue
            for pid, u in usinas.items():
                if not isinstance(u, dict):
                    continue
                for e in (u.get("eventos") or []):
                    out.append([dia, fonte, pid, u.get("usina") or "",
                                json.dumps(e, ensure_ascii=False)])
    return out


def _linhas_perdas_dias(d):
    """A ESTRUTURA de dia/fonte/usina, com ou sem evento.

    Sem isto, um dia em que nenhuma usina teve perda desaparece na volta: a reconstrução só
    cria dia a partir de evento. Pego no teste de ida e volta — 62 dias iam, 61 voltavam, com
    os 37.270 eventos completos. "Dia sem evento" e "dia não processado" são coisas diferentes,
    e confundi-las faria a plataforma reprocessar um dia que já estava fechado."""
    out = []
    for dia, fontes in (d or {}).items():
        for fonte, usinas in (fontes or {}).items():
            if not isinstance(usinas, dict):
                out.append([dia, fonte, "", "", None])
                continue
            if not usinas:
                out.append([dia, fonte, "", "", None])
            for pid, u in usinas.items():
                if isinstance(u, dict):
                    out.append([dia, fonte, pid, u.get("usina") or "", u.get("ts")])
    return out


def _linhas_lista_json(v):
    """Lista de dicts heterogêneos → uma linha com o JSON. Usado onde os campos variam e
    fixar colunas perderia informação silenciosamente (history do tracker_issues)."""
    return [[json.dumps(x, ensure_ascii=False)] for x in (v or [])]


def _linhas_paradas(d):
    """{fonte: {ts, ini, fim, rows:[...], dias_classificados}} → uma linha por row."""
    out = []
    for fonte, blk in (d or {}).items():
        if not isinstance(blk, dict):
            continue
        for r in (blk.get("rows") or []):
            out.append([fonte, blk.get("ini") or "", blk.get("fim") or "",
                        json.dumps(r, ensure_ascii=False)])
    return out


def _linhas_paradas_fontes(d):
    """Os campos do BLOCO de cada fonte — tudo menos as `rows`.

    Duas perdas silenciosas de uma vez, achadas no teste de ida e volta:
    (1) fonte com ZERO rows sumia, porque a reconstrução criava fonte a partir de row;
    (2) `ts` e `dias_classificados` nunca chegavam à cópia — eu só levava ini/fim/rows.
    `dias_classificados` é o que diz quais dias já foram classificados; perdê-lo faria a
    plataforma reclassificar tudo."""
    return [[fonte, blk.get("ts"), blk.get("ini") or "", blk.get("fim") or "",
             json.dumps(blk.get("dias_classificados"), ensure_ascii=False)]
            for fonte, blk in (d or {}).items() if isinstance(blk, dict)]


ESTRUTURAS = [
    # (arquivo,           chave dentro do json,  aba,                cabeçalho,                    extrator)
    ("ufv_state.json", "strings_trancadas", "strings_trancadas", ["chave"], _linhas_lista),
    ("ufv_state.json", "verified", "verified", ["chave"], _linhas_lista),
    ("ufv_state.json", "manutencao", "manutencao", ["chave"], _linhas_lista),
    ("ufv_state.json", "tracking", "tracking", ["chave", "valor"], _linhas_dict),
    ("ufv_state.json", "etm_tickets", "etm_tickets", ["chave", "valor"], _linhas_dict),
    ("ufv_state.json", "os_atribuidas", "os_atribuidas", ["chave", "valor"], _linhas_dict),
    ("ufv_state.json", "comments_thread", "comentarios", ["chave", "autor", "quando", "texto"],
     _linhas_threads),
    ("ufv_state.json", "comments", "comments", ["chave", "valor"], _linhas_dict),
    ("trackers_notas.json", None, "trackers_notas", ["chave", "valor"], _linhas_dict),
    ("trackers_garantia.json", None, "trackers_garantia", ["chave", "valor"], _linhas_dict),
    ("trackers_notas.local.json", None, "trackers_notas_local", ["chave", "valor"], _linhas_dict),
    ("trackers_garantia.local.json", None, "trackers_garantia_local", ["chave", "valor"],
     _linhas_dict),
    ("string_notas.json", None, "string_notas", ["chave", "valor"], _linhas_dict),
]

# ── Grupo 2: a SÉRIE TEMPORAL (etapa 4b) ──────────────────────────────────────
# Workbook e cadência próprios porque a natureza é outra: são ~200 mil linhas que crescem um
# dia por dia, contra os 8,9 mil do estado, que muda a cada clique. Publicar isto de hora em
# hora seria reescrever 200 mil linhas para acrescentar algumas dezenas.
SERIES = [
    ("trk_eventos.json", None, "trk_eventos",
     ["dia", "plant_id", "usina", "tracker", "parada", "retorno", "dur_min"], _linhas_trk_eventos),
    ("trk_eventos.json", None, "trk_cobertura",
     ["dia", "plant_id", "usina", "ts", "cobertura", "n_eventos"], _linhas_trk_cobertura),
    ("tracker_issues.json", "history", "issues_history", ["json"], _linhas_lista_json),
    ("tracker_issues.json", "active", "issues_active", ["chave", "valor"], _linhas_dict),
    ("perdas_strings.json", None, "perdas_strings",
     ["dia", "fonte", "plant_id", "usina", "evento"], _linhas_perdas),
    ("perdas_strings.json", None, "perdas_dias",
     ["dia", "fonte", "plant_id", "usina", "ts"], _linhas_perdas_dias),
    ("paradas_book.json", None, "paradas_book", ["fonte", "ini", "fim", "row"], _linhas_paradas),
    ("paradas_book.json", None, "paradas_fontes",
     ["fonte", "ts", "ini", "fim", "dias_classificados"], _linhas_paradas_fontes),
]

GRUPOS = {
    "estado": (ESTRUTURAS, "plataforma_estado", "Estado da Plataforma"),
    "series": (SERIES, "plataforma_series", "Series da Plataforma"),
}


def _dados_dir():
    return os.environ.get("GRIDCO_DADOS_DIR") or _AQUI


def _ler(arquivo):
    try:
        with io.open(os.path.join(_dados_dir(), arquivo), encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return None


def montar_xlsx(grupo: str = "estado") -> tuple:
    """(bytes do .xlsx, resumo {aba: n_linhas}). Uma aba por estrutura do grupo."""
    import openpyxl
    estruturas = GRUPOS[grupo][0]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    resumo = {}
    for arquivo, chave, aba, cab, extrator in estruturas:
        d = _ler(arquivo)
        if d is None:
            continue
        v = d.get(chave) if chave else d
        if chave is None and isinstance(v, dict):
            # metadados de documentação do próprio arquivo não são dado
            v = {k: x for k, x in v.items() if not k.startswith("_")}
        linhas = extrator(v)
        ws = wb.create_sheet(title=aba[:31])
        ws.append(cab)
        for ln in linhas:
            # String VAZIA vira None. O openpyxl grava "" como uma célula inlineStr sem o
            # elemento <is>, e o importador da API recusa o arquivo inteiro com
            # "Célula B2 é inlineStr sem elemento <is>" — um HTTP 500 por causa de um autor
            # de comentário em branco. Célula vazia de verdade é o que se quer dizer aqui.
            ws.append([None if x == "" else x for x in ln])
        resumo[aba] = len(linhas)
    # carimbo: sem isto não dá para saber, olhando o banco, de quando é a cópia
    ws = wb.create_sheet(title="_meta")
    ws.append(["chave", "valor"])
    ws.append(["gerado_em", dt.datetime.now().isoformat(timespec="seconds")])
    ws.append(["origem", os.environ.get("COMPUTERNAME") or "?"])
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue(), resumo


def publicar(grupo: str = "estado") -> dict:
    """Sobe a cópia do grupo. Idempotente: `replace=true` troca o conteúdo inteiro."""
    with _lock:
        chave_wb, nome_wb = GRUPOS[grupo][1], GRUPOS[grupo][2]
        dados, resumo = montar_xlsx(grupo)
        # PERGUNTAR se existe, em vez de tentar criar e engolir o erro: a API responde 500
        # (não 409) quando a chave já existe, e adivinhar código de erro é frágil — o primeiro
        # envio real morreu exatamente aí, num 500 que parecia falha do sync e era só
        # "já criei isso antes".
        existentes = {w.get("key") for w in (_req("/api/workbooks") or [])}
        if chave_wb not in existentes:
            _req("/api/workbooks", json.dumps({"key": chave_wb,
                                               "display_name": nome_wb}).encode(), "POST")
        corpo, sep = [], b"--gridco"
        cab = (b'--gridco\r\nContent-Disposition: form-data; name="file"; '
               b'filename="estado.xlsx"\r\nContent-Type: '
               b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n")
        corpo = cab + dados + b"\r\n--gridco--\r\n"
        r = _req(f"/api/workbooks/{chave_wb}/sync-xlsx?replace=true", corpo, "POST",
                 "multipart/form-data; boundary=gridco", timeout=1800)
        return {"grupo": grupo, "enviado": r, "abas": resumo,
                "KB": round(len(dados) / 1024, 1)}


# ── A VOLTA: reconstruir o arquivo a partir das linhas ────────────────────────
# Cada extrator do grupo `series` achata de um jeito próprio, então cada um precisa do seu
# inverso. Sem isto a cópia serviria só para consulta — e o motivo dela existir é justamente
# poder trazer de volta.
def _rec_trk(linhas_ev, linhas_cob):
    """trk_eventos + trk_cobertura → {dia: {plant: {nome, ts, cobertura, eventos:[]}}}"""
    out = {}
    for x in linhas_cob:
        if len(x) < 5 or x[0] is None:
            continue
        out.setdefault(str(x[0]), {})[str(x[1])] = {
            "nome": x[2] or "", "ts": x[3], "cobertura": x[4], "eventos": []}
    for x in linhas_ev:
        if len(x) < 7 or x[0] is None:
            continue
        u = out.setdefault(str(x[0]), {}).setdefault(str(x[1]),
                                                     {"nome": x[2] or "", "ts": None,
                                                      "cobertura": None, "eventos": []})
        u["eventos"].append({"tracker": x[3], "parada": x[4], "retorno": x[5],
                             "dur_min": x[6]})
    return out


def _rec_perdas(linhas, estrutura=()):
    """perdas_strings → {dia: {fonte: {plant: {usina, ts, eventos:[]}}}}.

    A ESTRUTURA vem primeiro (aba `perdas_dias`) e os eventos depois: é o que preserva o dia
    em que ninguém teve perda."""
    out = {}
    for x in estrutura:
        if len(x) < 4 or x[0] is None:
            continue
        f = out.setdefault(str(x[0]), {}).setdefault(str(x[1]), {})
        if x[2]:
            f.setdefault(str(x[2]), {"usina": x[3] or "", "ts": x[4] if len(x) > 4 else None,
                                     "eventos": []})
    for x in linhas:
        if len(x) < 5 or x[0] is None:
            continue
        u = out.setdefault(str(x[0]), {}).setdefault(str(x[1]), {}).setdefault(
            str(x[2]), {"usina": x[3] or "", "ts": None, "eventos": []})
        try:
            u["eventos"].append(json.loads(x[4]))
        except (ValueError, TypeError):
            pass
    return out


def _rec_paradas(linhas, fontes=()):
    """A estrutura das fontes vem primeiro; as rows entram depois."""
    out = {}
    for x in fontes:
        if not x or x[0] is None:
            continue
        try:
            dias = json.loads(x[4]) if len(x) > 4 and isinstance(x[4], str) else None
        except ValueError:
            dias = None
        out[str(x[0])] = {"ts": x[1] if len(x) > 1 else None,
                          "ini": (x[2] if len(x) > 2 else "") or "",
                          "fim": (x[3] if len(x) > 3 else "") or "",
                          "rows": [], "dias_classificados": dias}
    for x in linhas:
        if len(x) < 4 or x[0] is None:
            continue
        b = out.setdefault(str(x[0]), {"ini": x[1], "fim": x[2], "rows": []})
        try:
            b["rows"].append(json.loads(x[3]))
        except (ValueError, TypeError):
            pass
    return out


def _linhas_do_banco(sheet_id):
    out, off = [], 0
    while True:
        d = _req(f"/api/sheets/{sheet_id}/rows?offset={off}&limit=1000")
        got = d.get("rows") or []
        out.extend(got)
        if len(got) < 1000:
            return out
        off += 1000


def restaurar_series(destino=None, escrever=False) -> dict:
    """Remonta os arquivos do grupo `series`. Mesma regra da `restaurar`: nunca sobrescreve."""
    abas = {x["sheet_name"]: x for x in (_req("/api/sheets") or [])
            if x.get("workbook_key") == GRUPOS["series"][1]}
    if not abas:
        return {"_erro": "sem cópia de séries no banco"}

    def vals(nome):
        s_ = abas.get(nome)
        return [r.get("values") or [] for r in _linhas_do_banco(s_["id"])] if s_ else []

    montado = {
        "trk_eventos.json": _rec_trk(vals("trk_eventos"), vals("trk_cobertura")),
        "perdas_strings.json": _rec_perdas(vals("perdas_strings"), vals("perdas_dias")),
        "paradas_book.json": _rec_paradas(vals("paradas_book"), vals("paradas_fontes")),
    }
    hist = []
    for x in vals("issues_history"):
        try:
            hist.append(json.loads(x[0]))
        except (ValueError, TypeError, IndexError):
            pass
    ativos = {}
    for x in vals("issues_active"):
        if len(x) >= 2 and x[0] is not None:
            try:
                ativos[x[0]] = json.loads(x[1])
            except (ValueError, TypeError):
                ativos[x[0]] = x[1]
    montado["tracker_issues.json"] = {"active": ativos, "history": hist,
                                      "ultima_atualizacao": dt.datetime.now().isoformat(
                                          timespec="seconds")}
    if escrever:
        base = destino or _dados_dir()
        os.makedirs(base, exist_ok=True)
        for arquivo, conteudo in montado.items():
            alvo = os.path.join(base, arquivo)
            if os.path.exists(alvo):
                print(f"[estado_backup] {arquivo} JA EXISTE — não sobrescrevo")
                continue
            tmp = alvo + ".tmp"
            with io.open(tmp, "w", encoding="utf-8") as f:
                json.dump(conteudo, f, ensure_ascii=False)
            os.replace(tmp, alvo)
            print(f"[estado_backup] {arquivo} restaurado do banco")
    return montado


def restaurar(destino=None, escrever=False) -> dict:
    """Remonta os arquivos a partir da cópia no banco. Devolve {arquivo: conteúdo}.

    `escrever=False` por padrão — e isso é de propósito. Restauração sobrescreve o trabalho de
    gente; ela existe para o servidor NOVO, cujo disco está vazio, não para "sincronizar" por
    cima de um estado vivo. Quem chama com `escrever=True` assume que o destino está vazio.
    """
    abas = {s["sheet_name"]: s for s in (_req("/api/sheets") or [])
            if s.get("workbook_key") == WORKBOOK}
    if not abas:
        return {"_erro": "sem cópia no banco"}

    montado = {}
    for arquivo, chave, aba, cab, extrator in ESTRUTURAS:
        s = abas.get(aba)
        if not s:
            continue
        vals = [r.get("values") or [] for r in _linhas_do_banco(s["id"])]
        if extrator is _linhas_lista:
            v = [x[0] for x in vals if x and x[0] is not None]
        elif extrator is _linhas_dict:
            v = {}
            for x in vals:
                if len(x) >= 2 and x[0] is not None:
                    try:
                        v[x[0]] = json.loads(x[1]) if isinstance(x[1], str) else x[1]
                    except ValueError:
                        v[x[0]] = x[1]
        else:                                     # _linhas_threads
            v = {}
            for x in vals:
                if not x or x[0] is None:
                    continue
                v.setdefault(x[0], []).append({
                    "nick": x[1] if len(x) > 1 and x[1] is not None else "",
                    "ts": x[2] if len(x) > 2 and x[2] is not None else "",
                    "texto": x[3] if len(x) > 3 and x[3] is not None else ""})
        alvo = montado.setdefault(arquivo, {})
        if chave:
            alvo[chave] = v
        else:
            alvo.update(v if isinstance(v, dict) else {})

    if escrever:
        base = destino or _dados_dir()
        os.makedirs(base, exist_ok=True)
        for arquivo, conteudo in montado.items():
            p = os.path.join(base, arquivo)
            if os.path.exists(p):
                print(f"[estado_backup] {arquivo} JA EXISTE — não sobrescrevo")
                continue
            tmp = p + ".tmp"
            with io.open(tmp, "w", encoding="utf-8") as f:
                json.dump(conteudo, f, ensure_ascii=False, indent=1)
            os.replace(tmp, p)
            print(f"[estado_backup] {arquivo} restaurado do banco")
    return montado


#`ufv_state.json` é o SINAL de que existe estado nesta máquina: é ele que guarda strings
# trancadas, comentários e verificações, e é criado no primeiro uso. Os demais podem faltar
# legitimamente — `string_notas.json`, por exemplo, nunca foi criado em 3 meses de operação.
# Disparar por "falta algum arquivo" fazia a restauração rodar em TODO boot por causa dele.
_SINAL = "ufv_state.json"


def restaurar_se_vazio() -> bool:
    """Boot de servidor novo: se não há estado nenhum em disco, traz do banco. True se trouxe.
    Nunca sobrescreve — arquivo presente é sempre mais confiável que a cópia."""
    if os.path.exists(os.path.join(_dados_dir(), _SINAL)):
        return False
    print(f"[estado_backup] disco sem {_SINAL} — restaurando estado e séries do banco")
    restaurar(escrever=True)
    try:
        restaurar_series(escrever=True)
    except Exception as e:            # noqa: BLE001 - sem série ainda é melhor que sem subir
        print(f"[estado_backup] séries não vieram: {e}")
    return True


def estado() -> dict:
    dados, resumo = montar_xlsx()
    return {"abas": resumo, "KB": round(len(dados) / 1024, 1),
            "total_linhas": sum(resumo.values())}


if __name__ == "__main__":
    import sys
    if "--publicar" in sys.argv:
        print(json.dumps(publicar(), ensure_ascii=False, indent=1))
    else:
        print(json.dumps(estado(), ensure_ascii=False, indent=1))
        print("\n(ensaio — use --publicar para enviar)")
