"""Gera o espelho local das planilhas a partir da Gridco Performance API.

POR QUE EXISTE
O objetivo é tirar a plataforma do OneDrive, para ela poder rodar num servidor com link fixo.
Hoje as bases vêm de arquivos sincronizados na máquina do analista — o que impede hospedar, trava
o arquivo (`Errno 13` quando o Excel está aberto) e deixa a coleta à mercê do sync sobrescrever o
que acabou de ser gravado.

POR QUE ESPELHO, E NÃO TROCAR AS LEITURAS
O `app.py` lê essas planilhas em ~20 lugares, cada um com seu cabeçalho (`header=0/1/2/3`), uns por
pandas e outros por openpyxl. Reescrever os 20 seria arriscado e não traria nada: o que precisamos
é trocar a ORIGEM do arquivo, não a forma de ler. Então baixamos as linhas da API e
MATERIALIZAMOS um .xlsx equivalente em `plataforma/bases/` — que é o diretório de espelho que o
`_base_espelho()` já consulta quando não há OneDrive. Nenhum ponto de leitura muda.

A REGRA QUE FAZ ISSO FUNCIONAR
Cada linha é escrita na sua posição ORIGINAL (`row_number` que a API devolve). É isso que mantém
`header=2` apontando para a linha certa sem ajuste nenhum — a aba Equipamentos, por exemplo,
começa na linha 3, e escrever "da linha 1 em diante" quebraria todos os leitores em silêncio.

NOTA SOBRE TABELAS NOMEADAS: a API não as expõe, e o espelho tampouco as recria. Os leitores
foram adaptados (25/08, opção 2 do Levi): o `dashboard_thopen` sintetiza a "tabela" da própria
aba quando não há tabela nomeada e resolve os 3 nomes fixos por assinatura de colunas; o
`load_thopen_meta` do app.py idem. Verificado contra o OneDrive pelos leitores REAIS:
carteira/registro/meta iguais; daily igual salvo 1 célula em que a API estava MAIS completa
(cache de fórmula não calculado no arquivo local).
"""
import datetime as _dt
import io
import json
import os
import re
import threading
import time
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor

# A API serializa datas como TEXTO ("2026-02-01"). Gravar isso como texto no espelho parece
# inofensivo e NÃO É: o `load_metas` lê a coluna Mês com `pd.to_datetime(..., dayfirst=True)`, e
# com dayfirst o texto "2026-02-01" vira DIA 2 do mês 1. Medido em 25/08: as 11 linhas mensais de
# cada usina colapsavam todas para (2026, 1) e o gerencial ficaria com uma meta só, repetida —
# sem erro nenhum na tela. Então o espelho tem de restaurar o TIPO, não só o conteúdo.
_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})"
                  r"(?:[T ](\d{2}):(\d{2}):(\d{2})(?:\.\d+)?)?(?:Z|[+-]\d{2}:?\d{2})?$")


def _valor(v):
    """Converte o que a API manda de volta ao tipo do Excel. Hoje só datas — o resto (números,
    texto, None) já vem no tipo certo pelo JSON."""
    if not isinstance(v, str):
        return v
    m = _ISO.match(v.strip())
    if not m:
        return v
    a, mes, d, hh, mm, ss = m.groups()
    try:
        if hh is None:
            return _dt.date(int(a), int(mes), int(d))
        return _dt.datetime(int(a), int(mes), int(d), int(hh), int(mm), int(ss))
    except ValueError:
        # Parece data mas não é (mês 13, dia 32, texto de código). Devolve como veio: célula
        # esquisita não pode derrubar o espelho inteiro — o valor original ainda é a verdade.
        return v

_AQUI = os.path.dirname(os.path.abspath(__file__))
_BASES_DIR = os.path.join(_AQUI, "bases")
_MARCA = os.path.join(_BASES_DIR, "_bd_api_sync.json")   # o que já foi gerado, e de qual versão

API_BASE = os.environ.get("GRIDCO_DB_API", "https://app.gridco.com.br/db_performace").rstrip("/")

# 500 é o teto do endpoint: pedir 2000 devolve ZERO linhas (medido 25/08) em vez de erro — um
# limite maior faria o espelho nascer VAZIO sem nenhum aviso.
_PAGINA = 500
# 12 threads: o suficiente para o download deixar de ser o gargalo sem virar rajada na API.
# No dashboard Thopen, 16 threads levaram 41 mil linhas de 67 s para 5 s; aqui a conta e' a
# mesma, com folga para a API respirar.
_THREADS = int(os.environ.get("BD_API_THREADS", "12"))

# workbook_key na API -> nome de arquivo que o `_base_espelho()` procura
ESPELHO = {
    "bd_performance":      "BD_Performance.xlsx",
    "bd_thopen":           "BD_Thopen.xlsx",
    "tickets_performance": "Tickets de Performance (atualizada).xlsx",
}

_lock = threading.Lock()


def _get(caminho: str, tentativas: int = 3):
    """GET com repetição. A rede da máquina oscila e uma falha isolada não pode abortar um
    espelho inteiro no meio — quem chama já está num laço longo."""
    erro = None
    for n in range(tentativas):
        try:
            req = urllib.request.Request(f"{API_BASE}{caminho}",
                                         headers={"Accept": "application/json"})
            with urllib.request.urlopen(req, timeout=90) as r:
                return json.loads(r.read().decode("utf-8"))
        except Exception as e:            # noqa: BLE001 - qualquer falha de rede/parse re-tenta
            erro = e
            time.sleep(2 * (n + 1))
    raise RuntimeError(f"API {caminho} falhou após {tentativas} tentativas: {erro}")


def workbooks() -> list:
    d = _get("/api/workbooks")
    return d if isinstance(d, list) else (d.get("items") or [])


def sheets() -> list:
    d = _get("/api/sheets")
    return d if isinstance(d, list) else (d.get("items") or [])


def linhas_da_aba(sheet_id: int) -> list:
    """Todas as linhas de uma aba, paginadas. Devolve os dicts crus da API (row_number/values)."""
    out, off = [], 0
    while True:
        d = _get(f"/api/sheets/{sheet_id}/rows?offset={off}&limit={_PAGINA}")
        got = d.get("rows") or []
        out.extend(got)
        if len(got) < _PAGINA:
            return out
        off += _PAGINA


# Versão do GERADOR do espelho. Entra na marca junto com o `updated_at` da API, e é o que faz um
# CONSERTO no código valer sem esperar a fonte mudar. Sem isto, espelho gerado com bug fica
# congelado para sempre: o `sincronizar` compara só a versão da fonte, vê "sem mudança" e nunca
# reexecuta o gerador — foi exatamente o que houve entre 25 e 26/08/2026. O conserto do cabeçalho
# (ver gerar_xlsx) já estava escrito, correto, e passou um dia sem NUNCA rodar, com a plataforma
# inteira lendo um espelho sem cabeçalho. BUMPAR sempre que mexer na lógica de geração.
_GERADOR_VER = 2


def gerar_xlsx(workbook_key: str, destino: str | None, abas_permitidas=None) -> dict:
    """Materializa um workbook da API num .xlsx. Devolve um resumo {abas, linhas, segundos}.

    `destino=None` devolve os BYTES do arquivo em `bytes` em vez de gravar em disco.

    Grava em temporário e só troca o arquivo no fim (os.replace, atômico) DEPOIS de reabrir com o
    openpyxl para provar que o arquivo é íntegro — mesmo cuidado do `/api/admin/base/<nome>`: um
    espelho truncado por cima de uma base boa derrubaria o carregamento inteiro da plataforma.
    """
    import openpyxl

    t0 = time.time()
    todas = [s for s in sheets() if s.get("workbook_key") == workbook_key]
    if not todas:
        raise RuntimeError(f"workbook '{workbook_key}' não tem abas na API")
    if abas_permitidas:
        todas = [s for s in todas if s.get("sheet_name") in abas_permitidas]

    # BAIXAR EM PARALELO, MONTAR EM SÉRIE. O download é espera de rede e escala com threads;
    # o openpyxl não é thread-safe, então a montagem continua sequencial. Medido em 30/08: as
    # três bases levavam 160 s em série (48 + 90 + 22) — tempo que o processo web pagaria a
    # cada TTL. Com 12 threads a mesma carga cai para poucos segundos, e é o que torna viável
    # o web ter as bases em memória em vez de ler o arquivo do worker.
    with ThreadPoolExecutor(max_workers=_THREADS) as ex:
        baixadas = dict(zip([s["id"] for s in todas],
                            ex.map(lambda s: linhas_da_aba(s["id"]), todas)))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)                      # a aba padrão não pertence a nenhuma planilha real
    total_linhas = 0
    for s in todas:
        ws = wb.create_sheet(title=str(s["sheet_name"])[:31])
        if str(s.get("visibility") or "").lower() == "hidden":
            ws.sheet_state = "hidden"
        escritas, hdr = set(), None
        for ln in baixadas[s["id"]]:
            r = ln.get("row_number")
            if not r:
                continue
            escritas.add(r)
            if hdr is None and ln.get("headers"):
                hdr = ln["headers"]
            for ci, v in enumerate(ln.get("values") or [], start=1):
                if v is not None:
                    ws.cell(row=r, column=ci, value=_valor(v))
            total_linhas += 1
        # A LINHA DE CABEÇALHO PRECISA EXISTIR NO ARQUIVO. Em 25/08 ~14:01 um re-sync mudou o
        # contrato da API: o cabeçalho deixou de vir como LINHA DE DADOS (a 1ª linha passou a ser
        # a de dados, ex. row_number=4 na Equipamentos) e os nomes ficaram só no campo `headers`.
        # O espelho gerado sem esta linha derrubou TODOS os leitores em silêncio — `header=2`
        # caía numa linha vazia, tudo virava coluna "Unnamed", o portão Full O&M desligava (146
        # usinas na tela) e o gerencial perdia as metas. Então: se a `header_row` que a API
        # declara não veio como dado, materializamos o cabeçalho nela — e se um sync futuro
        # voltar a mandá-lo como dado, nada é escrito em dobro.
        _hr = s.get("header_row")
        if hdr and _hr and _hr not in escritas:
            for ci, nome in enumerate(hdr, start=1):
                if nome not in (None, ""):
                    ws.cell(row=_hr, column=ci, value=str(nome))

    if destino is None:
        # Modo MEMÓRIA: devolve os bytes do .xlsx em vez de gravar. É o que permite ao processo
        # WEB ter as bases sem depender de arquivo nenhum — ver `carregar()`.
        buf = io.BytesIO()
        wb.save(buf)
        wb.close()
        dados = buf.getvalue()
        # mesma prova de integridade do disco: reabrir antes de entregar. Bytes truncados
        # derrubariam TODOS os leitores de uma vez, e em silêncio.
        _v = openpyxl.load_workbook(io.BytesIO(dados), read_only=True)
        n_abas = len(_v.sheetnames)
        _v.close()
        return {"abas": n_abas, "linhas": total_linhas, "bytes": dados,
                "segundos": round(time.time() - t0, 1)}

    os.makedirs(os.path.dirname(destino) or ".", exist_ok=True)
    # o temporário PRECISA terminar em .xlsx: o openpyxl recusa abrir ".tmp" pela extensão, e a
    # prova de integridade abaixo falharia sempre — publicando nunca, ou pior, publicando sem prova.
    tmp = destino + ".novo.xlsx"
    wb.save(tmp)
    wb.close()
    # prova de integridade antes de publicar
    try:
        _v = openpyxl.load_workbook(tmp, read_only=True)
        n_abas = len(_v.sheetnames)
        _v.close()
    except Exception as e:                    # noqa: BLE001
        os.remove(tmp)
        raise RuntimeError(f"espelho gerado saiu corrompido, NÃO publiquei: {e}") from e
    os.replace(tmp, destino)
    return {"abas": n_abas, "linhas": total_linhas, "segundos": round(time.time() - t0, 1)}


# ── Bases EM MEMÓRIA (o que tira o disco do caminho) ──────────────────────────
# O `_bd_api_loop` roda só no WORKER, e web e worker são PROCESSOS SEPARADOS: memória não se
# compartilha entre eles. Hoje o arquivo em `bases/` é o canal entre os dois. Com isto cada
# processo carrega a sua cópia direto da API — 2,9 MB para o BD_Performance, uma vez a cada
# TTL — e o web deixa de depender do worker (e de disco) para ter as bases.
#
# Guarda-se os BYTES do .xlsx, não um Workbook do openpyxl: os leitores do app.py já esperam
# "algo que se abre como arquivo", e um Workbook em memória custaria centenas de MB (no
# dashboard Thopen mediu-se 292 MB) contra 7,3 MB somando as três bases.
_MEM_TTL = int(os.environ.get("BD_API_MEM_TTL", "1800"))     # s
_mem = {}            # chave -> {"bytes":…, "versao":…, "ts":…, "abas":…, "linhas":…}
_mem_lock = threading.Lock()


def versao_em_memoria(workbook_key: str):
    """Marca de versão da carga atual (`updated_at` da API + versão do gerador), ou None.

    Serve para o chamador saber se uma revalidação MUDOU alguma coisa — comparar antes e depois
    é mais honesto do que o `carregar` devolver um booleano: ele tem três saídas possíveis
    (cache válido, revalidado sem mudança, rebaixado) e achatá-las em sim/não esconderia
    justamente o caso "a API não respondeu e estamos servindo a carga anterior"."""
    with _mem_lock:
        ent = _mem.get(workbook_key)
        return ent["versao"] if ent else None


def carregar(workbook_key: str, forcar: bool = False, revalidar: bool = False):
    """Bytes do .xlsx do workbook, de um cache em memória. None se a API não responder e não
    houver carga anterior.

    A revalidação é BARATA: uma chamada a /api/workbooks compara o `updated_at`. Só quando a
    fonte mudou é que as dezenas de milhares de linhas são baixadas de novo.

    Três modos, e a diferença entre eles importa:
      * padrão      — dentro do TTL devolve o cache SEM nem perguntar à API. É o do laço.
      * revalidar   — pula o atalho do TTL mas mantém a comparação de `updated_at`: pergunta
                      sempre, rebaixa só se mudou. É o do botão "Atualizar" (02/09/2026), que
                      precisa ser correto na hora e barato quando não houve alteração.
      * forcar      — rebaixa incondicionalmente, mesmo sem mudança. Só para depurar; usar isto
                      no botão faria 25 mil linhas trafegarem a cada clique, à toa.
    """
    with _mem_lock:
        ent = _mem.get(workbook_key)
        agora = time.time()
        if ent and not forcar and not revalidar and (agora - ent["ts"]) < _MEM_TTL:
            return ent["bytes"]
        try:
            w = {x.get("key"): x for x in workbooks()}.get(workbook_key)
            if not w:
                raise RuntimeError(f"workbook '{workbook_key}' ausente na API")
            versao = f"{w.get('updated_at') or ''}|g{_GERADOR_VER}"
            if ent and ent["versao"] == versao and not forcar:
                ent["ts"] = agora            # nada mudou na fonte: só renova o relógio
                return ent["bytes"]
            r = gerar_xlsx(workbook_key, None)
            _mem[workbook_key] = {"bytes": r["bytes"], "versao": versao, "ts": agora,
                                  "abas": r["abas"], "linhas": r["linhas"]}
            print(f"[bd_api] {workbook_key} em memória: {r['abas']} abas, {r['linhas']} linhas, "
                  f"{len(r['bytes'])/1e6:.1f} MB em {r['segundos']}s")
            return r["bytes"]
        except Exception as e:               # noqa: BLE001
            # Carga anterior é melhor que nada: a alternativa é a plataforma ficar cega porque
            # a API piscou. Quem chamou decide se cai para o arquivo.
            print(f"[bd_api] carga em memória de {workbook_key} falhou ({e})"
                  f"{' — mantendo a anterior' if ent else ''}")
            return ent["bytes"] if ent else None


def estado_memoria() -> dict:
    with _mem_lock:
        return {k: {"abas": v["abas"], "linhas": v["linhas"],
                    "MB": round(len(v["bytes"]) / 1e6, 1),
                    "idade_s": round(time.time() - v["ts"])} for k, v in _mem.items()}


def _marca_ler() -> dict:
    try:
        with io.open(_MARCA, encoding="utf-8") as f:
            return json.load(f)
    except Exception:                          # noqa: BLE001
        return {}


def _marca_gravar(d: dict):
    os.makedirs(_BASES_DIR, exist_ok=True)
    tmp = _MARCA + ".tmp"
    with io.open(tmp, "w", encoding="utf-8") as f:
        json.dump(d, f, ensure_ascii=False, indent=1)
    os.replace(tmp, _MARCA)


def sincronizar(force: bool = False) -> dict:
    """Regenera o espelho das planilhas cujo `updated_at` mudou na API.

    Barato quando nada mudou: uma chamada a /api/workbooks e pronto. O trabalho pesado (dezenas de
    milhares de linhas) só acontece quando a fonte de fato mudou.
    """
    with _lock:
        marca = _marca_ler()
        resumo = {}
        try:
            wbs = {w.get("key"): w for w in workbooks()}
        except Exception as e:                 # noqa: BLE001
            return {"erro": str(e)}
        for chave, arquivo in ESPELHO.items():
            w = wbs.get(chave)
            if not w:
                resumo[chave] = "ausente na API"
                continue
            # A marca é "versão da FONTE + versão do GERADOR": mudou qualquer um dos dois, refaz.
            versao = f"{w.get('updated_at') or ''}|g{_GERADOR_VER}"
            destino = os.path.join(_BASES_DIR, arquivo)
            if not force and marca.get(chave) == versao and os.path.exists(destino):
                resumo[chave] = "sem mudança"
                continue
            try:
                r = gerar_xlsx(chave, destino)
                marca[chave] = versao
                resumo[chave] = r
                print(f"[bd_api] {arquivo}: {r['abas']} abas, {r['linhas']} linhas "
                      f"em {r['segundos']}s (versão {versao})")
            except Exception as e:             # noqa: BLE001
                # NÃO limpa a marca nem o arquivo: espelho velho é melhor que espelho nenhum,
                # e a próxima rodada tenta de novo.
                resumo[chave] = f"falhou: {e}"
                print(f"[bd_api] {arquivo}: FALHOU — mantendo o espelho anterior. {e}")
        _marca_gravar(marca)
        return resumo


if __name__ == "__main__":
    import sys
    print(json.dumps(sincronizar(force="--force" in sys.argv), ensure_ascii=False, indent=1))
