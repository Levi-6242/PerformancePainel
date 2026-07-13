# -*- coding: utf-8 -*-
"""Dashboard Thopen — Performance da Usina (réplica do Power BI), fonte = BD_Thopen.xlsx.

Lê o BD_Thopen.xlsx sincronizado (o MESMO arquivo que o Power BI consome via SharePoint)
e reproduz o relatório "Performance da Usina":

  • Página Mensal: barras Produzida × Meta + tabela (Meta, Produzida, Diferença %,
    FC Meta, FC Real).  Fonte: tabela diária da usina (soma por mês) + `Historico_2026`
    (Meta/FC/PR) + `T_Usinas` (Potência MWp para o FC Real).
  • Página Diária: barras de energia/dia + linha de irradiação + linha de meta diária,
    KPIs (acumulado mês, irradiação real/meta, disponibilidade), produção anual
    (2023→2026, via `Historico`) e tabela de comentários.

Contas idênticas ao Power BI (ver as queries M de referência):
  • Produzida (MWh)  = Σ "Energia Produzida (kWh)" do mês / 1000
  • FC Real          = Produzida_kWh / (Pot_kWp × 24 × dias)   [dias = dia de hoje no
                       mês corrente, senão dias do mês] — espelha "Historico Atual (%)"
  • Diferença (%)    = (Produzida − Meta) / Meta
  • Meta diária      = Meta_mensal / dias_do_mês  (linha plana do gráfico diário)

Foco V1: Altair (mas o seletor lista todas as usinas que têm tabela diária).
Recarrega sozinho quando o xlsx muda (mtime). Porta 5080.
"""
import os
import json
import shutil
import tempfile
import threading
import datetime as dt
from calendar import monthrange

import openpyxl
from openpyxl.utils import range_boundaries
from flask import Flask, jsonify, render_template, request

# ── Localização dos dados ──────────────────────────────────────────────────────
# Na nuvem (Railway) definimos THOPEN_DATA_DIR=data → lê o SNAPSHOT das planilhas empacotado no repo
# (data/BD_Thopen.xlsx, data/Polaris, data/Matrix, data/Copel). Local, sem a env, segue lendo AO VIVO
# do OneDrive (dados sempre atuais). Atualizar a nuvem = novo push da pasta data/.
_DATA_DIR = os.environ.get("THOPEN_DATA_DIR")
if not _DATA_DIR and any(k.startswith("RAILWAY_") for k in os.environ):
    _DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")  # auto na nuvem Railway

_CANDIDATOS = [
    os.environ.get("BD_THOPEN_PATH"),
    os.path.join(_DATA_DIR, "BD_Thopen.xlsx") if _DATA_DIR else None,
    r"C:\Users\Levi Maia\OneDrive - GRID CO\Grid Co_ - 17. Acesso Externo Thopen"
    r"\1. Registro usinas Thopen\BD_Thopen.xlsx",
    r"C:\Users\Levi Maia\OneDrive - GRID CO\BD_Thopen.xlsx",
    r"C:\Users\Levi Maia\OneDrive - GRID CO\Área de Trabalho\BD_Thopen.xlsx",
]

ANO = 2026   # relatório do ano corrente (tabela Historico_2026 é 2026-específica)

# ── Comentários adicionais (campo digitável por usina, aba Diário) ───────────────
# ÚNICO dado que o app GRAVA (o resto é só leitura). Guardado num JSON {usina: texto}.
# Em servidor efêmero (Railway) o disco some a cada re-deploy → apontar COMENTARIOS_DIR para um
# VOLUME PERSISTENTE (ex.: /data). Sem a env, grava ao lado do script (ok local / máquina própria).
_COMENTARIOS_DIR = os.environ.get("COMENTARIOS_DIR") or os.path.dirname(os.path.abspath(__file__))
_COMENTARIOS_PATH = os.path.join(_COMENTARIOS_DIR, "comentarios_adicionais.json")
_com_lock = threading.Lock()


def _load_comentarios():
    try:
        with open(_COMENTARIOS_PATH, encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except (FileNotFoundError, ValueError, OSError):
        return {}


def _save_comentarios(data):
    with _com_lock:
        tmp = _COMENTARIOS_PATH + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        os.replace(tmp, _COMENTARIOS_PATH)   # troca atômica


def _bd_path():
    for c in _CANDIDATOS:
        if c and os.path.exists(c):
            return c
    raise FileNotFoundError("BD_Thopen.xlsx não encontrado (defina BD_THOPEN_PATH).")


def _planilha_em():
    """Data/hora de modificação do BD_Thopen.xlsx, formatada (ou None)."""
    try:
        return dt.datetime.fromtimestamp(os.path.getmtime(_bd_path())).strftime("%d/%m/%Y %H:%M")
    except OSError:
        return None


def _open_wb(path, tmpname="bd_thopen_dash.xlsx"):
    """Lê o workbook de uma CÓPIA temporária — assim funciona mesmo com o arquivo
    ABERTO no Excel ou sincronizando no OneDrive (a leitura direta dá PermissionError,
    mas o Windows permite copiar com leitura compartilhada). `tmpname` separa as cópias
    (BD_Thopen vs Budget Polaris) pra não se sobrescreverem."""
    tmp = os.path.join(tempfile.gettempdir(), tmpname)
    try:
        shutil.copy2(path, tmp)
        return openpyxl.load_workbook(tmp, data_only=True)
    except Exception:
        return openpyxl.load_workbook(path, data_only=True)   # fallback: leitura direta


# ── Cache do workbook por mtime ────────────────────────────────────────────────
_lock = threading.Lock()
_state = {"mtime": None, "path": None, "wb": None, "tbl": {}, "daily": {}, "df": {}}


def _ref_of(obj):
    return obj if isinstance(obj, str) else obj.ref


def _range_rows(ws, ref):
    c1, r1, c2, r2 = range_boundaries(ref)
    return list(ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2,
                             values_only=True))


def _header(ws, ref):
    c1, r1, c2, r2 = range_boundaries(ref)
    row = next(ws.iter_rows(min_row=r1, max_row=r1, min_col=c1, max_col=c2,
                            values_only=True), ())
    return [str(h).strip() if h is not None else "" for h in row]


def _usina_da_aba(ws, ref, hdr):
    """Nome OFICIAL da usina = valor da coluna 'Usina' DENTRO da aba (o dado real, que casa
    com as tabelas de meta/cadastro). É a chave correta — o título da aba é só um rótulo e
    às vezes vem abreviado/sem acento. Cai pro título da aba se não houver coluna 'Usina'."""
    iu = next((i for i, h in enumerate(hdr) if h.lower() == "usina"), None)
    if iu is None:
        return ws.title
    c1, r1, _c2, r2 = range_boundaries(ref)
    col = c1 + iu
    for row in ws.iter_rows(min_row=r1 + 1, max_row=min(r1 + 50, r2),
                            min_col=col, max_col=col, values_only=True):
        if row[0] not in (None, ""):
            return str(row[0]).strip()
    return ws.title


def _wb():
    """Workbook em cache; recarrega quando o arquivo muda. Indexa tabelas nomeadas
    e detecta as tabelas diárias (header com 'Energia Produzida')."""
    with _lock:
        path = _bd_path()
        m = os.path.getmtime(path)
        if _state["wb"] is None or _state["mtime"] != m or _state["path"] != path:
            try:
                wb = _open_wb(path)
            except Exception:
                if _state["wb"] is not None:
                    return _state["wb"]     # arquivo travado num instante: mantém a última leitura boa
                raise
            tbl, daily = {}, {}
            for ws in wb.worksheets:
                try:
                    names = list(ws.tables.keys())
                except Exception:
                    names = []
                for nm in names:
                    ref = _ref_of(ws.tables[nm])
                    tbl[nm] = (ws.title, ref)
                    hdr = _header(ws, ref)
                    hl = [h.lower() for h in hdr]
                    # aba diária = tem 'Energia Produzida' E 'Data' (exclui resumos como "Clientes");
                    # chave = nome da coluna 'Usina' (oficial), não o rótulo da aba
                    if any("energia produzida" in h for h in hl) and "data" in hl:
                        daily[_usina_da_aba(ws, ref, hdr)] = (ws.title, ref)
            _state.update(mtime=m, path=path, wb=wb, tbl=tbl, daily=daily, df={})
        return _state["wb"]


def _cols(rows):
    """rows[0] = header. Devolve (header_list, data_rows)."""
    if not rows:
        return [], []
    hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
    return hdr, rows[1:]


def _ci(hdr, *subs):
    """Índice da 1ª coluna cujo nome (minúsculo) contém TODOS os pedaços."""
    for i, h in enumerate(hdr):
        hl = h.lower()
        if all(s in hl for s in subs):
            return i
    return None


def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return None
    if "," in s:                       # formato pt-BR: 1.038.500,67
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _table(name):
    """(header, data_rows) de uma tabela nomeada, em cache."""
    wb = _wb()
    if name in _state["df"]:
        return _state["df"][name]
    out = ([], [])
    if name in _state["tbl"]:
        sheet, ref = _state["tbl"][name]
        out = _cols(_range_rows(wb[sheet], ref))
    _state["df"][name] = out
    return out


# ── Polaris: os ACTUALS (geração/irradiação/disponibilidade DIÁRIAS) vêm de um Excel de
#    Budget separado (alimentado semanalmente), NÃO do BD_Thopen. Meta, histórico (2023-2025)
#    e cadastro continuam vindo do BD_Thopen — igual às outras carteiras. ────────────────────
_POLARIS_DIR = (os.path.join(_DATA_DIR, "Polaris") if _DATA_DIR else
                r"C:\Users\Levi Maia\OneDrive - GRID CO\Grid Co_ - 17. Acesso Externo Thopen"
                r"\3. Polaris")
# nome no Budget (coluna "UFV 1")  →  nome de coleta canônico (igual T_Usinas/CARTEIRAS).
# De-para EXPLÍCITO (em vez da cadeia frágil de substituições do Power Query).
_POLARIS_NOME = {
    "UFV Aparecida do Taboado 2 1": "Aparecida do Taboado 1",
    "UFV Aparecida do Taboado 2 2": "Aparecida do Taboado 2",
    "UFV Aparecida III 1": "Aparecida 3",
    "UFV Araçoiaba da Serra I 1": "Araçoiaba da Serra 1",
    "UFV Araçoiaba da Serra I 2": "Araçoiaba da Serra 2",
    "UFV Boa Esperança do Sul I 1": "Boa Esperança do Sul 1",
    "UFV Boa Esperança do Sul I 2": "Boa Esperança do Sul 2",
    "UFV Caxambu I 1": "Caxambu",
    "UFV Goytacazes 1 1": "Goytacazes 1",
    "UFV Guaratinguetá 5 1": "Guaratinguetá V",
    "UFV Ibaté 1 1": "Ibaté 1",
    "UFV Ibaté 2 1": "Ibaté 2",
    "UFV Ipixuna 1 1": "Ipixuna 1",
    "UFV Ipixuna 2 1": "Ipixuna 2",
    "UFV Marajoara 1 1": "Marajoara 1",
    "UFV Piracicaba 1 1": "Piracicaba 1",
    "UFV Salto Pirapora 3 1": "Salto Pirapora 3",
    "UFV Santa Bárbara 1 1": "Santa Bárbara I",
    "UFV Santarem 1 1": "Santarém 1",
    "UFV Santarem 1 2": "Santarém 2",
    "UFV Santo Inácio 12 1": "Santo Inácio XII",
    "UFV São Bento 5 1": "São Bento V",
    "UFV Vargem Grande 1 1": "Vargem Grande 1",
    "UFV Araci 1 1": "Araci 1",
    "UFV Betânia 1 1": "Betânia 1",
    "UFV Boa Viagem 2 1": "Boa Viagem 2 1",
    "UFV Boa Viagem I 1": "Boa Viagem I 1",
    "UFV Ceará Mirim I 1": "Ceará Mirim I 1",
    "UFV Ceará Mirim I 2": "Ceará Mirim I 2",
    "UFV Delmiro Gouvea 1 1": "Delmiro Gouvea 1",
    "UFV Delmiro Gouvea 1 2": "Delmiro Gouvea 2",
    "UFV Delmiro Gouvea 1 3": "Delmiro Gouvea 3",
    "UFV Delmiro Gouvea 1 4": "Delmiro Gouvea 4",
    "UFV Goytacazes 4 2": "Goytacazes 4 2",
    "UFV Marajoara 2 1": "Marajoara 2 1",
    "UFV Piancó 1 1": "Piancó 1",
    "UFV Porteiras 1 1": "Porteiras 1",
    "UFV Porto Real 2 1": "Porto Real 2 1",
    "UFV Porto Real 3 1": "Porto Real 3",
    "UFV Urupês 1 1": "Urupês 1",
}
# lookup robusto (caixa-insensível) aceitando tanto o nome do Budget ("UFV …") quanto o já
# canônico — os Comentários Polaris usam o nome canônico ("Aparecida 3"), às vezes em CAIXA ALTA.
_POLARIS_LOOKUP = {}
for _k, _v in _POLARIS_NOME.items():
    _POLARIS_LOOKUP[_k.lower()] = _v
    _POLARIS_LOOKUP[_v.lower()] = _v


def _polaris_budget_path():
    """Budget Polaris mais recente (Budget_2025_UFVs_Raizen*.xlsx). Escolhe pela DATA no nome
    (…GridCo - AAAAMMDD.xlsx) e não por mtime — na nuvem o mtime é a hora do checkout, não a real."""
    import glob
    cands = glob.glob(os.path.join(_POLARIS_DIR, "Budget_2025_UFVs_Raizen*.xlsx"))
    return max(cands) if cands else None


def _polaris_coments():
    """{(usina, date): texto} a partir de 'Comentários Polaris.xlsx' (Tabela6)."""
    out = {}
    cpath = os.path.join(_POLARIS_DIR, "Comentários Polaris.xlsx")
    if not os.path.exists(cpath):
        return out
    try:
        wb = _open_wb(cpath, "polaris_coment_dash.xlsx")
        ws = wb["Comentários"]
        hdr, rows = _cols(_range_rows(ws, _ref_of(ws.tables["Tabela6"])))
    except Exception:
        return out
    iU = _ci(hdr, "ufv"); iD = _ci(hdr, "data"); iC = _ci(hdr, "coment")
    if None in (iU, iD, iC):
        return out
    for r in rows:
        u = _POLARIS_LOOKUP.get(str(r[iU]).strip().lower()) if r[iU] else None
        d = r[iD]
        if isinstance(d, dt.datetime):
            d = d.date()
        elif not isinstance(d, dt.date):                  # a data às vezes vem como texto "dd/mm/aaaa"
            try:
                d = dt.datetime.strptime(str(d).strip(), "%d/%m/%Y").date()
            except (ValueError, TypeError):
                d = None
        if u and d and r[iC]:
            txt = str(r[iC]).strip()
            out[(u, d)] = (out[(u, d)] + " | " + txt) if (u, d) in out else txt
    return out


def _polaris_records():
    """{usina_canônica: [{data, ger, ipoa, disp, com}]} a partir do Budget Polaris, no MESMO
    formato de `_daily_records`. Só o ano corrente (ANO), igual às abas diárias da Thopen.
    Cache invalidado pelo mtime do Budget."""
    path = _polaris_budget_path()
    if not path:
        return {}
    m = os.path.getmtime(path)
    cache = _state.get("polaris")
    if cache and cache.get("mtime") == m:
        return cache["recs"]
    wb = _open_wb(path, "polaris_budget_dash.xlsx")

    def wide(aba, tabela):
        """Tabela larga (UFV 1 | UFV 2 | <datas…>) → {usina: {date: valor}}."""
        ws = wb[aba]
        rows = _range_rows(ws, _ref_of(ws.tables[tabela]))
        if not rows:
            return {}
        header = rows[0]
        coldate = []
        for i in range(2, len(header)):
            h = header[i]
            d = h.date() if isinstance(h, dt.datetime) else (h if isinstance(h, dt.date) else None)
            if d is None and h is not None:
                try:
                    d = dt.datetime.strptime(str(h).strip(), "%d/%m/%Y").date()
                except ValueError:
                    d = None
            coldate.append((i, d))
        out = {}
        for r in rows[1:]:
            usina = _POLARIS_LOOKUP.get(str(r[0]).strip().lower()) if r[0] else None
            if not usina:
                continue
            s = out.setdefault(usina, {})
            for i, d in coldate:
                if d is None or d.year != ANO:
                    continue
                v = r[i]
                if isinstance(v, (int, float)):
                    s[d] = s.get(d, 0.0) + v
        return out

    ger = wide("Ger. Diaria", "GerDiaria")
    irr = wide("Irradiancia Diaria", "IrradDiaria")
    disp = wide("Disp. Diaria", "DispDiaria")
    com = _polaris_coments()
    recs = {}
    for u in (set(ger) | set(irr) | set(disp)):
        datas = sorted(set(ger.get(u, {})) | set(irr.get(u, {})) | set(disp.get(u, {})))
        lst = [{
            "data": d,
            "ger": ger.get(u, {}).get(d),
            "ipoa": irr.get(u, {}).get(d),
            "disp": disp.get(u, {}).get(d),
            "com": com.get((u, d)),
        } for d in datas]
        if lst:
            recs[u] = lst
    _state["polaris"] = {"mtime": m, "recs": recs}
    return recs


# ── Fontes de planilha externa (Matrix, Copel, …): contexto DIFERENCIADO ─────────────────────────
#    Os actuals diários (geração/irradiação/disp) de algumas carteiras NÃO vêm do BD_Thopen e sim de
#    um Excel separado. Regra de corte (_SHEET_CORTE = 01/06/2026):
#      • usina QUE EXISTE no BD_Thopen        → planilha p/ datas < corte + BD_Thopen p/ datas >= corte
#      • usina FORA do BD_Thopen (Caroá, Pharma II/III/IV) → planilha INTEIRA (todas as datas do ano)
#    Meta/histórico/cadastro sempre do BD_Thopen. Cada fonte tem 1 tabela por usina (Data/Usina/
#    geração/irradiação/disp) + 1 tabela larga de comentários. A coluna "Usina" já traz o canônico.
_THOPEN_EXT = r"C:\Users\Levi Maia\OneDrive - GRID CO\Grid Co_ - 17. Acesso Externo Thopen"
_MATRIX_DIR = os.path.join(_DATA_DIR, "Matrix") if _DATA_DIR else _THOPEN_EXT + r"\6. Matrix"
_COPEL_DIR = os.path.join(_DATA_DIR, "Copel") if _DATA_DIR else _THOPEN_EXT + r"\5. Copel"
_SHEET_CORTE = dt.date(2026, 6, 1)   # < corte: planilha | >= corte: BD_Thopen (se a usina existir lá)
_SHEET_SOURCES = [
    {"dir": _MATRIX_DIR, "glob": "Gera*Matrix*.xlsx", "tmp": "matrix_dash.xlsx"},
    {"dir": _COPEL_DIR, "glob": "Gera*Copel*.xlsx", "tmp": "copel_dash.xlsx"},
]
# nomes que o BD_Thopen grava diferente do cliente/planilha/meta → canoniza p/ tudo casar
_NOME_CANON = {"Santo Antonio do Platina": "Santo Antonio da Platina"}
_BD_ALIAS = {v: k for k, v in _NOME_CANON.items()}   # canônico -> nome da aba no BD_Thopen


def _sheet_path(src):
    import glob
    cands = [c for c in glob.glob(os.path.join(src["dir"], src["glob"]))
             if not os.path.basename(c).startswith("~")]
    return cands[0] if cands else None


def _sheet_coments(wb):
    """{(usina, date): texto} da tabela larga de comentários/ocorrências (Data + 1 coluna por usina).
    Remove o prefixo 'UFV ' (Copel) e o \\xa0 (Matrix) do nome da coluna."""
    out = {}
    alvo = None
    for w in wb.worksheets:
        for tn in (list(w.tables.keys()) if hasattr(w, "tables") else []):
            if "coment" in tn.lower() or "ocorr" in tn.lower():
                alvo = (w, tn)
                break
        if alvo:
            break
    if not alvo:
        return out
    ws, tn = alvo
    hdr, rows = _cols(_range_rows(ws, _ref_of(ws.tables[tn])))
    iD = _ci(hdr, "data")
    if iD is None:
        return out

    def _clean(h):
        h = h.replace("\xa0", " ").strip()
        return h[4:].strip() if h.upper().startswith("UFV ") else h

    cols = [(i, _clean(h)) for i, h in enumerate(hdr) if i != iD and h]
    for r in rows:
        dd = r[iD]
        if not isinstance(dd, (dt.datetime, dt.date)):
            continue
        day = dd.date() if isinstance(dd, dt.datetime) else dd
        for i, uname in cols:
            if r[i]:
                out[(uname, day)] = str(r[i]).strip()
    return out


def _sheet_records():
    """{usina_canônica: {date: {ger,ipoa,disp,com}}} unindo TODAS as fontes de planilha externa
    (Matrix, Copel, …). 1 tabela/usina em formato longo. Cache por mtimes das fontes."""
    paths = [(s, _sheet_path(s)) for s in _SHEET_SOURCES]
    paths = [(s, p) for s, p in paths if p]
    sig = tuple((p, os.path.getmtime(p)) for _, p in paths)
    cache = _state.get("sheets")
    if cache and cache.get("sig") == sig:
        return cache["recs"]
    out = {}
    for src, path in paths:
        wb = _open_wb(path, src["tmp"])
        for ws in wb.worksheets:
            for tn in (list(ws.tables.keys()) if hasattr(ws, "tables") else []):
                hdr, rows = _cols(_range_rows(ws, _ref_of(ws.tables[tn])))
                iD = _ci(hdr, "data"); iU = _ci(hdr, "usina")
                iG = _ci(hdr, "gera")
                if iG is None:
                    iG = _ci(hdr, "energia")   # Matrix usa "Geração"; Copel usa "Energia (kWh)"
                iI = _ci(hdr, "irradia"); iDp = _ci(hdr, "disp")
                if iD is None or iU is None:
                    continue   # pula a tabela de comentários (sem coluna "Usina")
                for r in rows:
                    u = str(r[iU]).strip() if r[iU] else None
                    dd = r[iD]
                    if not u or not isinstance(dd, (dt.datetime, dt.date)):
                        continue
                    u = _NOME_CANON.get(u, u)
                    day = dd.date() if isinstance(dd, dt.datetime) else dd
                    out.setdefault(u, {})[day] = {
                        "ger": _num(r[iG]) if iG is not None else None,
                        "ipoa": _num(r[iI]) if iI is not None else None,
                        "disp": _num(r[iDp]) if iDp is not None else None,
                        "com": None,
                    }
        for (u, day), txt in _sheet_coments(wb).items():
            u = _NOME_CANON.get(u, u)
            cel = out.setdefault(u, {}).setdefault(day, {"ger": None, "ipoa": None, "disp": None, "com": None})
            cel["com"] = txt
    _state["sheets"] = {"sig": sig, "recs": out}
    return out


# ── Leitura por domínio ─────────────────────────────────────────────────────────
def _daily_bd(usina):
    """Registros diários {data, ger, ipoa, disp, com} da tabela do BD_Thopen da usina."""
    key = ("recs", usina)
    if key in _state["df"]:
        return _state["df"][key]
    wb = _wb()
    entry = _state["daily"].get(usina)
    if entry is None and usina in _BD_ALIAS:
        entry = _state["daily"].get(_BD_ALIAS[usina])   # "Santo Antonio da Platina" -> aba "do Platina"
    recs = []
    if entry:
        sheet_title, ref = entry
        hdr, rows = _cols(_range_rows(wb[sheet_title], ref))
        iD = _ci(hdr, "data")
        iG = _ci(hdr, "energia produzida")
        iI = _ci(hdr, "ipoa")
        iDp = _ci(hdr, "disponibilidade", "usina")
        iC = _ci(hdr, "coment")
        for r in rows:
            d = r[iD] if iD is not None else None
            if not isinstance(d, (dt.datetime, dt.date)):
                continue
            com = None
            if iC is not None and r[iC] not in (None, ""):
                com = str(r[iC]).strip()
            recs.append({
                "data": d.date() if isinstance(d, dt.datetime) else d,
                "ger": _num(r[iG]) if iG is not None else None,
                "ipoa": _num(r[iI]) if iI is not None else None,
                "disp": _num(r[iDp]) if iDp is not None else None,
                "com": com,
            })
    _state["df"][key] = recs
    return recs


def _daily_records(usina):
    """{data, ger, ipoa, disp, com} por usina. Polaris = Budget; Matrix/Copel = planilha externa com
    corte (planilha < 01/06 + BD_Thopen >= 01/06 se a usina existir no BD; senão planilha inteira);
    o resto = BD_Thopen."""
    pol = _polaris_records()
    if usina in pol:
        return pol[usina]
    sheet = _sheet_records().get(usina)
    if sheet:
        bd = {r["data"]: r for r in _daily_bd(usina)}
        if bd:   # usina existe no BD_Thopen → planilha antes do corte + BD a partir do corte
            recs = [{"data": d, "ger": e.get("ger"), "ipoa": e.get("ipoa"),
                     "disp": e.get("disp"), "com": e.get("com")}
                    for d, e in sheet.items() if d.year == ANO and d < _SHEET_CORTE]
            recs += [r for r in bd.values() if r["data"] >= _SHEET_CORTE]
        else:    # usina fora do BD_Thopen → planilha inteira (todas as datas do ano)
            recs = [{"data": d, "ger": e.get("ger"), "ipoa": e.get("ipoa"),
                     "disp": e.get("disp"), "com": e.get("com")}
                    for d, e in sheet.items() if d.year == ANO]
        recs.sort(key=lambda r: r["data"])
        return recs
    return _daily_bd(usina)


def _meta2026(usina):
    """{mes: {meta(kWh), fc, metairr, pr}} para a usina (tabela Historico_2026)."""
    hdr, rows = _table("Historico_2026")
    out = {}
    if not hdr:
        return out
    iU = _ci(hdr, "usina")
    iMes = _ci(hdr, "mês") if _ci(hdr, "mês") is not None else _ci(hdr, "mes")
    iFC = _ci(hdr, "fc")
    iMeta = _ci(hdr, "meta (2026)")
    iIrr = _ci(hdr, "irradia")
    iPR = _ci(hdr, "pr")
    for r in rows:
        if iU is None or str(r[iU]).strip() != usina:
            continue
        mv = r[iMes]
        m = mv.month if isinstance(mv, (dt.datetime, dt.date)) else _num(mv)
        if not m:
            continue
        out[int(m)] = {
            "meta": _num(r[iMeta]) if iMeta is not None else None,
            "fc": _num(r[iFC]) if iFC is not None else None,
            "metairr": _num(r[iIrr]) if iIrr is not None else None,
            "pr": _num(r[iPR]) if iPR is not None else None,
        }
    return out


def _registro():
    """{usina: {nome, cidade, estado, cliente, pot_mwp}} (tabela T_Usinas)."""
    if ("reg",) in _state["df"]:
        return _state["df"][("reg",)]
    hdr, rows = _table("T_Usinas")
    out = {}
    if hdr:
        iU = _ci(hdr, "usina")
        iNome = _ci(hdr, "nome")
        iCid = _ci(hdr, "cidade")
        iEst = _ci(hdr, "estado")
        iCli = _ci(hdr, "cliente")
        iPot = _ci(hdr, "mwp")
        for r in rows:
            u = r[iU] if iU is not None else None
            if not u:
                continue
            out[str(u).strip()] = {
                "nome": r[iNome] if iNome is not None else None,
                "cidade": r[iCid] if iCid is not None else None,
                "estado": r[iEst] if iEst is not None else None,
                "cliente": r[iCli] if iCli is not None else None,
                "pot_mwp": _num(r[iPot]) if iPot is not None else None,
            }
    # Dado inconsistente no BD_Thopen: T_Usinas chama "Vargem Grande IB", mas meta/histórico/
    # relatório usam "Vargem Grande 1" (o de-para do Power BI normaliza p/ "Vargem Grande 1").
    # Alias p/ a potência/cadastro casar com o nome canônico que usamos.
    if "Vargem Grande IB" in out:
        out.setdefault("Vargem Grande 1", out["Vargem Grande IB"])
    _state["df"][("reg",)] = out
    return out


def _produzida_mensal(usina, ano):
    """{mes: kWh} de 'Produzida (<ano>)' (linhas COM mês) — para o comparativo mensal."""
    hdr, rows = _table("Historico")
    out = {}
    if not hdr:
        return out
    iU = _ci(hdr, "usina")
    iMes = _ci(hdr, "mês")
    iTipo = _ci(hdr, "tipo")
    iVal = _ci(hdr, "valor")
    alvo = str(ano)
    for r in rows:
        if iU is None or str(r[iU]).strip() != usina:
            continue
        tipo = str(r[iTipo]).strip().lower() if iTipo is not None else ""
        if "produzida" not in tipo or alvo not in tipo:
            continue
        mv = r[iMes] if iMes is not None else None
        if isinstance(mv, (dt.datetime, dt.date)):
            v = _num(r[iVal]) if iVal is not None else None
            if v is not None:
                out[mv.month] = out.get(mv.month, 0.0) + v
    return out


def _produzida_anual(usina):
    """{ano(str): kWh}. Usa o TOTAL anual explícito (linha com Mês vazio) quando existe;
    senão soma os meses. Evita contar em DOBRO quando há AMBOS (ex.: 2023/2024 ganharam
    linhas mensais mas mantiveram a linha de total anual)."""
    hdr, rows = _table("Historico")
    if not hdr:
        return {}
    iU = _ci(hdr, "usina")
    iAno = _ci(hdr, "ano")
    iMes = _ci(hdr, "mês")
    iTipo = _ci(hdr, "tipo")
    iVal = _ci(hdr, "valor")
    total, mensal = {}, {}
    for r in rows:
        if iU is None or str(r[iU]).strip() != usina:
            continue
        tipo = str(r[iTipo]).strip().lower() if iTipo is not None else ""
        if not tipo.startswith("produzida"):
            continue
        v = _num(r[iVal]) if iVal is not None else None
        if v is None:
            continue
        ano = str(r[iAno]).strip() if iAno is not None else ""
        mv = r[iMes] if iMes is not None else None
        if isinstance(mv, (dt.datetime, dt.date)):
            mensal[ano] = mensal.get(ano, 0.0) + v
        else:
            total[ano] = total.get(ano, 0.0) + v
    return {a: (total[a] if a in total else mensal.get(a)) for a in (set(total) | set(mensal))}


def _prod_mensal_ano(usina, ano):
    """{mes: kWh} de um ANO. Usinas de planilha externa (Matrix/Copel) puxam o histórico da
    PRÓPRIA planilha — não do BD Historico. É por isso que 'Produzida 2025' = o que a planilha
    registra em 2025 (só dez/2025, quando essas séries começam), igual ao Power BI; o BD Historico
    dessas usinas pode ter o ano inteiro real (ex.: Sarandi 2025 no BD = 11,6M) e NÃO deve ser usado."""
    e = _sheet_records().get(usina)
    if e is not None:
        out = {}
        for day, cel in e.items():
            if day.year == ano and cel.get("ger") is not None:
                out[day.month] = out.get(day.month, 0.0) + cel["ger"]
        return out
    return _produzida_mensal(usina, ano)


def _prod_anual_por_usina(usina):
    """{ano(str): kWh} p/ o gráfico anual. Usinas de planilha → soma da planilha (anos < ANO);
    demais → BD Historico."""
    if usina in _sheet_records():
        out = {}
        for y in (ANO - 3, ANO - 2, ANO - 1):
            mm = _prod_mensal_ano(usina, y)
            if mm:
                out[str(y)] = sum(mm.values())
        return out
    return _produzida_anual(usina)


# ── Endpoints ─────────────────────────────────────────────────────────────────
app = Flask(__name__)


# ── Carteiras (cliente dono do portfólio) → usinas pelo NOME DE COLETA (coluna "Usina") ──────
# De-para fixo aqui porque a coluna "Cliente" da aba "Dados Gerais Usinas" está vazia. Os nomes
# usados pelo cliente ("UFV …", "Ouro Branco X - Y", "Aparecida III"…) foram casados com a coluna
# de coleta via "Usina"/"Nome" do T_Usinas. Inclui usinas ainda NÃO coletadas (sem aba de dados):
# elas só aparecem no botão quando passarem a ter dados — o frontend filtra pela lista disponível.
CARTEIRAS = {
    "Thopen": [
        "Altair", "Alto Paraná 1", "Alto Paraná 2", "AP. do Taboado", "Areia Branca", "Aruanã",
        "Barretos", "Bernardino de Campos", "Brodowski", "Canarana 1", "Canarana 2", "Ceilândia 1",
        "Ceilândia 2", "Céu Azul", "Cidade Gaucha", "Colorado 1", "Colorado 2", "Coração 1",
        "Coração 2", "Embu Guaçu", "Fazenda Limão", "Fernandópolis", "Indaiatuba", "Junco",
        "Linhares", "Lyon", "Mandaguaçu", "Matão 1", "Matão 2", "Monte Aprazível", "Nova Iguaçu",
        "Nova Londrina", "Paranavaí", "Parelhas", "Poconé 1", "Primavera", "Ribeirão Cascalheiras",
        "Rodrigues", "Rondonópolis", "Sapopema", "Saturnino 1", "Senador", "Sitio Bonfim",
        "Sitio dos Nogueiras", "Sorocaba", "Tanabi",
    ],
    "Copel": [
        "Pharma II", "Pharma III", "Pharma IV", "Santo Antonio do Platina",
        "Santo Antonio da Platina", "Sarandi", "Segredo",
    ],
    "Matrix": [
        "Belo Jardim", "Caroá", "Inhapi", "Ouro Branco I", "Ouro Branco II", "Ouro Branco III",
        "Ouro Branco IV", "Ouro Branco V", "Santana do Ipanema", "São Bento do Una", "Vertentes",
    ],
    "Polaris": [
        "Aparecida do Taboado 1", "Aparecida do Taboado 2", "Aparecida 3", "Araci 1",
        "Araçoiaba da Serra 1", "Araçoiaba da Serra 2", "Betânia 1", "Boa Esperança do Sul 1",
        "Boa Esperança do Sul 2", "Boa Viagem 2 1", "Boa Viagem I 1", "Caxambu", "Ceará Mirim I 1",
        "Ceará Mirim I 2", "Delmiro Gouvea 1", "Delmiro Gouvea 2", "Delmiro Gouvea 3",
        "Delmiro Gouvea 4", "Goytacazes 1", "Goytacazes 4 2", "Guaratinguetá V", "Ibaté 1",
        "Ibaté 2", "Ipixuna 1", "Ipixuna 2", "Marajoara 1", "Marajoara 2 1", "Piancó 1",
        "Piracicaba 1", "Porteiras 1", "Porto Real 2 1", "Porto Real 3", "Salto Pirapora 3",
        "Santa Bárbara I", "Santarém 1", "Santarém 2", "Santo Inácio XII", "São Bento V",
        "Urupês 1", "Vargem Grande 1",
    ],
}
CARTEIRA_ORDEM = ["Thopen", "Copel", "Matrix", "Polaris"]
_CARTEIRA_DE = {u: c for c in CARTEIRA_ORDEM for u in CARTEIRAS[c]}  # usina -> carteira


@app.route("/api/t/usinas")
def usinas():
    _wb()
    us = sorted({_NOME_CANON.get(u, u) for u in
                 (set(_state["daily"].keys()) | set(_polaris_records().keys()) | set(_sheet_records().keys()))})
    default = "Altair" if "Altair" in us else (us[0] if us else None)
    carteira_de = {u: _CARTEIRA_DE.get(u) for u in us}  # carteira de cada usina disponível
    today = dt.date.today()
    mes_max = today.month if today.year == ANO else 12  # p/ o filtro de mês GLOBAL do frontend
    return jsonify({"usinas": us, "default": default,
                    "carteiras": CARTEIRA_ORDEM, "carteira_de": carteira_de,
                    "ano": ANO, "mes_max": mes_max})


@app.route("/api/t/overview")
def overview():
    usina = request.args.get("usina", "Altair")
    recs = _daily_records(usina)
    reg = _registro().get(usina, {})
    pot_mwp = reg.get("pot_mwp")
    meta = _meta2026(usina)
    today = dt.date.today()

    prod = {}
    for r in recs:
        if r["data"].year == ANO and r["ger"] is not None:
            prod[r["data"].month] = prod.get(r["data"].month, 0.0) + r["ger"]
    p2023 = _prod_mensal_ano(usina, 2023)  # anos anteriores, mês a mês (planilha p/ Matrix/Copel)
    p2024 = _prod_mensal_ano(usina, 2024)
    p2025 = _prod_mensal_ano(usina, 2025)

    meses = []
    for m in range(1, 13):
        mt = meta.get(m, {})
        meta_kwh = mt.get("meta")
        pz = prod.get(m)
        pz23 = p2023.get(m)
        pz24 = p2024.get(m)
        pz25 = p2025.get(m)
        dias = today.day if (ANO == today.year and m == today.month) else monthrange(ANO, m)[1]
        fc_real = (pz / (pot_mwp * 1000 * 24 * dias)) if (pz is not None and pot_mwp) else None
        dif = ((pz - meta_kwh) / meta_kwh) if (pz is not None and meta_kwh) else None
        meses.append({
            "mes": m,
            "meta": (meta_kwh / 1000) if meta_kwh else None,
            "produzida": (pz / 1000) if pz is not None else None,
            "produzida_2023": (pz23 / 1000) if pz23 else None,
            "produzida_2024": (pz24 / 1000) if pz24 else None,
            "produzida_2025": (pz25 / 1000) if pz25 else None,
            "dif": dif,
            "fc_meta": mt.get("fc"),
            "fc_real": fc_real,
        })

    pa = _prod_anual_por_usina(usina)
    prod2026 = sum(prod.values())
    # Meta 2026 = meta acumulada (YTD) dos meses em que a usina gerou — espelha a medida DAX
    # do Power BI (SUM(meta) WHERE Date <= mês) p/ usinas contínuas (casa exato c/ Altair).
    # Acumula SÓ nos meses com produção de propósito: não cobra meta de meses em que a usina
    # ainda não existia (usinas que entraram no meio do ano). Ver nota p/ alternar p/ calendário.
    meta2026_ytd = sum((meta.get(m, {}).get("meta") or 0) for m in prod)
    anual = [
        {"label": "Produzida 2023", "valor": pa.get("2023")},
        {"label": "Produzida 2024", "valor": pa.get("2024")},
        {"label": "Produzida 2025", "valor": pa.get("2025")},
        {"label": "Meta 2026", "valor": meta2026_ytd or None},
        {"label": "Produzida 2026", "valor": prod2026 or None},
    ]

    per = sorted({(r["data"].year, r["data"].month) for r in recs if r["ger"] is not None})
    periodos = [{"ano": a, "mes": mm} for a, mm in per]

    return jsonify({
        "usina": usina, "nome": reg.get("nome") or usina, "cidade": reg.get("cidade"),
        "estado": reg.get("estado"), "cliente": reg.get("cliente"),
        "potencia_mwp": pot_mwp, "ano": ANO, "meses": meses,
        "anual": anual, "periodos": periodos, "planilha_em": _planilha_em(),
        "comentario_extra": _load_comentarios().get(usina, ""),
    })


@app.route("/api/t/comentario_extra", methods=["POST"])
def comentario_extra():
    """Salva/atualiza o texto livre 'Comentários adicionais' de uma usina (aba Diário)."""
    data = request.get_json(force=True, silent=True) or {}
    usina = (data.get("usina") or "").strip()
    texto = (data.get("texto") or "").strip()[:5000]   # teto p/ evitar abuso (campo é aberto)
    if not usina:
        return jsonify({"ok": False, "erro": "usina não informada"}), 400
    com = _load_comentarios()
    if texto:
        com[usina] = texto
    else:
        com.pop(usina, None)   # texto vazio = apaga a entrada
    _save_comentarios(com)
    return jsonify({"ok": True})


@app.route("/api/t/diario")
def diario():
    usina = request.args.get("usina", "Altair")
    ano = int(request.args.get("ano", ANO))
    mes = int(request.args.get("mes", 1))
    recs = [r for r in _daily_records(usina)
            if r["data"].year == ano and r["data"].month == mes]
    recs.sort(key=lambda r: r["data"])

    meta = _meta2026(usina).get(mes, {}) if ano == ANO else {}
    dias_mes = monthrange(ano, mes)[1]
    meta_kwh = meta.get("meta")
    meta_dia = (meta_kwh / dias_mes) if meta_kwh else None

    pontos = [{"dia": r["data"].day, "ger": r["ger"], "ipoa": r["ipoa"], "com": r["com"]}
              for r in recs]
    disp_vals = [r["disp"] for r in recs if r["disp"] is not None]
    kpis = {
        "acum_prod": sum((r["ger"] or 0) for r in recs),
        "acum_meta": meta_kwh,
        "irr_real": round(sum((r["ipoa"] or 0) for r in recs), 2),
        "irr_meta": meta.get("metairr"),
        "disp": (sum(disp_vals) / len(disp_vals)) if disp_vals else None,
    }
    # Agrupa dias CONSECUTIVOS com o MESMO comentário num intervalo
    # (ex.: "01/05/2026 a 06/05/2026"). Dia sem comentário ou texto diferente quebra a sequência.
    grupos = []
    prev_d = None
    for r in recs:
        c, d = r["com"], r["data"]
        if c:
            if (grupos and grupos[-1]["com"] == c and prev_d is not None
                    and (d - prev_d).days == 1 and grupos[-1]["fim"] == prev_d):
                grupos[-1]["fim"] = d
            else:
                grupos.append({"ini": d, "fim": d, "com": c})
        prev_d = d

    def _drange(g):
        a = g["ini"].strftime("%d/%m/%Y")
        b = g["fim"].strftime("%d/%m/%Y")
        return a if a == b else a + " a " + b

    coment = [{"data": _drange(g), "com": g["com"]} for g in grupos]

    return jsonify({"usina": usina, "ano": ano, "mes": mes, "pontos": pontos,
                    "meta_dia": meta_dia, "kpis": kpis, "comentarios": coment})


def _resumo_usina(usina, ano, mes):
    """Linha do resumo (Visão Geral) de uma usina num mês — espelha as medidas DAX:
    Produzida/Meta/Irradiação/Disp do mês + FC = Produzida / (Pot_kWp × 24 × dias)."""
    recs = [r for r in _daily_records(usina) if r["data"].year == ano and r["data"].month == mes]
    gers = [r["ger"] for r in recs if r["ger"] is not None]
    ipoas = [r["ipoa"] for r in recs if r["ipoa"] is not None]
    disps = [r["disp"] for r in recs if r["disp"] is not None]
    produzida = sum(gers) if gers else None
    irr_med = sum(ipoas) if ipoas else None
    disp = (sum(disps) / len(disps)) if disps else None
    mt = _meta2026(usina).get(mes, {}) if ano == ANO else {}
    meta, irr_esp = mt.get("meta"), mt.get("metairr")
    pot_kwp = (_registro().get(usina, {}).get("pot_mwp") or 0) * 1000 or None
    today = dt.date.today()
    dias = today.day if (ano == today.year and mes == today.month) else monthrange(ano, mes)[1]
    fc = (produzida / (pot_kwp * 24 * dias)) if (produzida is not None and pot_kwp) else None
    dif_prod = ((produzida - meta) / meta) if (produzida is not None and meta) else None
    dif_irr = ((irr_med - irr_esp) / irr_esp) if (irr_med and irr_esp) else None
    return {"usina": usina, "pot_kwp": pot_kwp, "meta": meta, "produzida": produzida,
            "dif_prod": dif_prod, "irr_esp": irr_esp, "irr_med": irr_med, "dif_irr": dif_irr,
            "disp": disp, "fc": fc}


@app.route("/api/t/geral")
def geral():
    """Resumo (Visão Geral) de uma CARTEIRA num mês de referência: 1 linha por usina em operação,
    + a data de corte (p/ a capa). Espelha o relatório 'Performance UFVs - Geração'."""
    carteira = request.args.get("carteira", "Thopen")
    ano = int(request.args.get("ano", ANO))
    mes = int(request.args.get("mes", dt.date.today().month))
    _wb()
    disponiveis = {_NOME_CANON.get(u, u) for u in
                   (set(_state["daily"].keys()) | set(_polaris_records().keys())
                    | set(_sheet_records().keys()))}
    nomes = sorted({u for u in CARTEIRAS.get(carteira, []) if u in disponiveis})
    linhas = [r for r in (_resumo_usina(u, ano, mes) for u in nomes) if r["produzida"] is not None]
    today = dt.date.today()
    corte = (today if (ano == today.year and mes == today.month)
             else dt.date(ano, mes, monthrange(ano, mes)[1]))
    nome_carteira = next((c for c in CARTEIRA_ORDEM if c == carteira), carteira)
    return jsonify({"carteira": nome_carteira, "ano": ano, "mes": mes,
                    "corte": corte.strftime("%d/%m/%Y"), "linhas": linhas})


@app.route("/api/t/reload")
def reload_bd():
    """Força reler o BD_Thopen.xlsx (limpa o cache) — usado pelo botão Atualizar."""
    with _lock:
        _state["wb"] = None
        _state["mtime"] = None
        _state["df"] = {}
        _state["polaris"] = None   # força reler o Budget Polaris também
        _state["sheets"] = None    # e as planilhas externas (Matrix, Copel)
    _wb()  # recarrega agora (chamado FORA do lock — _wb() readquire o lock)
    return jsonify({"ok": True, "planilha_em": _planilha_em(),
                    "atualizado_em": dt.datetime.now().strftime("%H:%M:%S")})


@app.route("/")
def index():
    return render_template("dashboard_thopen.html")


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5080, debug=False, threaded=True)
